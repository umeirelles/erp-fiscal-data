"""Testes de apply_invoice_mapping: do relatório de mapeamento à planilha entregue.

Os relatórios de teste são montados por ``mapping_rows()`` com nomes e valores
internos e gravados no esquema da vez pela fixture ``write_mapping``, que é
parametrizada nos dois esquemas: o legado em inglês, que é o dos relatórios
arquivados (inclusive os 26 pareceres de sap_procosa), e o em português, montado
com ``report_schema.to_external_frame``. Assim todo teste que grava um relatório
roda nos dois. O entregável sai em português nos dois casos.
"""

from __future__ import annotations

import re
from datetime import time

import openpyxl
import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import report_schema
from apply_invoice_mapping import (
    CONFIDENCE_ROWS,
    FIRST_DATA_EXCEL_ROW,
    MAPPING_METADATA,
    STATE_ABSENT,
    STATE_DELIVERED,
    STATE_PENDING,
    _describe_header_conflict_pt,
    build_mapped_workbook,
    build_parser,
    main,
)
from invoice_mapper import semantic_header_conflict

MAPPED_SHEET = "Dados Mapeados"
COVERAGE_SHEET = "Cobertura"


def write_workbook(path, headers: list[str], rows: list[list[object]]) -> None:
    pd.DataFrame([headers, *rows]).to_excel(path, header=False, index=False)


def mapping_rows() -> pd.DataFrame:
    """Relatório de três linhas, com nomes e valores internos (iguais aos do legado)."""
    return pd.DataFrame(
        [
            {
                "model_column_number": 1,
                "model_column": "CFOP",
                "client_column_number": 2,
                "client_column": "CFOP source",
                "name_confidence": 0.98,
                "name_runner_up_confidence": 0.70,
                "name_margin": 0.28,
                "data_confidence": None,
                "llm_data_confidence": None,
                "llm_final_confidence": None,
                "final_confidence": 0.98,
                "method": "name",
                "status": "accepted",
                "reason": "Clear header match.",
                "reviewer_status": "",
                "correct_client_column": "",
                "reviewer_note": "",
            },
            {
                "model_column_number": 2,
                "model_column": "Valor COFINS",
                "client_column_number": 3,
                "client_column": "COFI value",
                "name_confidence": 0.82,
                "name_runner_up_confidence": 0.80,
                "name_margin": 0.02,
                "data_confidence": 0.91,
                "llm_data_confidence": 0.93,
                "llm_final_confidence": 0.89,
                "final_confidence": 0.88,
                "method": "data",
                "status": "review",
                "reason": "Names and numeric values are compatible.",
                "reviewer_status": "",
                "correct_client_column": "",
                "reviewer_note": "Check this value field.",
            },
            {
                "model_column_number": 3,
                "model_column": "Nº do NCM",
                "client_column_number": None,
                "client_column": None,
                "name_confidence": 0.40,
                "name_runner_up_confidence": 0.38,
                "name_margin": 0.02,
                "data_confidence": 0.10,
                "llm_data_confidence": 0.12,
                "llm_final_confidence": 0.08,
                "final_confidence": 0.0,
                "method": "data",
                "status": "no_match",
                "reason": "No automatic match.",
                "reviewer_status": "incorrect",
                "correct_client_column": "Corrected NCM",
                "reviewer_note": "Reviewer supplied the source column.",
            },
        ]
    )


def external_report(frame: pd.DataFrame, schema: str) -> pd.DataFrame:
    """Relatório com nomes internos -> como ele fica gravado no esquema dado."""
    if schema == report_schema.PT_BR:
        return report_schema.to_external_frame(frame, "mappings")
    return frame.copy()


def save_report(path, report: pd.DataFrame, schema: str, metadata: dict | None = None) -> None:
    """Grava um relatório já no esquema dado, com a aba de metadados se houver."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        report.to_excel(
            writer, sheet_name=report_schema.sheet_name("mappings", schema), index=False
        )
        if metadata is not None:
            if schema == report_schema.PT_BR:
                frame = report_schema.to_external_metadata(metadata)
            else:
                frame = pd.DataFrame(list(metadata.items()), columns=["key", "value"])
            frame.to_excel(
                writer, sheet_name=report_schema.sheet_name("run_metadata", schema), index=False
            )


@pytest.fixture(
    params=[report_schema.LEGACY_EN, report_schema.PT_BR], ids=["legado_en", "pt_br"]
)
def schema(request) -> str:
    """Esquema em que o relatório de teste é gravado."""
    return request.param


@pytest.fixture
def write_mapping(schema):
    """Grava um relatório montado com nomes internos no esquema parametrizado."""

    def write(frame: pd.DataFrame, path, metadata: dict | None = None) -> None:
        save_report(path, external_report(frame, schema), schema, metadata)

    return write


def label(field: str, schema: str) -> str:
    """Cabeçalho de Mapeamentos no esquema dado, como as mensagens de erro o citam."""
    return report_schema.header_label("mappings", field, schema)


def banner_row(field: str) -> int:
    """Índice (base 0, no quadro lido sem cabeçalho) da linha do banner de um campo."""
    return 1 + [name for _, name in MAPPING_METADATA].index(field)


def banner_excel_row(field: str) -> int:
    """Linha do Excel (base 1) da linha do banner de um campo."""
    return 2 + [name for _, name in MAPPING_METADATA].index(field)


def test_build_mapped_workbook_defaults_to_confirmed_mappings(tmp_path, write_mapping) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"

    write_workbook(
        model_path,
        ["CFOP", "Valor COFINS", "Nº do NCM"],
        [["5102", 10.0, "85167910"]],
    )
    write_workbook(
        client_path,
        ["Other", "CFOP source", "COFI value", "Corrected NCM"],
        [["A", "5102", 12.5, "85167910"], ["B", "6102", 7.5, "73211100"]],
    )
    write_mapping(mapping_rows(), mapping_path)

    build_mapped_workbook(
        client_path=client_path,
        mapping_path=mapping_path,
        output_path=output_path,
        model_path=model_path,
    )

    raw = pd.read_excel(output_path, sheet_name=MAPPED_SHEET, header=None, dtype=object)
    assert raw.iloc[0].tolist() == [
        "__info_mapeamento__",
        "CFOP",
        "Valor COFINS",
        "Nº do NCM",
    ]
    assert raw.iloc[1 : len(MAPPING_METADATA) + 1, 0].tolist() == [
        label for label, _ in MAPPING_METADATA
    ]

    client_column_row = banner_row("client_column")
    assert raw.iloc[client_column_row, 1] == "CFOP source"
    assert pd.isna(raw.iloc[client_column_row, 2])
    assert raw.iloc[client_column_row, 3] == "Corrected NCM"

    first_data_row = len(MAPPING_METADATA) + 1
    assert raw.iloc[first_data_row, 1] == "5102"
    assert pd.isna(raw.iloc[first_data_row, 2])
    assert raw.iloc[first_data_row, 3] == "85167910"
    assert raw.iloc[first_data_row + 1, 1] == "6102"
    assert pd.isna(raw.iloc[first_data_row + 1, 2])
    assert raw.iloc[first_data_row + 1, 3] == "73211100"

    workbook = load_workbook(output_path)
    worksheet = workbook[MAPPED_SHEET]
    assert worksheet.freeze_panes == f"B{len(MAPPING_METADATA) + 2}"
    assert worksheet.cell(banner_excel_row("name_confidence"), 2).number_format == "0.00%"
    llm_confidence_excel_row = banner_excel_row("llm_final_confidence")
    assert worksheet.cell(llm_confidence_excel_row, 3).value == 0.89
    assert worksheet.cell(llm_confidence_excel_row, 3).number_format == "0.00%"


def test_banner_labels_are_the_portuguese_headers_and_confidences_are_keyed_by_field() -> None:
    # O rótulo da coluna A é o cabeçalho em português do campo; a formatação
    # de porcentagem é escolhida pelo campo, então trocar o texto do rótulo
    # não pode quebrá-la.
    labels = {field: text for text, field in MAPPING_METADATA}
    assert labels == report_schema.COLUMNS["banner"]
    assert CONFIDENCE_ROWS <= {field for _, field in MAPPING_METADATA}
    assert FIRST_DATA_EXCEL_ROW == 23


def test_deliverable_is_written_in_portuguese_from_either_schema(tmp_path, write_mapping) -> None:
    # O mesmo relatório, legado ou em português, dá o mesmo entregável em
    # português: valores de máquina e etapas traduzidos, texto livre intacto.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    mapping = mapping_rows()
    mapping["catalog_name"] = ["cfop", "", "ibge_municipality"]
    mapping["decided_by"] = ["lexical", "resolver", ""]
    mapping["supported_by"] = ["lexical + catalog (agrees)", "resolver + profile", ""]
    mapping["vetoed_by"] = ["", "", "vacuity"]
    write_mapping(mapping, mapping_path)

    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=output_path, model_path=model_path)

    book = openpyxl.load_workbook(output_path)
    assert book.sheetnames == [MAPPED_SHEET, COVERAGE_SHEET]
    sheet = book[MAPPED_SHEET]
    assert sheet["A1"].value == "__info_mapeamento__"

    def banner(field: str) -> list[object]:
        return [sheet.cell(banner_excel_row(field), column).value for column in (2, 3, 4)]

    assert banner("status") == ["aceito", "revisao", "sem_correspondencia"]
    assert banner("method") == ["nome", "dados", "dados"]
    assert banner("catalog_name") == ["cfop", None, "municipio_ibge"]
    assert banner("decided_by") == ["lexico", "resolvedor_ia", None]
    assert banner("supported_by") == [
        "lexico + catalogo (concorda)", "resolvedor_ia + perfil", None,
    ]
    assert banner("vetoed_by") == [None, None, "coluna_vazia"]
    assert banner("reviewer_status") == [None, None, "incorreto"]
    # Texto livre (justificativa e observação do revisor) é copiado como está.
    assert banner("reason") == mapping["reason"].tolist()
    assert banner("reviewer_note") == [None, "Check this value field.",
                                       "Reviewer supplied the source column."]


def test_legacy_report_prints_a_one_line_migration_notice(
    tmp_path, write_mapping, schema, capsys
) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    write_mapping(mapping_rows(), mapping_path)

    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=tmp_path / "mapped.xlsx", model_path=model_path)

    captured = capsys.readouterr()
    assert "colunas_modelo=3 linhas_cliente=1 colunas_mapeadas=2" in captured.out
    if schema == report_schema.LEGACY_EN:
        notice = captured.err.strip()
        assert "\n" not in notice
        assert "mapping.xlsx" in notice and "migrate_reports.py" in notice
    else:
        assert captured.err == ""


def test_include_review_populates_unconfirmed_proposal(tmp_path, write_mapping) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(
        model_path,
        ["CFOP", "Valor COFINS", "Nº do NCM"],
        [["5102", 10.0, "85167910"]],
    )
    write_workbook(
        client_path,
        ["Other", "CFOP source", "COFI value", "Corrected NCM"],
        [["A", "5102", 12.5, "85167910"]],
    )
    write_mapping(mapping_rows(), mapping_path)

    build_mapped_workbook(
        client_path=client_path,
        mapping_path=mapping_path,
        output_path=output_path,
        model_path=model_path,
        accepted_only=False,
    )

    raw = pd.read_excel(output_path, sheet_name=MAPPED_SHEET, header=None, dtype=object)
    first_data_row = len(MAPPING_METADATA) + 1
    assert raw.iloc[first_data_row, 2] == 12.5


def test_include_review_allows_duplicate_unconfirmed_proposals(tmp_path, write_mapping) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["Base CBS", "Base IBS"], [[100.0, 100.0]])
    write_workbook(client_path, ["Generic base"], [[125.0]])
    mapping = mapping_rows().iloc[:2].copy()
    mapping["model_column_number"] = [1, 2]
    mapping["model_column"] = ["Base CBS", "Base IBS"]
    mapping["client_column_number"] = [1, 1]
    mapping["client_column"] = ["Generic base", "Generic base"]
    mapping["status"] = ["review", "review"]
    mapping["reviewer_status"] = ["", ""]
    write_mapping(mapping, mapping_path)

    build_mapped_workbook(
        client_path=client_path,
        mapping_path=mapping_path,
        output_path=output_path,
        model_path=model_path,
        accepted_only=False,
    )

    raw = pd.read_excel(output_path, sheet_name=MAPPED_SHEET, header=None, dtype=object)
    first_data_row = len(MAPPING_METADATA) + 1
    assert raw.iloc[first_data_row, 1:].tolist() == [125, 125]


def test_mapped_workbook_preserves_excel_time_values(tmp_path, write_mapping) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["Criado às:"], [[time(8, 0, 0)]])
    write_workbook(client_path, ["Created at"], [[None]])
    client_workbook = load_workbook(client_path)
    client_workbook.active["A2"] = time(15, 30, 11)
    client_workbook.active["A2"].number_format = "hh:mm:ss"
    client_workbook.save(client_path)
    mapping = mapping_rows().iloc[[0]].copy()
    mapping["model_column"] = "Criado às:"
    mapping["client_column_number"] = 1
    mapping["client_column"] = "Created at"
    write_mapping(mapping, mapping_path)

    build_mapped_workbook(
        client_path=client_path,
        mapping_path=mapping_path,
        output_path=output_path,
        model_path=model_path,
    )

    worksheet = load_workbook(output_path, data_only=True)[MAPPED_SHEET]
    cell = worksheet.cell(len(MAPPING_METADATA) + 2, 2)
    assert cell.value == time(15, 30, 11)
    assert cell.number_format == "hh:mm:ss"


def test_accepted_only_leaves_unreviewed_review_mapping_blank(tmp_path, write_mapping) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(
        model_path,
        ["CFOP", "Valor COFINS", "Nº do NCM"],
        [["5102", 10.0, "85167910"]],
    )
    write_workbook(
        client_path,
        ["Other", "CFOP source", "COFI value", "Corrected NCM"],
        [["A", "5102", 12.5, "85167910"]],
    )
    mapping = mapping_rows()
    mapping.loc[2, "reviewer_status"] = "no_mapping_exists"
    # A linha do fixture traz uma correção para o parecer 'incorrect' que ela
    # tinha. Só 'incorrect' usa correção, então ela sai junto com o parecer.
    mapping.loc[2, "correct_client_column"] = ""
    write_mapping(mapping, mapping_path)

    build_mapped_workbook(
        client_path=client_path,
        mapping_path=mapping_path,
        output_path=output_path,
        model_path=model_path,
        accepted_only=True,
    )

    raw = pd.read_excel(output_path, sheet_name=MAPPED_SHEET, header=None, dtype=object)
    first_data_row = len(MAPPING_METADATA) + 1
    assert raw.iloc[first_data_row, 1] == "5102"
    assert pd.isna(raw.iloc[first_data_row, 2])
    assert pd.isna(raw.iloc[first_data_row, 3])


@pytest.mark.parametrize("typed", ["incorrect", "Incorreto", " INCORRETO ", "incorreto"])
def test_reviewer_status_is_accepted_in_either_language(tmp_path, schema, typed) -> None:
    # O parecer é digitado por pessoas: vale nas duas línguas, com qualquer
    # caixa e espaços nas pontas, num relatório de qualquer esquema.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    report = external_report(mapping_rows(), schema)
    report.loc[2, label("reviewer_status", schema)] = typed
    save_report(mapping_path, report, schema)

    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=output_path, model_path=model_path)

    sheet = load_workbook(output_path)[MAPPED_SHEET]
    assert sheet.cell(FIRST_DATA_EXCEL_ROW, 4).value == "85167910"
    assert sheet.cell(banner_excel_row("reviewer_status"), 4).value == "incorreto"


def test_unknown_reviewer_status_is_refused(tmp_path, schema) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    report = external_report(mapping_rows(), schema)
    report.loc[0, label("reviewer_status", schema)] = "talvez"
    save_report(mapping_path, report, schema)

    with pytest.raises(
        ValueError, match=f"Valor desconhecido em {label('reviewer_status', schema)}: 'talvez'"
    ):
        build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                              output_path=tmp_path / "out.xlsx", model_path=model_path)


def test_incorrect_review_requires_a_corrected_column(tmp_path, write_mapping, schema) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [[1, 2, 3]])
    write_workbook(
        client_path,
        ["Other", "CFOP source", "COFI value", "Corrected NCM"],
        [["A", "5102", 12.5, "85167910"]],
    )
    mapping = mapping_rows()
    mapping.loc[2, "correct_client_column"] = ""
    write_mapping(mapping, mapping_path)

    with pytest.raises(
        ValueError, match=f"{label('correct_client_column', schema)} está em branco"
    ):
        build_mapped_workbook(
            client_path=client_path,
            mapping_path=mapping_path,
            output_path=output_path,
            model_path=model_path,
        )


def test_unreviewed_accepted_semantic_conflict_is_rejected(tmp_path, write_mapping) -> None:
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["Base ICMS_ST"], [[100.0]])
    write_workbook(client_path, ["ICMS Base"], [[100.0]])
    mapping = mapping_rows().iloc[[0]].copy()
    mapping["model_column"] = "Base ICMS_ST"
    mapping["client_column_number"] = 1
    mapping["client_column"] = "ICMS Base"
    write_mapping(mapping, mapping_path)

    with pytest.raises(ValueError, match="viola as regras atuais") as raised:
        build_mapped_workbook(
            client_path=client_path,
            mapping_path=mapping_path,
            output_path=output_path,
            model_path=model_path,
        )
    # A descrição do conflito sai em português, embora o texto que vai à IA
    # continue em inglês.
    assert "conflito de qualificador do tributo (ST vs normal)" in str(raised.value)


@pytest.mark.parametrize(
    ("left", "right", "expected"),
    [
        ("Base ICMS_ST", "ICMS Base", "conflito de qualificador do tributo (ST vs normal)"),
        ("Valor PIS", "Valor COFINS", "conflito de família de tributo (PIS vs COFINS)"),
        ("Valor IBS", "Valor CBS", "conflito de família de tributo (IBS vs CBS)"),
        ("Base ICMS", "Valor ICMS", "conflito de papel semântico (base de cálculo vs valor)"),
        ("Alíquota IPI", "Valor IPI", "conflito de papel semântico (alíquota vs valor)"),
        ("Vlr Isen ICM", "Vlr Outr ICM", "conflito de papel semântico (isentas vs outras)"),
        (
            "CST ICMS", "Valor ICMS",
            "conflito de papel semântico (situação tributária vs valor)",
        ),
        ("CFOP", "CFOP source", None),
        ("Base PIS/COFINS", "Base PIS", None),
    ],
)
def test_portuguese_header_conflict_agrees_with_the_model_input_text(left, right, expected) -> None:
    # A gêmea em português precisa achar o mesmo conflito, do mesmo tipo, que
    # semantic_header_conflict, cujo texto em inglês é entrada da IA.
    kinds = {
        "tax-family conflict": "conflito de família de tributo",
        "tax-qualifier conflict": "conflito de qualificador do tributo",
        "semantic-role conflict": "conflito de papel semântico",
    }
    original = semantic_header_conflict(left, right)
    portuguese = _describe_header_conflict_pt(left, right)
    assert portuguese == expected
    if original is None:
        assert portuguese is None
    else:
        kind = next(english for english in kinds if original.startswith(english))
        assert portuguese.startswith(kinds[kind])


def test_mapping_run_metadata_detects_changed_inputs(tmp_path) -> None:
    # Canário do esquema legado: as abas Mappings e Run Metadata, com key/value
    # e a chave client_sha256, como nos relatórios arquivados.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP"], [["5102"]])
    write_workbook(client_path, ["CFOP source"], [["5102"]])
    mapping = mapping_rows().iloc[[0]].copy()
    mapping["client_column_number"] = 1
    mapping["client_column"] = "CFOP source"
    with pd.ExcelWriter(mapping_path, engine="openpyxl") as writer:
        mapping.to_excel(writer, sheet_name="Mappings", index=False)
        pd.DataFrame(
            [["client_sha256", "not-the-current-hash"]],
            columns=["key", "value"],
        ).to_excel(writer, sheet_name="Run Metadata", index=False)

    with pytest.raises(ValueError, match="client.xlsx difere do arquivo usado"):
        build_mapped_workbook(
            client_path=client_path,
            mapping_path=mapping_path,
            output_path=output_path,
            model_path=model_path,
        )


def test_portuguese_run_metadata_detects_changed_inputs(tmp_path) -> None:
    # Gêmea em português do canário: abas Mapeamentos e Metadados da Execução,
    # com chave/valor e a chave sha256_cliente. Sem reconhecer os nomes, a
    # verificação de hash seria pulada em silêncio.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP"], [["5102"]])
    write_workbook(client_path, ["CFOP source"], [["5102"]])
    mapping = mapping_rows().iloc[[0]].copy()
    mapping["client_column_number"] = 1
    mapping["client_column"] = "CFOP source"
    with pd.ExcelWriter(mapping_path, engine="openpyxl") as writer:
        report_schema.to_external_frame(mapping, "mappings").to_excel(
            writer, sheet_name="Mapeamentos", index=False
        )
        pd.DataFrame(
            [["sha256_cliente", "nao-e-o-hash-atual"]],
            columns=["chave", "valor"],
        ).to_excel(writer, sheet_name="Metadados da Execução", index=False)

    with pytest.raises(ValueError, match="client.xlsx difere do arquivo usado"):
        build_mapped_workbook(
            client_path=client_path,
            mapping_path=mapping_path,
            output_path=output_path,
            model_path=model_path,
        )


@pytest.mark.parametrize(
    ("schema_of_mappings", "metadata_sheet", "metadata_columns"),
    [
        (report_schema.PT_BR, "Run Metadata", ["key", "value"]),
        (report_schema.LEGACY_EN, "Metadados da Execução", ["chave", "valor"]),
    ],
    ids=["mapeamentos_pt_metadados_en", "mapeamentos_en_metadados_pt"],
)
def test_metadata_sheet_from_the_other_schema_is_refused(
    tmp_path, schema_of_mappings, metadata_sheet, metadata_columns
) -> None:
    # Uma aba de metadados com o nome do outro esquema não pode ser ignorada:
    # a verificação de hash seria pulada sem aviso.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    write_workbook(model_path, ["CFOP"], [["5102"]])
    write_workbook(client_path, ["CFOP source"], [["5102"]])
    mapping = mapping_rows().iloc[[0]].copy()
    mapping["client_column_number"] = 1
    mapping["client_column"] = "CFOP source"
    with pd.ExcelWriter(mapping_path, engine="openpyxl") as writer:
        external_report(mapping, schema_of_mappings).to_excel(
            writer, sheet_name=report_schema.sheet_name("mappings", schema_of_mappings),
            index=False,
        )
        pd.DataFrame([["client_sha256", "x"]], columns=metadata_columns).to_excel(
            writer, sheet_name=metadata_sheet, index=False
        )

    with pytest.raises(ValueError, match="Esquema misto"):
        build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                              output_path=tmp_path / "out.xlsx", model_path=model_path)


def test_a_correction_outside_the_incorrect_branch_is_refused(tmp_path, write_mapping) -> None:
    # Só o ramo 'incorrect' lê correct_client_column. Em 'correct' a linha
    # continua sendo entregue, com a coluna do mapeador e não a do revisor;
    # um parecer escolhido errado punha dado ERRADO do cliente no modelo base,
    # em vez de deixá-lo em branco, e saía com código 0.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    for status in ("correct", "no_mapping_exists", ""):
        mapping = mapping_rows()
        mapping.loc[0, "reviewer_status"] = status
        mapping.loc[0, "correct_client_column"] = 4
        write_mapping(mapping, mapping_path)
        with pytest.raises(ValueError, match="usa uma correção"):
            build_mapped_workbook(
                client_path=client_path,
                mapping_path=mapping_path,
                output_path=tmp_path / "out.xlsx",
                model_path=model_path,
            )


def test_confirming_a_row_that_names_no_column_is_refused(tmp_path, write_mapping, schema) -> None:
    # 'correct' afirma que a proposta está certa, então precisa haver proposta.
    # Sem a trava a linha é confirmada sobre nada e entrega uma coluna em
    # branco assinada pelo revisor, com código de saída 0.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    mapping = mapping_rows()
    mapping.loc[2, "reviewer_status"] = "correct"
    mapping.loc[2, "correct_client_column"] = ""
    write_mapping(mapping, mapping_path)
    with pytest.raises(
        ValueError, match=f"{label('client_column_number', schema)} está em branco"
    ):
        build_mapped_workbook(
            client_path=client_path,
            mapping_path=mapping_path,
            output_path=tmp_path / "out.xlsx",
            model_path=model_path,
        )


def test_duplicated_header_is_selected_by_either_occurrence_suffix(
    tmp_path, write_mapping, schema
) -> None:
    # Cabeçalho repetido na exportação: em memória a segunda ocorrência é
    # 'Valor ICMS (occurrence 2)'; para pessoas é 'Valor ICMS (ocorrência 2)'.
    # O relatório em português grava o sufixo em português, e o revisor pode
    # digitar qualquer uma das formas, com ou sem acento, em qualquer esquema.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["Valor ICMS", "CFOP"], [[1.0, "5102"]])
    write_workbook(client_path, ["Valor ICMS", "Valor ICMS"], [[10.0, 20.0]])

    proposal = mapping_rows().iloc[:2].copy()
    proposal["model_column_number"] = [1, 2]
    proposal["model_column"] = ["Valor ICMS", "CFOP"]
    proposal["client_column_number"] = [2, None]
    proposal["client_column"] = ["Valor ICMS (occurrence 2)", None]
    proposal["status"] = ["accepted", "no_match"]

    # A proposta aceita aponta a segunda ocorrência pelo nome de exibição.
    write_mapping(proposal, mapping_path)
    if schema == report_schema.PT_BR:
        stored = pd.read_excel(mapping_path, dtype=object)
        assert stored.loc[0, "coluna_cliente"] == "Valor ICMS (ocorrência 2)"
    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=output_path, model_path=model_path)
    sheet = load_workbook(output_path)[MAPPED_SHEET]
    assert sheet.cell(FIRST_DATA_EXCEL_ROW, 2).value == 20.0
    assert sheet.cell(banner_excel_row("client_column"), 2).value == "Valor ICMS (ocorrência 2)"
    coverage = list(load_workbook(output_path)[COVERAGE_SHEET].values)
    delivered = dict(zip(coverage[0], coverage[1]))
    assert delivered["coluna_cliente"] == "Valor ICMS (ocorrência 2)"

    # A correção do revisor aceita os dois sufixos e o número da coluna.
    for typed in (
        "Valor ICMS (ocorrência 2)",
        "valor icms (OCORRENCIA 2)",
        "Valor ICMS (occurrence 2)",
        "2",
    ):
        corrected = proposal.copy()
        corrected["client_column_number"] = [1, None]
        corrected["client_column"] = ["Valor ICMS", None]
        corrected["status"] = ["review", "no_match"]
        corrected["reviewer_status"] = ["incorrect", ""]
        corrected["correct_client_column"] = [typed, ""]
        write_mapping(corrected, mapping_path)
        build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                              output_path=output_path, model_path=model_path)
        sheet = load_workbook(output_path)[MAPPED_SHEET]
        assert sheet.cell(FIRST_DATA_EXCEL_ROW, 2).value == 20.0, typed

    # O nome sem sufixo é o nome de exibição da primeira ocorrência e a
    # seleciona (comportamento documentado em docs/DATA.md).
    corrected["correct_client_column"] = ["Valor ICMS", ""]
    write_mapping(corrected, mapping_path)
    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=output_path, model_path=model_path)
    assert load_workbook(output_path)[MAPPED_SHEET].cell(FIRST_DATA_EXCEL_ROW, 2).value == 10.0

    # Uma ocorrência que não existe não é trocada por outra: é recusada.
    corrected["correct_client_column"] = ["Valor ICMS (ocorrência 3)", ""]
    write_mapping(corrected, mapping_path)
    with pytest.raises(
        ValueError, match=re.escape("'Valor ICMS (ocorrência 3)' não foi encontrada")
    ):
        build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                              output_path=output_path, model_path=model_path)


def test_columns_with_no_data_are_hidden_not_removed(tmp_path, write_mapping) -> None:
    # Um cliente preenche uma minoria do modelo, e isso é o desenho. Mas as
    # colunas entregues ficam espalhadas: as 19 de datasul vão de AI a GW, com
    # 101 colunas vazias entre duas delas. Ocultar as põe lado a lado sem
    # remover nada nem mudar valor algum.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM"], [["5102", 10.0, "85167910"]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    write_mapping(mapping_rows(), mapping_path)
    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=output_path, model_path=model_path)
    book = openpyxl.load_workbook(output_path)
    sheet = book[MAPPED_SHEET]
    hidden = {
        get_column_letter(c): sheet.column_dimensions[get_column_letter(c)].hidden
        for c in range(2, 5)
    }
    # O fixture entrega duas: 'CFOP', aceita, e 'Nº do NCM', por correção do
    # revisor, que é o propósito do ciclo de revisão e é como sap_procosa
    # chega a 62 entregues a partir de 45 aceites. Só a linha em revisão não
    # confirmada, 'Valor COFINS', fica sem dado.
    assert hidden["B"] is False
    assert hidden["C"] is True
    assert hidden["D"] is False
    # ocultar não pode remover nada: as três colunas do modelo continuam lá
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    assert list(header[1:4]) == ["CFOP", "Valor COFINS", "Nº do NCM"]


def test_coverage_sheet_separates_absence_from_a_pending_decision(tmp_path, write_mapping) -> None:
    # A aba mapeada não responde "o que não está aqui e deveria estar", porque
    # toda coluna que ela preenche é aceita por construção. Uma coluna vazia
    # quer dizer que o cliente não tem esse conceito (uma resposta completa)
    # ou que há um candidato esperando o revisor. Só o segundo caso é trabalho.
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    mapping_path = tmp_path / "mapping.xlsx"
    output_path = tmp_path / "mapped.xlsx"
    write_workbook(model_path, ["CFOP", "Valor COFINS", "Nº do NCM", "Valor IPI"],
                   [["5102", 10.0, "85167910", 1.0]])
    write_workbook(client_path, ["Other", "CFOP source", "COFI value", "Corrected NCM"],
                   [["A", "5102", 12.5, "85167910"]])
    mapping = pd.concat([mapping_rows(), mapping_rows().iloc[[2]]], ignore_index=True)
    mapping.loc[3, ["model_column_number", "model_column"]] = [4, "Valor IPI"]
    mapping.loc[3, ["reviewer_status", "correct_client_column", "reviewer_note"]] = ["", "", ""]
    write_mapping(mapping, mapping_path)
    build_mapped_workbook(client_path=client_path, mapping_path=mapping_path,
                          output_path=output_path, model_path=model_path)
    book = openpyxl.load_workbook(output_path)
    assert COVERAGE_SHEET in book.sheetnames
    rows = list(book[COVERAGE_SHEET].values)
    header = list(rows[0])
    assert header == list(report_schema.COLUMNS["coverage"].values())
    data = [dict(zip(header, r)) for r in rows[1:]]
    states = [r["situacao_cobertura"] for r in data]
    # ordenado para o trabalho vir primeiro: entregue, depois pendente, depois ausência
    assert states == [
        "entregue", "entregue", "aguardando decisão", "não consta na exportação do cliente",
    ]
    assert data[0]["coluna_modelo"] == "CFOP" and data[0]["coluna_excel"] == "B"
    # uma correção do revisor conta como entregue e cita a coluna onde caiu
    assert data[1]["coluna_modelo"] == "Nº do NCM"
    assert data[1]["coluna_cliente"] == "Corrected NCM"
    # a linha pendente ainda precisa citar o candidato que o revisor vai julgar
    assert data[2]["coluna_cliente"] == "COFI value"
    # a justificativa é copiada do relatório como está
    assert data[2]["justificativa"] == "Names and numeric values are compatible."
    assert data[3]["coluna_modelo"] == "Valor IPI"


def test_coverage_states_are_the_portuguese_of_the_legacy_states() -> None:
    states = report_schema.ENUMS["coverage_state"]
    assert states["delivered"] == STATE_DELIVERED
    assert states["awaiting decision"] == STATE_PENDING
    assert states["not in this client's export"] == STATE_ABSENT


def test_command_line_is_in_portuguese(capsys) -> None:
    help_text = build_parser().format_help()
    assert help_text.startswith("uso: ")
    assert "opções:" in help_text
    assert "Exportação .xlsx do cliente" in help_text
    # As flags continuam em inglês.
    assert "--include-review" in help_text and "--output" in help_text

    with pytest.raises(SystemExit) as raised:
        main(["c.xlsx", "r.xlsx", "--output", "o.xlsx", "--accepted-only", "--include-review"])
    assert str(raised.value) == "--accepted-only e --include-review não podem ser usados juntos"

    with pytest.raises(SystemExit) as raised:
        main([])
    assert raised.value.code == 2
    assert "os seguintes argumentos são obrigatórios" in capsys.readouterr().err
