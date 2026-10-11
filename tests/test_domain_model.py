"""Verifica os contratos compartilhados de campos, revisão e seleção de origens."""
from dataclasses import asdict
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import pandas as pd
import pytest

import invoice_mapper as mapper
import report_schema
from apply_invoice_mapping import build_mapped_workbook
from mapeval.score import score_report
from mapeval.truth import AnswerKey, load_answer_key
from reference_catalog import expected_catalogs
from spreadsheet import SheetTable
from target_schema import target_field


def table(headers, rows):
    result = SheetTable.from_headers(headers)
    result.data = pd.DataFrame(rows)
    return result


class Scores:
    def __init__(self, rows):
        self.rows = rows

    def scores(self, model_headers, client_headers):
        return np.array([self.rows[name] for name in model_headers])


class FailingResolver:
    def resolve(self, **kwargs):
        raise AssertionError("Este cenário deve ser resolvido sem consultar a API")


def report_files(tmp_path, rows):
    client = tmp_path / "client.xlsx"
    report = tmp_path / "report.xlsx"
    pd.DataFrame([["Source A", "Source B", "Source C"], [11, 22, 33]]).to_excel(
        client, header=False, index=False,
    )
    defaults = dict(client_column_number=1, client_column="Source A", status="review",
                    final_confidence=.85, name_confidence=.85, name_runner_up_confidence=.4,
                    name_margin=.45, data_confidence=.8, method="data", reason="Proposta",
                    reviewer_status="", correct_client_column="", reviewer_note="")
    frame = pd.DataFrame([{**defaults, **row} for row in rows])
    report_schema.to_external_frame(frame, "mappings").to_excel(
        report, sheet_name="Mapeamentos", index=False,
    )
    return client, report


def test_evaluation_is_invariant_to_report_row_order(tmp_path):
    rows = [dict(model_column_number=1, model_column="Target A", status="accepted"),
            dict(model_column_number=2, model_column="Target B", client_column_number=2,
                 client_column="Source B", status="accepted")]
    client, report = report_files(tmp_path, rows)
    key = AnswerKey("synthetic", {1: 1, 2: 2})
    before = score_report(report, client, key)
    report_files(tmp_path, rows[::-1])
    after = score_report(report, client, key)
    assert asdict(before) == asdict(after)
    assert after.delivered_correct == 2


@pytest.mark.parametrize("numbers", [[1, 1], [1, 0], [1, 1.5], [1, None], [1, "invalid"], [1, 3]])
def test_evaluation_rejects_invalid_field_identity(tmp_path, numbers):
    client, report = report_files(tmp_path, [
        dict(model_column_number=n, model_column=f"Target {i}") for i, n in enumerate(numbers)
    ])
    with pytest.raises(ValueError, match="numero_coluna_modelo"):
        score_report(report, client, AnswerKey("synthetic", {}))


def test_apply_coverage_and_evaluation_use_the_same_review_decision(tmp_path):
    rows = [
        dict(model_column_number=1, model_column="Target A", status="accepted", reviewer_status="no_mapping_exists"),
        dict(model_column_number=2, model_column="Target B", client_column_number=2,
             client_column="Source B", reviewer_status="correct"),
        dict(model_column_number=3, model_column="Target C", reviewer_status="incorrect",
             correct_client_column="Source C"),
    ]
    client, report = report_files(tmp_path, rows)
    model, output = tmp_path / "model.xlsx", tmp_path / "output.xlsx"
    pd.DataFrame([["Target A", "Target B", "Target C"]]).to_excel(model, index=False, header=False)
    build_mapped_workbook(client, report, output, model)
    coverage = pd.read_excel(output, sheet_name="Cobertura").set_index("numero_coluna_modelo")
    assert coverage.loc[1, "situacao_cobertura"] == "não consta na exportação do cliente"
    assert coverage.loc[2, "situacao_cobertura"] == "entregue"
    assert coverage.loc[3, "numero_coluna_cliente"] == 3
    score = score_report(report, client, AnswerKey("synthetic", {1: 1, 2: 2, 3: 3}))
    assert [r.status for r in score.rows] == ["no_match", "accepted", "accepted"]
    assert [r.client_column for r in score.rows] == [None, 2, 3]
    assert score.delivered_correct == 2
    assert not score.withheld_correct


def test_reviewer_cannot_confirm_the_same_source_for_two_fields(tmp_path):
    client, report = report_files(tmp_path, [
        dict(model_column_number=1, model_column="Target A", reviewer_status="correct"),
        dict(model_column_number=2, model_column="Target B", reviewer_status="correct"),
    ])
    with pytest.raises(ValueError, match="mais de um destino"):
        score_report(report, client, AnswerKey("synthetic", {}))


def test_answer_key_reuse_is_flagged_without_changing_the_key():
    for name, source, targets in [("datasul", 8, [39, 147]), ("oracle", 23, [100, 190])]:
        key = load_answer_key(f"data/keys/de_para_{name}.xlsx", name)
        assert key.source_conflicts == {source: targets}
        assert all(key.present[target] == source for target in targets)
    assert AnswerKey("synthetic", {1: 1, 2: 2}).source_conflicts == {}


def test_explicit_registry_covers_every_canonical_column():
    raw = json.loads(Path("target_fields.json").read_text())
    model = SheetTable.from_excel("data/model.xlsx")
    assert [(row[0], row[1]) for row in raw["fields"]] == [
        (c.index + 1, c.name) for c in model.columns
    ]
    assert all(target_field(c.name) is not None for c in model.columns)


@pytest.mark.parametrize("header,catalog", [
    ("Tipo de Imposto CBS", "ibs_cbs_cst"), ("Tipo de Imposto IBSM", "ibs_cbs_cst"),
    ("Tipo de Imposto PIS", "cst_pis_cofins"), ("Tipo de Imposto IPI", "cst_ipi"),
])
def test_canonical_tax_status_fields_have_explicit_catalogs(header, catalog):
    assert catalog in expected_catalogs(header)


@pytest.mark.parametrize("cpf,cnpj", [(0, 1), (1, 0), (.5, .5)])
def test_allowed_identifiers_do_not_depend_on_target_sample_mix(cpf, cnpj):
    kinds = target_field("Nº ID fiscal 1").identifier_kinds
    source = dict(valid_cpf_fraction=cpf, valid_cnpj_fraction=cnpj)
    for sample in ({}, dict(valid_cpf_fraction=.914, valid_cnpj_fraction=.086)):
        assert mapper.identifier_match_score(sample, source, kinds=kinds) == 1
    assert mapper.identifier_match_score({}, source, kinds=()) == 0


@pytest.mark.parametrize("reverse", [False, True])
def test_stronger_header_claim_wins_independently_of_target_order(reverse):
    headers = ["Target A", "Target B"]
    if reverse:
        headers.reverse()
    results = mapper.map_columns(
        table(headers, [[1, 2]]), table(["Source X", "Source Y"], [[11, 22]]),
        Scores({"Target A": [.91, .70], "Target B": [.99, .70]}), None,
    )
    accepted = {r.model_column: r.client_column for r in results if r.status == "accepted"}
    assert accepted == {"Target B": "Source X"}


@pytest.mark.parametrize("reverse", [False, True])
def test_ambiguous_header_claims_stay_for_review_online(reverse):
    headers = ["Target A", "Target B"][::(-1 if reverse else 1)]
    class Resolver:
        def resolve(self, **kwargs):
            return mapper.DataMappingDecision(client_column_id="column_001", data_confidence=.99,
                                              final_confidence=.99, reason="Shared proposal.")
    results = mapper.map_columns(
        table(headers, [[1, 2]]), table(["Source X", "Source Y"], [[11, 22]]),
        Scores({"Target A": [.95, .70], "Target B": [.94, .70]}), Resolver(),
    )
    assert all(r.status == "review" for r in results)
    assert all(r.reason == mapper.REASON_SOURCE_COMPETITION for r in results)


def test_unique_arithmetic_succeeds_without_a_resolver_call():
    base = [100., 200., 300., 400., 500.] * 4
    weaker = [100., 200., 300., 400., 625.] * 4
    rate = [10.] * 20
    value = [10., 20., 30., 40., 50.] * 4
    result = mapper.map_columns(
        table(["Base ICMS", "Aliquota ICMS", "Valor ICMS"], list(zip(base, rate, value))),
        table(["X1", "X2", "ICMS Taxa", "ICMS Valor"], list(zip(base, weaker, rate, value))),
        Scores({"Base ICMS": [.84, .86, .20, .20], "Aliquota ICMS": [.20, .20, .96, .20],
                "Valor ICMS": [.20, .20, .20, .96]}), FailingResolver(),
    )[0]
    assert (result.client_column, result.status) == ("X1", "accepted")
    assert result.llm_final_confidence is None
    assert "resolver" not in result.supported_by


def test_resolver_reuses_the_six_prepared_profiles():
    class Responses:
        def parse(self, **kwargs):
            return SimpleNamespace(output_parsed=mapper.DataMappingDecision(
                client_column_id=None, data_confidence=0, final_confidence=0, reason="No choice."))
    with patch.object(mapper, "summarize_values", wraps=mapper.summarize_values) as summarize:
        mapper.map_columns(
            table(["Target A", "Target B", "Target C"], [["one", "two", "three"]]),
            table(["Source X", "Source Y", "Source Z"], [["red", "green", "blue"]]),
            Scores({"Target A": [.85, .84, .83], "Target B": [.84, .85, .83],
                    "Target C": [.83, .84, .85]}),
            mapper.OpenAIDataResolver(SimpleNamespace(responses=Responses())),
        )
    assert summarize.call_count == 6


@pytest.mark.parametrize("reverse", [False, True])
@pytest.mark.parametrize("weaker,expected", [(.75, {"Target B": "Source X"}), (.87, {})])
def test_second_pass_compares_all_source_claims_before_accepting(reverse, weaker, expected):
    headers = ["Target A", "Target B"][::(-1 if reverse else 1)]
    class Resolver:
        def resolve(self, **kwargs):
            return mapper.DataMappingDecision(client_column_id="column_001", data_confidence=.99,
                                              final_confidence=.99, reason="Shared proposal.")
    values = ["distinctive first example", "distinctive second example", "distinctive third example"]
    results = mapper.map_columns(
        table(headers, [[value, value] for value in values]),
        table(["Source X", "Source Y"], [[value, "other data"] for value in values]),
        Scores({"Target A": [weaker, .4], "Target B": [.88, .4]}), Resolver(),
    )
    accepted = {r.model_column: r.client_column for r in results if r.status == "accepted"}
    assert accepted == expected
    assert all(r.client_column == "Source X" for r in results)


def test_report_accounts_for_every_client_column_and_formats_confidence(tmp_path):
    import openpyxl
    model = table(["Target A", "Target B", "Target C", "Target D"], [[1, 2, 3, 4]])
    client = table(["Source A", "Source B", "Unused empty source"], [[11, 22, None]])
    scorer = Scores({"Target A": [.99, .4, 0], "Target B": [.4, .85, 0],
                     "Target C": [.4, .75, 0], "Target D": [.4, .7, 0]})
    results = mapper.map_columns(model, client, scorer, None)
    unused = mapper.summarize_unmatched_columns(
        model_table=model, client_table=client, results=dict(enumerate(results)),
        similarities=scorer.scores([c.name for c in model.columns], []),
        reference_catalogs=None, client_tax_trios={}, sample_size=8,
    )
    accepted = {r.client_column_number for r in results if r.status == "accepted"}
    unassigned = {r.client_column_number for r in unused}
    assert accepted.isdisjoint(unassigned)
    assert accepted | unassigned == {1, 2, 3}
    assert {r.client_column_number: r.disposition for r in unused} == {
        2: "proposed_not_accepted", 3: "never_referenced",
    }
    output = tmp_path / "report.xlsx"
    mapper.write_results(results, output, unmatched=unused, queue=[])
    workbook = openpyxl.load_workbook(output)
    try:
        sheet = workbook["Mapeamentos"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        confidence = sheet.cell(2, headers["confianca_final"])
        assert confidence.value == .99
        assert confidence.number_format == "0.00%"
        assert sheet.cell(2, headers["decidido_por"]).value == "semantico"
        assert workbook["Colunas do Cliente sem Uso"].max_row == 3
    finally:
        workbook.close()


@pytest.mark.parametrize("reverse", [False, True])
def test_withdrawn_competing_proposal_does_not_veto_the_remaining_claim(reverse):
    headers = ["Target A", "Target B"][::(-1 if reverse else 1)]
    class Resolver:
        def resolve(self, **kwargs):
            chosen = "column_001" if kwargs["model_column"].name == "Target A" else None
            return mapper.DataMappingDecision(client_column_id=chosen, data_confidence=.99 if chosen else 0,
                                              final_confidence=.99 if chosen else 0, reason="Fixed proposal.")
    results = mapper.map_columns(
        table(headers, [[f"first {n}", f"second {n}"] for n in range(30)]),
        table(["Source X", "Source Y"], [[f"third {n}", f"fourth {n}"] for n in range(30)]),
        Scores({"Target A": [.9414, .7], "Target B": [.9116, .7]}), Resolver(),
    )
    assert {r.model_column: r.status for r in results} == {"Target A": "accepted", "Target B": "no_match"}


@pytest.mark.parametrize("reverse", [False, True])
def test_newly_confirmed_values_enable_local_arithmetic_without_more_api_calls(reverse):
    base = [1e9 * n for n in range(1, 31)]
    values = {"Valor ICMS": [b * .1 for b in base], "Aliquota ICMS": [10.] * 30,
              "Base ICMS": [None] * 30}
    headers = list(values)[::(-1 if reverse else 1)]
    calls = []
    class Resolver:
        def resolve(self, **kwargs):
            name = kwargs["model_column"].name
            calls.append(name)
            return mapper.DataMappingDecision(
                client_column_id={"Valor ICMS": "column_001", "Base ICMS": "column_003"}[name],
                data_confidence=.99, final_confidence=.99, reason="Fixed proposal.")
    results = mapper.map_columns(
        table(headers, list(zip(*(values[h] for h in headers)))),
        table(["ICMS Valor", "ICMS Taxa", "ICMS Base"],
              list(zip(values["Valor ICMS"], values["Aliquota ICMS"], base))),
        Scores({"Valor ICMS": [.85, .1, .1], "Aliquota ICMS": [.1, .98, .1],
                "Base ICMS": [.1, .1, .85]}), Resolver(),
    )
    by_name = {r.model_column: r for r in results}
    assert all(r.status == "accepted" for r in results)
    assert len({r.client_column_number for r in results}) == 3
    assert sorted(calls) == ["Base ICMS", "Valor ICMS"]
    assert by_name["Base ICMS"].final_confidence == .94
    assert by_name["Base ICMS"].decided_by == "arithmetic"
    assert by_name["Valor ICMS"].decided_by == "profile"
    assert by_name["Aliquota ICMS"].decided_by == "semantic"
    assert "arithmetic" in by_name["Valor ICMS"].supported_by


def test_target_descriptions_have_distinct_explicit_identities():
    assert target_field("CFOP_descrição").number == 35
    assert target_field("item/material_descrição").number == 195
    assert target_field("Descrição") is None
    assert target_field("CFOP_descrição", number=195) is None
    assert target_field("CFOP_descrição", number=35).role == "text"


@pytest.mark.parametrize("include_review", [False, True])
def test_unused_sources_follow_review_and_keep_pending_and_empty_columns(tmp_path, include_review):
    client, report = report_files(tmp_path, [
        dict(model_column_number=1, model_column="Target A", status="accepted",
             reviewer_status="incorrect", correct_client_column="Source B"),
        dict(model_column_number=2, model_column="Target B", client_column_number=3,
             client_column="Source C", status="review"),
    ])
    pd.DataFrame([["Source A", "Source B", "Source C", "Empty"], [11, 22, 33, None]]).to_excel(
        client, header=False, index=False)
    model, output = tmp_path / "model.xlsx", tmp_path / "output.xlsx"
    pd.DataFrame([["Target A", "Target B"]]).to_excel(model, header=False, index=False)
    original_report = report.read_bytes()
    build_mapped_workbook(client, report, output, model, accepted_only=not include_review)
    unused = report_schema.to_internal_frame(
        pd.read_excel(output, sheet_name="Colunas do Cliente sem Uso"), "unmatched")
    assert dict(zip(unused.client_column_number, unused.disposition)) == {
        1: "never_referenced", 3: "proposed_not_accepted", 4: "never_referenced"}
    assert list(unused.rows_with_value) == [1, 1, 0]
    assert report.read_bytes() == original_report


@pytest.mark.parametrize("confidence", [0.0, 0.01, 0.99])
def test_local_identifier_evidence_does_not_depend_on_resolver_confidence(confidence):
    model = table(["Nº ID fiscal 1"], [["model A"], ["model B"], ["model C"]])
    client = table(["Source X"], [["52998224725"], ["11144477735"], ["52998224725"]])
    raw = mapper.DataMappingDecision(client_column_id="column_001", data_confidence=confidence,
                                     final_confidence=confidence, reason="Controlled API response")
    api = SimpleNamespace(responses=SimpleNamespace(parse=lambda **kwargs: SimpleNamespace(output_parsed=raw)))
    result = mapper.map_columns(model, client, Scores({"Nº ID fiscal 1": [.8]}),
                                mapper.OpenAIDataResolver(api))[0]
    assert (result.status, result.client_column_number, result.decided_by) == ("accepted", 1, "identifier")
    assert result.final_confidence == .92
    assert result.llm_final_confidence == confidence
    assert result.llm_data_confidence == confidence


def test_weak_rows_preserve_raw_resolver_confidence_in_the_audit():
    raw = mapper.DataMappingDecision(client_column_id="column_001", data_confidence=.1,
                                     final_confidence=.1, reason="Weak data")
    api = SimpleNamespace(responses=SimpleNamespace(parse=lambda **kwargs: SimpleNamespace(output_parsed=raw)))
    result = mapper.map_columns(table(["Target A"], [[0], [0], [0]]),
                                table(["Source X"], [[0], [0], [0]]),
                                Scores({"Target A": [.88]}), mapper.OpenAIDataResolver(api))[0]
    assert result.llm_final_confidence == .1
    assert result.llm_data_confidence == .1
    assert result.final_confidence == .88
    assert result.status == "review"


def test_report_generation_reuses_prepared_client_profiles():
    from dataclasses import asdict
    model = table(["Target A", "Target B"], [["model A", "model B"]])
    client = table(["Source X", "Source Y"], [["source X", "source Y"]])
    class Resolver:
        def resolve(self, **kwargs):
            return mapper.DataMappingDecision(client_column_id="column_001", data_confidence=.5,
                                               final_confidence=.5, reason="Proposal")
    run = mapper._start_run(model, client, Scores({"Target A": [.86, .7], "Target B": [.86, .7]}),
                            mapper.MapperConfig(), None, online=True)
    results = mapper._map_run(run, Resolver())
    def reports(profiles=None):
        unused = mapper.summarize_unmatched_columns(
            model_table=model, client_table=client, results=run.results, similarities=run.similarities,
            reference_catalogs=None, client_tax_trios=run.client_tax_trios, sample_size=8,
            client_profiles=profiles,
        )
        queue = mapper.summarize_adjudication_queue(client_table=client, results=results,
                    final_threshold=.9, sample_size=8, client_profiles=profiles)
        return [asdict(row) for row in unused], [asdict(row) for row in queue]
    expected = reports()
    with patch.object(mapper, "summarize_values", side_effect=AssertionError("Profile recomputed")):
        assert reports(run.client_profiles) == expected


@pytest.mark.parametrize("alternatives, expected", [(1, "lexical + arithmetic"), (2, "lexical")])
def test_arithmetic_support_is_independent_of_explanation_wording(alternatives, expected):
    run = SimpleNamespace(client_tax_trios={}, header_step="lexical")
    column = mapper.ColumnRef("target_001", 0, "Valor PIS", 1)
    with patch.object(mapper, "trio_corroboration", return_value={"alternatives": alternatives}), \
         patch.object(mapper, "_trio_proof_text", return_value="Uma explicação reescrita livremente"):
        evidence = mapper._evidence_attribution(run, column, 0, decided_by="lexical", name_confidence=.95)
    assert evidence["supported_by"] == expected
    assert evidence["arithmetic_proof"] == "Uma explicação reescrita livremente"


@pytest.mark.parametrize("confidence", [0.0, .99])
def test_joint_arithmetic_acceptance_uses_local_evidence(confidence):
    rows = [[100+10*i, 1.65, (100+10*i)*.0165, (100+10*i)*.0165] for i in range(40)]
    model = table(["Base PIS", "Aliquota PIS", "Valor PIS"], [row[:3] for row in rows])
    client = table(["PIS base", "PIS rate", "Candidate A", "Candidate B"], rows)
    class Resolver:
        def resolve(self, **kwargs):
            return mapper.DataMappingDecision(client_column_id="column_003", data_confidence=confidence,
                                               final_confidence=confidence, reason="Selected candidate")
    result = mapper.map_columns(model, client, Scores({"Base PIS": [.99,.1,.1,.1],
        "Aliquota PIS": [.1,.99,.1,.1], "Valor PIS": [.1,.1,.89,.7]}), Resolver())[2]
    assert (result.status, result.client_column_number, result.final_confidence) == ("accepted", 3, .956)
    assert result.decided_by == "arithmetic"


@pytest.mark.parametrize("column_id", [None, "unknown_column"])
def test_zero_confidence_fix_does_not_invent_a_selected_column(column_id):
    class Resolver:
        def resolve(self, **kwargs):
            return mapper.DataMappingDecision(client_column_id=column_id, data_confidence=0,
                                               final_confidence=0, reason="No selection")
    result = mapper.map_columns(table(["Nº ID fiscal 1"], [["model"]]),
        table(["Source X"], [["52998224725"]]), Scores({"Nº ID fiscal 1": [.8]}), Resolver())[0]
    assert result.status == "no_match"
    assert result.client_column_number is None


@pytest.mark.parametrize("reverse", [False, True])
def test_zero_confidence_candidates_still_participate_in_source_competition(reverse):
    headers = ["Target A", "Target B"][::(-1 if reverse else 1)] + ["Anchor"]
    model = table(headers, [["model A", "model B", "anchor"]])
    client = table(["Source X", "Source Y"], [["source X", "source Y"]])
    class Resolver:
        def resolve(self, **kwargs):
            confidence = 0 if kwargs["model_column"].name == "Target A" else .5
            return mapper.DataMappingDecision(client_column_id="column_001", data_confidence=confidence,
                                               final_confidence=confidence, reason="Proposal")
    results = mapper.map_columns(model, client, Scores({
        "Target A": [.9009, .89], "Target B": [.92, .89], "Anchor": [.2, .99],
    }), Resolver())
    by_name = {r.model_column: r for r in results}
    assert by_name["Anchor"].status == "accepted"
    for name in ["Target A", "Target B"]:
        assert by_name[name].status == "review"
        assert by_name[name].vetoed_by == "exclusivity"
        assert by_name[name].decided_by == ""
