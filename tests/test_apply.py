"""Testes de ``mapper/apply.py``: parecer do revisor, planilha mapeada e Cobertura."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import time
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd
import pytest

from mapper import schema
from mapper.apply import (
    CONFIDENCE_ROWS,
    FIRST_DATA_EXCEL_ROW,
    MAPPING_METADATA,
    STATE_ABSENT,
    STATE_DELIVERED,
    STATE_PENDING,
    build_mapped_workbook,
    build_parser,
    main,
)
from mapper.evaluate import AnswerKey, score_report

from support import report_files

MAPPED_SHEET = "Dados Mapeados"
COVERAGE_SHEET = "Cobertura"
#: Índice da primeira linha de dados no DataFrame lido sem cabeçalho (base 0).
FIRST_DATA_ROW = len(MAPPING_METADATA) + 1

#: Modelo e exportação que ``mapping_rows`` descreve.
MODEL_HEADERS = ["CFOP", "Valor COFINS", "Nº do NCM"]
CLIENT_HEADERS = ["Other", "CFOP source", "COFI value", "Corrected NCM"]
CLIENT_ROW = ["A", "5102", 12.5, "85167910"]


@dataclass
class Files:
    model: Path
    client: Path
    mapping: Path
    output: Path


@pytest.fixture
def files(tmp_path) -> Files:
    return Files(
        model=tmp_path / "model.xlsx",
        client=tmp_path / "client.xlsx",
        mapping=tmp_path / "mapping.xlsx",
        output=tmp_path / "mapped.xlsx",
    )


def write_workbook(path, headers: list[str], rows: list[list[object]]) -> None:
    pd.DataFrame([headers, *rows]).to_excel(path, header=False, index=False)


def write_inputs(files: Files, client_rows: list[list[object]] | None = None) -> None:
    """Grava o modelo de três colunas e a exportação que ``mapping_rows`` descreve."""
    write_workbook(files.model, MODEL_HEADERS, [["5102", 10.0, "85167910"]])
    write_workbook(files.client, CLIENT_HEADERS, client_rows or [CLIENT_ROW])


def mapping_rows() -> pd.DataFrame:
    """Relatório de três linhas, com nomes e valores internos."""
    return pd.DataFrame(
        [
            {
                "model_column_number": 1,
                "model_column": "CFOP",
                "client_column_number": 2,
                "client_column": "CFOP source",
                "name_confidence": 0.98,
                "name_runner_up_confidence": 0.70,
                "name_margin": 0.28,
                "data_confidence": None,
                "llm_data_confidence": None,
                "llm_final_confidence": None,
                "final_confidence": 0.98,
                "method": "name",
                "status": "accepted",
                "reason": "Clear header match.",
                "reviewer_status": "",
                "correct_client_column": "",
                "reviewer_note": "",
            },
            {
                "model_column_number": 2,
                "model_column": "Valor COFINS",
                "client_column_number": 3,
                "client_column": "COFI value",
                "name_confidence": 0.82,
                "name_runner_up_confidence": 0.80,
                "name_margin": 0.02,
                "data_confidence": 0.91,
                "llm_data_confidence": 0.93,
                "llm_final_confidence": 0.89,
                "final_confidence": 0.88,
                "method": "data",
                "status": "review",
                "reason": "Names and numeric values are compatible.",
                "reviewer_status": "",
                "correct_client_column": "",
                "reviewer_note": "Check this value field.",
            },
            {
                "model_column_number": 3,
                "model_column": "Nº do NCM",
                "client_column_number": None,
                "client_column": None,
                "name_confidence": 0.40,
                "name_runner_up_confidence": 0.38,
                "name_margin": 0.02,
                "data_confidence": 0.10,
                "llm_data_confidence": 0.12,
                "llm_final_confidence": 0.08,
                "final_confidence": 0.0,
                "method": "data",
                "status": "no_match",
                "reason": "No automatic match.",
                "reviewer_status": "incorrect",
                "correct_client_column": "Corrected NCM",
                "reviewer_note": "Reviewer supplied the source column.",
            },
        ]
    )


def one_row_mapping(model_column: str, client_column: str) -> pd.DataFrame:
    """A linha aceita de ``mapping_rows``, ligando ``model_column`` à coluna 1 do cliente."""
    mapping = mapping_rows().iloc[[0]].copy()
    mapping["model_column"] = model_column
    mapping["client_column_number"] = 1
    mapping["client_column"] = client_column
    return mapping


def external_report(frame: pd.DataFrame) -> pd.DataFrame:
    """Converte os valores internos de um relatório de teste para o schema de saída."""
    return schema.to_external_frame(frame, "mappings")


def save_report(path, report: pd.DataFrame) -> None:
    """Grava um relatório já em português."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        report.to_excel(writer, sheet_name=schema.sheet_name("mappings"), index=False)


def write_mapping(frame: pd.DataFrame, path) -> None:
    """Grava em português um relatório montado com nomes internos."""
    save_report(path, external_report(frame))


def apply(files: Files, **options) -> None:
    build_mapped_workbook(
        client_path=files.client,
        mapping_path=files.mapping,
        output_path=files.output,
        model_path=files.model,
        **options,
    )


def read_mapped(path) -> pd.DataFrame:
    """Lê Dados Mapeados sem cabeçalho; o índice 0 do DataFrame corresponde à linha 1 do Excel."""
    return pd.read_excel(path, sheet_name=MAPPED_SHEET, header=None, dtype=object)


def label(field: str) -> str:
    """Retorna o cabeçalho usado nas mensagens de erro da aba Mapeamentos."""
    return schema.header_label("mappings", field)


def banner_row(field: str) -> int:
    """Retorna o índice do campo no banner, a partir de zero, para leitura pelo DataFrame."""
    return 1 + [name for _, name in MAPPING_METADATA].index(field)


def banner_excel_row(field: str) -> int:
    """Retorna a linha do campo no banner, a partir de 1, para acesso pelo Excel."""
    return banner_row(field) + 1


def test_build_mapped_workbook_defaults_to_confirmed_mappings(files) -> None:
    write_inputs(files, [CLIENT_ROW, ["B", "6102", 7.5, "73211100"]])
    write_mapping(mapping_rows(), files.mapping)

    apply(files)

    raw = read_mapped(files.output)
    assert raw.iloc[0].tolist() == [
        "__info_mapeamento__",
        "CFOP",
        "Valor COFINS",
        "Nº do NCM",
    ]
    assert raw.iloc[1 : len(MAPPING_METADATA) + 1, 0].tolist() == [
        label for label, _ in MAPPING_METADATA
    ]

    client_column_row = banner_row("client_column")
    assert raw.iloc[client_column_row, 1] == "CFOP source"
    assert pd.isna(raw.iloc[client_column_row, 2])
    assert raw.iloc[client_column_row, 3] == "Corrected NCM"

    assert raw.iloc[FIRST_DATA_ROW, 1] == "5102"
    assert pd.isna(raw.iloc[FIRST_DATA_ROW, 2])
    assert raw.iloc[FIRST_DATA_ROW, 3] == "85167910"
    assert raw.iloc[FIRST_DATA_ROW + 1, 1] == "6102"
    assert pd.isna(raw.iloc[FIRST_DATA_ROW + 1, 2])
    assert raw.iloc[FIRST_DATA_ROW + 1, 3] == "73211100"

    worksheet = load_workbook(files.output)[MAPPED_SHEET]
    assert worksheet.freeze_panes == f"B{len(MAPPING_METADATA) + 2}"
    assert worksheet.cell(banner_excel_row("name_confidence"), 2).number_format == "0.00%"
    llm_confidence_excel_row = banner_excel_row("llm_final_confidence")
    assert worksheet.cell(llm_confidence_excel_row, 3).value == 0.89
    assert worksheet.cell(llm_confidence_excel_row, 3).number_format == "0.00%"


def test_banner_labels_are_the_portuguese_headers_and_confidences_are_keyed_by_field() -> None:
    # A porcentagem é escolhida pelo campo, então trocar o rótulo não a quebra.
    labels = {field: text for text, field in MAPPING_METADATA}
    assert labels == schema.COLUMNS["banner"]
    assert CONFIDENCE_ROWS <= {field for _, field in MAPPING_METADATA}
    assert FIRST_DATA_EXCEL_ROW == 23


def test_deliverable_is_written_in_portuguese(files) -> None:
    write_inputs(files)
    mapping = mapping_rows()
    mapping["catalog_name"] = ["cfop", "", "ibge_municipality"]
    mapping["decided_by"] = ["lexical", "resolver", ""]
    mapping["supported_by"] = ["lexical + catalog (agrees)", "resolver + profile", ""]
    mapping["vetoed_by"] = ["", "", "vacuity"]
    write_mapping(mapping, files.mapping)

    apply(files)

    book = load_workbook(files.output)
    assert book.sheetnames == [MAPPED_SHEET, COVERAGE_SHEET, "Colunas do Cliente sem Uso"]
    sheet = book[MAPPED_SHEET]
    assert sheet["A1"].value == "__info_mapeamento__"

    def banner(field: str) -> list[object]:
        return [sheet.cell(banner_excel_row(field), column).value for column in (2, 3, 4)]

    assert banner("status") == ["aceito", "revisao", "sem_correspondencia"]
    assert banner("method") == ["nome", "dados", "dados"]
    assert banner("catalog_name") == ["cfop", None, "municipio_ibge"]
    assert banner("decided_by") == ["lexico", "resolvedor_ia", None]
    assert banner("supported_by") == [
        "lexico + catalogo (concorda)", "resolvedor_ia + perfil", None,
    ]
    assert banner("vetoed_by") == [None, None, "coluna_vazia"]
    assert banner("reviewer_status") == [None, None, "incorreto"]
    # Texto livre (justificativa e observação do revisor) é copiado como está.
    assert banner("reason") == mapping["reason"].tolist()
    assert banner("reviewer_note") == [None, "Check this value field.",
                                       "Reviewer supplied the source column."]


def test_apply_prints_the_counts_and_no_warning(files, capsys) -> None:
    write_inputs(files)
    write_mapping(mapping_rows(), files.mapping)

    apply(files)

    captured = capsys.readouterr()
    assert "colunas_modelo=3 linhas_cliente=1 colunas_mapeadas=2" in captured.out
    assert captured.err == ""


def test_include_review_populates_unconfirmed_proposal(files) -> None:
    write_inputs(files)
    write_mapping(mapping_rows(), files.mapping)

    apply(files, accepted_only=False)

    assert read_mapped(files.output).iloc[FIRST_DATA_ROW, 2] == 12.5


def test_include_review_allows_duplicate_unconfirmed_proposals(files) -> None:
    write_workbook(files.model, ["Base CBS", "Base IBS"], [[100.0, 100.0]])
    write_workbook(files.client, ["Generic base"], [[125.0]])
    mapping = mapping_rows().iloc[:2].copy()
    mapping["model_column_number"] = [1, 2]
    mapping["model_column"] = ["Base CBS", "Base IBS"]
    mapping["client_column_number"] = [1, 1]
    mapping["client_column"] = ["Generic base", "Generic base"]
    mapping["status"] = ["review", "review"]
    mapping["reviewer_status"] = ["", ""]
    write_mapping(mapping, files.mapping)

    apply(files, accepted_only=False)

    assert read_mapped(files.output).iloc[FIRST_DATA_ROW, 1:].tolist() == [125, 125]


def test_mapped_workbook_preserves_excel_time_values(files) -> None:
    write_workbook(files.model, ["Criado às:"], [[time(8, 0, 0)]])
    write_workbook(files.client, ["Created at"], [[None]])
    client_workbook = load_workbook(files.client)
    client_workbook.active["A2"] = time(15, 30, 11)
    client_workbook.active["A2"].number_format = "hh:mm:ss"
    client_workbook.save(files.client)
    write_mapping(one_row_mapping("Criado às:", "Created at"), files.mapping)

    apply(files)

    worksheet = load_workbook(files.output, data_only=True)[MAPPED_SHEET]
    cell = worksheet.cell(len(MAPPING_METADATA) + 2, 2)
    assert cell.value == time(15, 30, 11)
    assert cell.number_format == "hh:mm:ss"


def test_accepted_only_leaves_unreviewed_review_mapping_blank(files) -> None:
    write_inputs(files)
    mapping = mapping_rows()
    mapping.loc[2, "reviewer_status"] = "no_mapping_exists"
    # Só 'incorrect' usa correção, então a da linha sai junto com o parecer.
    mapping.loc[2, "correct_client_column"] = ""
    write_mapping(mapping, files.mapping)

    apply(files, accepted_only=True)

    raw = read_mapped(files.output)
    assert raw.iloc[FIRST_DATA_ROW, 1] == "5102"
    assert pd.isna(raw.iloc[FIRST_DATA_ROW, 2])
    assert pd.isna(raw.iloc[FIRST_DATA_ROW, 3])


def test_normalized_reviewer_status_is_applied_and_displayed(files) -> None:
    # O banner mostra a grafia canônica, seja qual for a digitada.
    write_inputs(files)
    report = external_report(mapping_rows())
    report.loc[2, label("reviewer_status")] = " INCORRÉTO "
    save_report(files.mapping, report)

    apply(files)

    sheet = load_workbook(files.output)[MAPPED_SHEET]
    assert sheet.cell(FIRST_DATA_EXCEL_ROW, 4).value == "85167910"
    assert sheet.cell(banner_excel_row("reviewer_status"), 4).value == "incorreto"


def test_unknown_reviewer_status_is_refused(files) -> None:
    write_inputs(files)
    report = external_report(mapping_rows())
    report.loc[0, label("reviewer_status")] = "talvez"
    save_report(files.mapping, report)

    with pytest.raises(ValueError) as raised:
        apply(files)
    message = str(raised.value)
    for detail in ("parecer_revisor", "talvez", "CFOP", "correto", "incorreto", "inexistente"):
        assert detail in message


def test_incorrect_review_requires_a_corrected_column(files) -> None:
    write_inputs(files)
    mapping = mapping_rows()
    mapping.loc[2, "correct_client_column"] = ""
    write_mapping(mapping, files.mapping)

    with pytest.raises(
        ValueError, match=f"{label('correct_client_column')} está em branco"
    ):
        apply(files)


def test_unreviewed_accepted_semantic_conflict_is_rejected(files) -> None:
    write_workbook(files.model, ["Base ICMS_ST"], [[100.0]])
    write_workbook(files.client, ["ICMS Base"], [[100.0]])
    write_mapping(one_row_mapping("Base ICMS_ST", "ICMS Base"), files.mapping)

    with pytest.raises(ValueError, match="viola as regras atuais") as raised:
        apply(files)
    # A descrição sai em português; o texto que vai à IA continua em inglês.
    assert "tratamentos tributários diferentes (ST vs normal)" in str(raised.value)


def write_report_with_metadata(files: Files, sheet_name: str, metadata: pd.DataFrame) -> None:
    """Cria modelo e exportação de uma coluna e grava um relatório com os metadados informados."""
    write_workbook(files.model, ["CFOP"], [["5102"]])
    write_workbook(files.client, ["CFOP source"], [["5102"]])
    with pd.ExcelWriter(files.mapping, engine="openpyxl") as writer:
        external_report(one_row_mapping("CFOP", "CFOP source")).to_excel(
            writer, sheet_name="Mapeamentos", index=False
        )
        metadata.to_excel(writer, sheet_name=sheet_name, index=False)


def test_run_metadata_detects_changed_inputs(files) -> None:
    # Nomes de aba, de coluna e de chave escritos à mão, como nos relatórios
    # arquivados: se não forem reconhecidos, a verificação é pulada em silêncio.
    write_report_with_metadata(
        files,
        "Metadados da Execução",
        pd.DataFrame([["sha256_cliente", "nao-e-o-hash-atual"]], columns=["chave", "valor"]),
    )

    with pytest.raises(ValueError, match="client.xlsx difere do arquivo usado"):
        apply(files)


@pytest.mark.parametrize(
    ("metadata_sheet", "metadata_columns", "message"),
    [
        ("Run Metadata", ["key", "value"], "Renomeie 'Run Metadata'"),
        ("Metadados da Execução", ["key", "value"], "Renomeie 'key' para 'chave'"),
    ],
    ids=["aba_antiga", "cabecalho_antigo"],
)
def test_metadata_with_old_names_is_refused(
    files, metadata_sheet, metadata_columns, message
) -> None:
    # Ignorada, a aba faria a verificação de hash ser pulada sem aviso.
    write_report_with_metadata(
        files, metadata_sheet, pd.DataFrame([["client_sha256", "x"]], columns=metadata_columns)
    )

    with pytest.raises(ValueError, match=re.escape(message)):
        apply(files)


def test_a_correction_outside_the_incorrect_branch_is_refused(files) -> None:
    # Com 'correct', descartar a correção entregaria a coluna do mapeador, e
    # não a do revisor, com exit code 0.
    write_inputs(files)
    for status in ("correct", "no_mapping_exists", ""):
        mapping = mapping_rows()
        mapping.loc[0, "reviewer_status"] = status
        mapping.loc[0, "correct_client_column"] = 4
        write_mapping(mapping, files.mapping)
        with pytest.raises(ValueError, match="usa uma correção"):
            apply(files)


def test_confirming_a_row_that_names_no_column_is_refused(files) -> None:
    # Sem o guard, a linha entregaria uma coluna vazia marcada como confirmada pelo revisor.
    write_inputs(files)
    mapping = mapping_rows()
    mapping.loc[2, "reviewer_status"] = "correct"
    mapping.loc[2, "correct_client_column"] = ""
    write_mapping(mapping, files.mapping)
    with pytest.raises(
        ValueError, match=f"{label('client_column_number')} está em branco"
    ):
        apply(files)


def test_duplicated_header_is_selected_by_the_occurrence_suffix(files) -> None:
    # Em memória a segunda ocorrência é 'Valor ICMS (occurrence 2)'; nas
    # planilhas, 'Valor ICMS (ocorrência 2)'.
    write_workbook(files.model, ["Valor ICMS", "CFOP"], [[1.0, "5102"]])
    write_workbook(files.client, ["Valor ICMS", "Valor ICMS"], [[10.0, 20.0]])

    proposal = mapping_rows().iloc[:2].copy()
    proposal["model_column_number"] = [1, 2]
    proposal["model_column"] = ["Valor ICMS", "CFOP"]
    proposal["client_column_number"] = [2, None]
    proposal["client_column"] = ["Valor ICMS (occurrence 2)", None]
    proposal["status"] = ["accepted", "no_match"]

    # A proposta aceita aponta a segunda ocorrência pelo nome de exibição.
    write_mapping(proposal, files.mapping)
    stored = pd.read_excel(files.mapping, dtype=object)
    assert stored.loc[0, "coluna_cliente"] == "Valor ICMS (ocorrência 2)"
    apply(files)
    sheet = load_workbook(files.output)[MAPPED_SHEET]
    assert sheet.cell(FIRST_DATA_EXCEL_ROW, 2).value == 20.0
    assert sheet.cell(banner_excel_row("client_column"), 2).value == "Valor ICMS (ocorrência 2)"
    coverage = list(load_workbook(files.output)[COVERAGE_SHEET].values)
    delivered = dict(zip(coverage[0], coverage[1]))
    assert delivered["coluna_cliente"] == "Valor ICMS (ocorrência 2)"

    # A correção do revisor aceita o sufixo, com ou sem acento, e o número da coluna.
    for typed in (
        "Valor ICMS (ocorrência 2)",
        "valor icms (OCORRENCIA 2)",
        "2",
    ):
        corrected = proposal.copy()
        corrected["client_column_number"] = [1, None]
        corrected["client_column"] = ["Valor ICMS", None]
        corrected["status"] = ["review", "no_match"]
        corrected["reviewer_status"] = ["incorrect", ""]
        corrected["correct_client_column"] = [typed, ""]
        write_mapping(corrected, files.mapping)
        apply(files)
        sheet = load_workbook(files.output)[MAPPED_SHEET]
        assert sheet.cell(FIRST_DATA_EXCEL_ROW, 2).value == 20.0, typed

    # O nome sem sufixo é o nome de exibição da primeira ocorrência e a
    # seleciona (comportamento documentado em docs/DATA.md).
    corrected["correct_client_column"] = ["Valor ICMS", ""]
    write_mapping(corrected, files.mapping)
    apply(files)
    assert load_workbook(files.output)[MAPPED_SHEET].cell(FIRST_DATA_EXCEL_ROW, 2).value == 10.0

    # Uma ocorrência que não existe não é trocada por outra: é recusada.
    corrected["correct_client_column"] = ["Valor ICMS (ocorrência 3)", ""]
    write_mapping(corrected, files.mapping)
    with pytest.raises(
        ValueError, match=re.escape("'Valor ICMS (ocorrência 3)' não foi encontrada")
    ):
        apply(files)


def test_columns_with_no_data_are_hidden_not_removed(files) -> None:
    write_inputs(files)
    write_mapping(mapping_rows(), files.mapping)
    apply(files)
    sheet = load_workbook(files.output)[MAPPED_SHEET]
    hidden = {
        get_column_letter(c): sheet.column_dimensions[get_column_letter(c)].hidden
        for c in range(2, 5)
    }
    # Entregues: 'CFOP', aceita, e 'Nº do NCM', pela correção do revisor. Só
    # 'Valor COFINS', em revisão não confirmada, fica sem dado.
    assert hidden["B"] is False
    assert hidden["C"] is True
    assert hidden["D"] is False
    # Ocultar não remove nada: as três colunas do modelo continuam lá.
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    assert list(header[1:4]) == ["CFOP", "Valor COFINS", "Nº do NCM"]


def test_coverage_sheet_separates_absence_from_a_pending_decision(files) -> None:
    # Distingue as colunas sem proposta das propostas que aguardam revisão.
    # Ambas ficam vazias na saída, mas recebem estados diferentes na cobertura.
    write_workbook(files.model, [*MODEL_HEADERS, "Valor IPI"],
                   [["5102", 10.0, "85167910", 1.0]])
    write_workbook(files.client, CLIENT_HEADERS, [CLIENT_ROW])
    mapping = pd.concat([mapping_rows(), mapping_rows().iloc[[2]]], ignore_index=True)
    mapping.loc[3, ["model_column_number", "model_column"]] = [4, "Valor IPI"]
    mapping.loc[3, ["reviewer_status", "correct_client_column", "reviewer_note"]] = ["", "", ""]
    write_mapping(mapping, files.mapping)
    apply(files)
    book = load_workbook(files.output)
    assert COVERAGE_SHEET in book.sheetnames
    rows = list(book[COVERAGE_SHEET].values)
    header = list(rows[0])
    assert header == list(schema.COLUMNS["coverage"].values())
    data = [dict(zip(header, r)) for r in rows[1:]]
    states = [r["situacao_cobertura"] for r in data]
    # Entregues primeiro, depois as pendentes, depois as ausentes.
    assert states == [
        "entregue", "entregue", "aguardando decisão", "não consta na exportação do cliente",
    ]
    assert data[0]["coluna_modelo"] == "CFOP" and data[0]["coluna_excel"] == "B"
    # Uma correção do revisor conta como entregue e cita a coluna onde caiu.
    assert data[1]["coluna_modelo"] == "Nº do NCM"
    assert data[1]["coluna_cliente"] == "Corrected NCM"
    # A linha pendente cita o candidato que o revisor vai julgar.
    assert data[2]["coluna_cliente"] == "COFI value"
    assert data[2]["justificativa"] == "Names and numeric values are compatible."
    assert data[3]["coluna_modelo"] == "Valor IPI"


def test_coverage_states_are_the_documented_portuguese_states() -> None:
    assert (STATE_DELIVERED, STATE_PENDING, STATE_ABSENT) == (
        "entregue", "aguardando decisão", "não consta na exportação do cliente",
    )


def test_command_line_is_in_portuguese(capsys) -> None:
    help_text = build_parser().format_help()
    assert help_text.startswith("uso: ")
    assert "opções:" in help_text
    assert "exportação .xlsx do cliente" in help_text
    # As flags continuam em inglês.
    assert "--include-review" in help_text and "--output" in help_text

    with pytest.raises(SystemExit) as raised:
        main(["c.xlsx", "r.xlsx", "--output", "o.xlsx", "--accepted-only", "--include-review"])
    assert str(raised.value) == "--accepted-only e --include-review não podem ser usados juntos"

    with pytest.raises(SystemExit) as raised:
        main([])
    assert raised.value.code == 2
    assert "os seguintes argumentos são obrigatórios" in capsys.readouterr().err


def test_apply_coverage_and_evaluation_use_the_same_review_decision(tmp_path):
    rows = [
        dict(model_column_number=1, model_column="Target A", status="accepted", reviewer_status="no_mapping_exists"),
        dict(model_column_number=2, model_column="Target B", client_column_number=2,
             client_column="Source B", reviewer_status="correct"),
        dict(model_column_number=3, model_column="Target C", reviewer_status="incorrect",
             correct_client_column="Source C"),
    ]
    client, report = report_files(tmp_path, rows)
    model, output = tmp_path / "model.xlsx", tmp_path / "output.xlsx"
    pd.DataFrame([["Target A", "Target B", "Target C"]]).to_excel(model, index=False, header=False)
    build_mapped_workbook(client, report, output, model)
    coverage = pd.read_excel(output, sheet_name="Cobertura").set_index("numero_coluna_modelo")
    assert coverage.loc[1, "situacao_cobertura"] == "não consta na exportação do cliente"
    assert coverage.loc[2, "situacao_cobertura"] == "entregue"
    assert coverage.loc[3, "numero_coluna_cliente"] == 3
    score = score_report(report, client, AnswerKey("synthetic", {1: 1, 2: 2, 3: 3}))
    assert [r.status for r in score.rows] == ["no_match", "accepted", "accepted"]
    assert [r.client_column for r in score.rows] == [None, 2, 3]
    assert score.delivered_correct == 2
    assert not score.withheld_correct


@pytest.mark.parametrize("include_review", [False, True])
def test_unused_sources_follow_review_and_keep_pending_and_empty_columns(tmp_path, include_review):
    client, report = report_files(tmp_path, [
        dict(model_column_number=1, model_column="Target A", status="accepted",
             reviewer_status="incorrect", correct_client_column="Source B"),
        dict(model_column_number=2, model_column="Target B", client_column_number=3,
             client_column="Source C", status="review"),
    ])
    pd.DataFrame([["Source A", "Source B", "Source C", "Empty"], [11, 22, 33, None]]).to_excel(
        client, header=False, index=False)
    model, output = tmp_path / "model.xlsx", tmp_path / "output.xlsx"
    pd.DataFrame([["Target A", "Target B"]]).to_excel(model, header=False, index=False)
    original_report = report.read_bytes()
    build_mapped_workbook(client, report, output, model, accepted_only=not include_review)
    unused = schema.to_internal_frame(
        pd.read_excel(output, sheet_name="Colunas do Cliente sem Uso"), "unmatched")
    assert dict(zip(unused.client_column_number, unused.disposition)) == {
        1: "never_referenced", 3: "proposed_not_accepted", 4: "never_referenced"}
    assert list(unused.rows_with_value) == [1, 1, 0]
    assert report.read_bytes() == original_report
