"""Testes do esquema externo em português (report_schema) e da CLI em português (cli_ptbr).

Três grupos:

* invariantes do vocabulário: nomes de aba válidos no Excel, cabeçalhos em
  snake_case ASCII, nenhum nome em português igual a um legado, valores
  distintos em cada enum;
* cobertura por introspecção: todo campo que o código grava tem tradução,
  de modo que um campo novo sem nome em português quebra o teste;
* conversões: ida e volta interno -> português -> interno, leitura do
  legado, esquema misto, parecer do revisor, nome de exibição, etapas,
  marcador da prova e números com vírgula decimal.
"""

from __future__ import annotations

import argparse
import ast
import dataclasses
import math
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

import apply_invoice_mapping
import invoice_mapper
import mapping_steps
import report_schema as rs
from cli_ptbr import PortugueseArgumentParser
from invoice_mapper import (
    AdjudicationRow,
    MappingResult,
    REVIEWER_COLUMNS,
    UnmatchedColumn,
)

ROOT = Path(__file__).resolve().parent.parent
SNAKE_CASE = re.compile(r"[a-z0-9]+(?:_[a-z0-9]+)*")
ARCHIVED_REPORTS = sorted((ROOT / "data" / "reports").glob("mapping_*_full.xlsx"))
requires_reports = pytest.mark.skipif(
    not ARCHIVED_REPORTS, reason="relatórios arquivados ausentes"
)
REPORT_TABLES = ("mappings", "run_metadata", "unmatched", "queue")


# ---------------------------------------------------------------------------
# Invariantes do vocabulário
# ---------------------------------------------------------------------------


def test_sheet_names_are_legal_in_excel(tmp_path) -> None:
    book = openpyxl.Workbook()
    for legacy, portuguese in rs.SHEETS.values():
        assert len(portuguese) <= 31, portuguese
        assert not set(portuguese) & set(":\\/?*[]"), portuguese
        assert not portuguese.startswith("'") and not portuguese.endswith("'")
        book.create_sheet(portuguese)
    path = tmp_path / "abas.xlsx"
    book.save(path)

    reloaded = openpyxl.load_workbook(path).sheetnames
    assert {portuguese for _, portuguese in rs.SHEETS.values()} <= set(reloaded)
    assert len({portuguese for _, portuguese in rs.SHEETS.values()}) == len(rs.SHEETS)


def test_headers_are_snake_case_ascii() -> None:
    for table, columns in rs.COLUMNS.items():
        for portuguese in columns.values():
            if table == "mapped_data":
                assert portuguese == "__info_mapeamento__"
                continue
            assert SNAKE_CASE.fullmatch(portuguese), (table, portuguese)
    for portuguese in rs.METADATA_KEYS.values():
        assert SNAKE_CASE.fullmatch(portuguese), portuguese


def test_no_portuguese_name_equals_a_legacy_one() -> None:
    legacy = (
        {field for columns in rs.COLUMNS.values() for field in columns}
        | set(rs.METADATA_KEYS)
        | {legacy for legacy, _ in rs.SHEETS.values()}
    )
    portuguese = (
        {name for columns in rs.COLUMNS.values() for name in columns.values()}
        | set(rs.METADATA_KEYS.values())
        | {portuguese for _, portuguese in rs.SHEETS.values()}
    )
    assert not legacy & portuguese, sorted(legacy & portuguese)


def test_headers_and_keys_are_distinct_within_each_table() -> None:
    for table, columns in rs.COLUMNS.items():
        assert len(set(columns.values())) == len(columns), table
    assert len(set(rs.METADATA_KEYS.values())) == len(rs.METADATA_KEYS)


def test_enum_values_are_distinct_and_translation_is_idempotent() -> None:
    for name, values in rs.ENUMS.items():
        assert len(set(values.values())) == len(values), name
        for internal, portuguese in values.items():
            # Um valor em português que também é chave do mesmo enum teria de
            # ser a tradução de si mesmo; senão traduzir duas vezes mudaria o valor.
            if portuguese in values:
                assert values[portuguese] == portuguese, (name, internal, portuguese)
    for name in rs.MACHINE_ENUMS:
        for portuguese in rs.ENUMS[name].values():
            assert portuguese == "" or SNAKE_CASE.fullmatch(portuguese), (name, portuguese)


def test_step_ids_are_snake_case_and_distinct() -> None:
    portuguese = list(rs.ALL_STEP_IDS.values())
    assert len(set(portuguese)) == len(portuguese)
    assert not set(portuguese) & set(rs.ALL_STEP_IDS)
    for step in portuguese:
        assert SNAKE_CASE.fullmatch(step), step


# ---------------------------------------------------------------------------
# Cobertura por introspecção
# ---------------------------------------------------------------------------


def test_every_mapping_result_field_has_a_header() -> None:
    fields = [field.name for field in dataclasses.fields(MappingResult)] + list(REVIEWER_COLUMNS)
    assert fields == list(rs.COLUMNS["mappings"])


def test_every_unmatched_column_field_has_a_header() -> None:
    fields = [field.name for field in dataclasses.fields(UnmatchedColumn)]
    assert fields == list(rs.COLUMNS["unmatched"])


def test_every_adjudication_row_field_has_a_header() -> None:
    fields = [field.name for field in dataclasses.fields(AdjudicationRow)]
    assert fields == list(rs.COLUMNS["queue"])


def test_every_banner_row_has_a_label() -> None:
    fields = [field for _, field in apply_invoice_mapping.MAPPING_METADATA]
    assert fields == list(rs.COLUMNS["banner"])
    assert set(fields) <= set(rs.COLUMNS["mappings"])


def test_every_reviewer_status_has_a_translation() -> None:
    assert set(apply_invoice_mapping.REVIEW_STATUSES) == set(rs.ENUMS["reviewer_status"])


def test_every_coverage_state_has_a_translation() -> None:
    states = rs.ENUMS["coverage_state"]
    known = set(states) | set(states.values())
    for state in (
        apply_invoice_mapping.STATE_DELIVERED,
        apply_invoice_mapping.STATE_PENDING,
        apply_invoice_mapping.STATE_ABSENT,
    ):
        assert state in known, state


def test_every_analysis_step_and_guard_has_a_translation() -> None:
    assert {step.id for step in mapping_steps.EVIDENCE_STEPS} == set(rs.STEP_IDS)
    assert {step.id for step in mapping_steps.GUARDS} == set(rs.GUARD_IDS)
    assert set(mapping_steps.STEPS_BY_ID) <= set(rs.ALL_STEP_IDS)


def test_every_status_and_method_literal_has_a_translation() -> None:
    source = (ROOT / "invoice_mapper.py").read_text(encoding="utf-8")
    statuses = set(re.findall(r'status="(\w+)"', source))
    methods = set(re.findall(r'method="(\w+)"', source))
    assert statuses and methods
    assert statuses <= set(rs.ENUMS["status"])
    assert methods <= set(rs.ENUMS["method"])


def test_every_queue_reason_has_a_translation() -> None:
    # As três frases de why_held, uma por ramo: abaixo do limiar, só nome
    # (confiança presa logo abaixo do limiar) e coluna proposta vazia.
    client = invoice_mapper.SheetTable(
        columns=[
            invoice_mapper.ColumnRef("column_001", 0, "Cheia", 1),
            invoice_mapper.ColumnRef("column_002", 1, "Vazia", 1),
        ],
        data=pd.DataFrame([[1, None], [2, None]]),
    )

    def proposal(number: int, client_number: int, final: float) -> MappingResult:
        return MappingResult(
            model_column_number=number, model_column=f"M{number}",
            client_column_number=client_number, client_column="x",
            name_confidence=final, name_runner_up_confidence=0.0, name_margin=0.0,
            data_confidence=None, final_confidence=final, method="data",
            status="review", reason="",
        )

    rows = invoice_mapper.summarize_adjudication_queue(
        client_table=client,
        results=[proposal(1, 1, 0.5), proposal(2, 1, 0.8999), proposal(3, 2, 0.95)],
        final_threshold=0.9,
        sample_size=8,
    )
    phrases = rs.ENUMS["why_held"]
    held = {row.why_held for row in rows}
    assert len(held) == 3
    assert held <= set(phrases) | set(phrases.values())


def test_every_run_metadata_key_and_sentinel_has_a_translation(tmp_path, monkeypatch) -> None:
    model_path = tmp_path / "modelo.xlsx"
    client_path = tmp_path / "cliente.xlsx"
    pd.DataFrame([["CFOP", "NCM"], ["5102", "84714900"]]).to_excel(
        model_path, header=False, index=False
    )
    pd.DataFrame([["CFOP"], ["5102"]]).to_excel(client_path, header=False, index=False)
    written: dict[str, object] = {}

    def capture(results, output_path, metadata=None, unmatched=None, queue=None):
        written.update(metadata or {})

    monkeypatch.setattr(invoice_mapper, "write_results", capture)
    invoice_mapper.main(
        [
            str(client_path),
            "--model", str(model_path),
            "--output", str(tmp_path / "relatorio.xlsx"),
            "--offline",
            "--no-reference-data",
        ]
    )

    assert written
    assert set(written) <= set(rs.METADATA_KEYS)
    for key in rs.METADATA_SENTINEL_KEYS:
        assert written[key] in rs.METADATA_SENTINELS, (key, written[key])


# ---------------------------------------------------------------------------
# Ida e volta
# ---------------------------------------------------------------------------


def _internal_mappings() -> pd.DataFrame:
    rows = [
        {
            "model_column_number": 1, "model_column": "CFOP",
            "client_column_number": 3, "client_column": "Cod. Fiscal",
            "name_confidence": 0.61, "status": "accepted", "method": "catalog",
            "reason": "Versioned official cfop codes cover 88.2% ...",
            "catalog_name": "ibge_municipality", "arithmetic_proof": None,
            "supported_by": "lexical + catalog", "decided_by": "catalog", "vetoed_by": "",
            "reviewer_status": "", "correct_client_column": "", "reviewer_note": "",
        },
        {
            "model_column_number": 2, "model_column": "Valor ICMS",
            "client_column_number": 12, "client_column": "Valor ICMS (occurrence 2)",
            "name_confidence": 1.0, "status": "review", "method": "data",
            "reason": "Prosa da IA, que passa intacta.",
            "catalog_name": None,
            "arithmetic_proof": "valor = base x aliquota holds for icms on 40 rows at fit 1.0",
            "supported_by": "resolver + semantic + arithmetic", "decided_by": "arithmetic",
            "vetoed_by": "family_conflict",
            "reviewer_status": "incorrect", "correct_client_column": "Unnamed column 4",
            "reviewer_note": "Nota do revisor, byte a byte.",
        },
        {
            "model_column_number": 3, "model_column": "Base PIS",
            "client_column_number": 4, "client_column": "Unnamed column 4",
            "name_confidence": 0.0, "status": "no_match", "method": "name_only",
            "reason": "", "catalog_name": "cst_icms_taxation", "arithmetic_proof": None,
            "supported_by": "lexical + catalog (agrees) + unrecorded", "decided_by": "unrecorded",
            "vetoed_by": "vacuity",
            "reviewer_status": "no_mapping_exists", "correct_client_column": 6,
            "reviewer_note": None,
        },
        {
            "model_column_number": 4, "model_column": "NCM",
            "client_column_number": None, "client_column": None,
            "name_confidence": 0.2, "status": "review", "method": "data_error",
            "reason": "", "catalog_name": "ibs_cbs_cst", "arithmetic_proof": None,
            "supported_by": "", "decided_by": "", "vetoed_by": "",
            "reviewer_status": "correct", "correct_client_column": None, "reviewer_note": "",
        },
    ]
    return pd.DataFrame(rows)


def _internal_unmatched() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "client_column_number": 7, "client_column": "Direito (occurrence 3)",
                "disposition": "proposed_not_accepted", "review_priority": "shortlist",
                "rows_with_value": 12, "distinct_values": 3,
                "signals": "arithmetic,catalog,cnpj_checkdigit,date",
                "signal_detail": "holds dates", "closest_model_column": "Data de Emissão",
                "closest_name_confidence": 0.5,
            },
            {
                "client_column_number": 9, "client_column": "Unnamed column 9",
                "disposition": "never_referenced", "review_priority": "low",
                "rows_with_value": 0, "distinct_values": 0, "signals": "",
                "signal_detail": "No values in this column.",
                "closest_model_column": "CFOP", "closest_name_confidence": 0.1,
            },
        ]
    )


def _internal_queue() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "rank": 1, "model_column_number": 2, "model_column": "Valor ICMS",
                "proposed_client_column_number": 12,
                "proposed_client_column": "Valor ICMS (occurrence 2)",
                "final_confidence": 0.8999, "gate_distance": 0.0001, "name_margin": 0.0,
                "rows_with_value": 40, "resolver_disagrees": "", "arithmetic_proof": "",
                "client_sample_values": "1 | 2", "why_held": "below the accept threshold",
            },
        ]
    )


@pytest.mark.parametrize(
    ("table", "frame"),
    [
        ("mappings", _internal_mappings()),
        ("unmatched", _internal_unmatched()),
        ("queue", _internal_queue()),
    ],
)
def test_internal_to_portuguese_and_back_is_lossless(table, frame) -> None:
    original = frame.copy()

    external = rs.to_external_frame(frame, table)

    assert list(external.columns) == [rs.COLUMNS[table][name] for name in frame.columns]
    back, schema = rs.to_internal_frame(external, table)
    assert schema == rs.PT_BR
    assert list(back.columns) == list(frame.columns)
    one_way = {
        field
        for (kind_table, field), kind in rs.FIELD_KIND.items()
        if kind_table == table and kind.startswith("phrase:")
    }
    for column in frame.columns:
        if column in one_way:
            continue
        assert back[column].equals(frame[column]), column
    # Nenhuma das conversões altera o quadro recebido.
    assert frame.equals(original)
    assert rs.to_internal_frame(external, table)[0].equals(back)


def _blank_as_none(values: list[object]) -> list[object]:
    # O Excel não distingue texto vazio de célula vazia.
    return [None if value == "" or pd.isna(value) else value for value in values]


def test_a_report_written_in_portuguese_reads_back_through_excel(tmp_path) -> None:
    path = tmp_path / "relatorio.xlsx"
    tables = {"mappings": _internal_mappings(), "unmatched": _internal_unmatched(), "queue": _internal_queue()}
    metadata = {"mapping_schema_version": rs.SCHEMA_VERSION_PT, "reference_manifest_sha256": "disabled"}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        rs.to_external_frame(tables["mappings"], "mappings").to_excel(
            writer, sheet_name=rs.sheet_name("mappings"), index=False
        )
        rs.to_external_metadata(metadata).to_excel(
            writer, sheet_name=rs.sheet_name("run_metadata"), index=False
        )
        for table in ("unmatched", "queue"):
            rs.to_external_frame(tables[table], table).to_excel(
                writer, sheet_name=rs.sheet_name(table), index=False
            )

    book = pd.read_excel(path, sheet_name=None, dtype=object)
    names = list(book)
    assert names == ["Mapeamentos", "Metadados da Execução", "Colunas do Cliente sem Uso", "Fila de Revisão"]
    schema = rs.detect_schema(book[names[0]].columns)
    assert schema == rs.PT_BR
    for table, frame in tables.items():
        back, found = rs.to_internal_frame(book[rs.find_sheet(names, table, schema)], table)
        assert found == schema
        for column in frame.columns:
            if rs.FIELD_KIND.get((table, column), "").startswith("phrase:"):
                continue
            assert _blank_as_none(back[column].tolist()) == _blank_as_none(frame[column].tolist()), column
    read_metadata = rs.internal_metadata(book[rs.find_sheet(names, "run_metadata", schema)], schema)
    assert read_metadata == {"mapping_schema_version": "3", "reference_manifest_sha256": "disabled"}


def test_the_external_mappings_frame_is_in_portuguese() -> None:
    external = rs.to_external_frame(_internal_mappings(), "mappings").to_dict("records")

    assert [row["situacao"] for row in external] == ["aceito", "revisao", "sem_correspondencia", "revisao"]
    assert [row["metodo"] for row in external] == ["catalogo", "dados", "somente_nome", "erro_dados"]
    assert external[0]["nome_catalogo"] == "municipio_ibge"
    assert external[2]["nome_catalogo"] == "cst_icms_tributacao"
    assert external[3]["nome_catalogo"] == "cst_ibs_cbs"
    assert external[1]["coluna_cliente"] == "Valor ICMS (ocorrência 2)"
    assert external[1]["coluna_cliente_correta"] == "Coluna sem nome 4"
    assert external[2]["coluna_cliente"] == "Coluna sem nome 4"
    assert external[2]["coluna_cliente_correta"] == 6
    assert external[1]["apoiado_por"] == "resolvedor_ia + semantico + aritmetica"
    assert external[2]["apoiado_por"] == "lexico + catalogo (concorda) + nao_registrado"
    assert external[2]["decidido_por"] == "nao_registrado"
    assert [row["vetado_por"] for row in external] == ["", "conflito_tributario", "coluna_vazia", ""]
    assert [row["parecer_revisor"] for row in external] == ["", "incorreto", "inexistente", "correto"]
    # Texto livre e dados passam intactos.
    assert external[1]["justificativa"] == "Prosa da IA, que passa intacta."
    assert external[1]["prova_aritmetica"].startswith("valor = base x aliquota holds")
    assert external[1]["observacao_revisor"] == "Nota do revisor, byte a byte."
    assert external[1]["coluna_modelo"] == "Valor ICMS"


def test_the_external_side_tables_are_in_portuguese() -> None:
    unmatched = rs.to_external_frame(_internal_unmatched(), "unmatched").to_dict("records")
    assert unmatched[0]["coluna_cliente"] == "Direito (ocorrência 3)"
    assert unmatched[0]["disposicao"] == "proposta_nao_aceita"
    assert unmatched[0]["prioridade_revisao"] == "prioritaria"
    assert unmatched[0]["sinais"] == "aritmetica,catalogo,dv_cnpj,data"
    assert unmatched[1]["coluna_cliente"] == "Coluna sem nome 9"
    assert unmatched[1]["disposicao"] == "nunca_referenciada"
    assert unmatched[1]["prioridade_revisao"] == "baixa"
    assert unmatched[1]["sinais"] == ""

    queue = rs.to_external_frame(_internal_queue(), "queue").to_dict("records")
    assert queue[0]["coluna_cliente_proposta"] == "Valor ICMS (ocorrência 2)"
    assert queue[0]["motivo_retencao"] == "abaixo do limiar de aceite"
    assert queue[0]["amostra_valores_cliente"] == "1 | 2"


def test_a_legacy_report_reads_as_internal_without_change() -> None:
    legacy = _internal_mappings()
    legacy.loc[0, "reviewer_status"] = " Correct "

    internal, schema = rs.to_internal_frame(legacy, "mappings")

    assert schema == rs.LEGACY_EN
    expected = _internal_mappings()
    expected.loc[0, "reviewer_status"] = "correct"
    for column in expected.columns:
        assert internal[column].equals(expected[column]), column


def test_a_hand_written_portuguese_report_reads_as_internal() -> None:
    frame = pd.DataFrame(
        [
            {
                "numero_coluna_modelo": 1, "coluna_modelo": "CFOP",
                "numero_coluna_cliente": 2, "coluna_cliente": "CFOP (ocorrencia 2)",
                "situacao": "aceito", "metodo": "nome", "nome_catalogo": "municipio_ibge",
                "apoiado_por": "lexico + catalog (agrees)", "decidido_por": "lexico",
                "vetado_por": "", "parecer_revisor": "Incorreto",
                "coluna_cliente_correta": "coluna sem nome 5",
                "colunaExtra": "fica como está",
            }
        ]
    )

    internal, schema = rs.to_internal_frame(frame, "mappings")
    row = internal.to_dict("records")[0]

    assert schema == rs.PT_BR
    assert row["model_column_number"] == 1
    assert row["client_column"] == "CFOP (occurrence 2)"
    assert row["status"] == "accepted"
    assert row["method"] == "name"
    assert row["catalog_name"] == "ibge_municipality"
    assert row["supported_by"] == "lexical + catalog (agrees)"
    assert row["decided_by"] == "lexical"
    assert row["reviewer_status"] == "incorrect"
    assert row["correct_client_column"] == "Unnamed column 5"
    assert row["colunaExtra"] == "fica como está"


def test_conversions_keep_every_cell_exactly_as_it_was() -> None:
    # Como em sap_procosa lido com dtype=object: os números de coluna
    # digitados pelo revisor são inteiros numa coluna object, ao lado de NaN.
    legacy = pd.DataFrame(
        {
            "model_column_number": [1, 2, 3],
            "client_column": [None, "Valor (occurrence 2)", "CFOP"],
            "reviewer_status": [math.nan, "incorrect", "incorrect"],
            "correct_client_column": pd.Series([math.nan, 6, 32], dtype=object),
            "reviewer_note": pd.Series([None, None, None], dtype=object),
        }
    )

    internal, _ = rs.to_internal_frame(legacy, "mappings")
    external = rs.to_external_frame(internal, "mappings")
    back, _ = rs.to_internal_frame(external, "mappings")

    for frame in (internal, back):
        assert frame["correct_client_column"].dtype == object
        assert [type(value) for value in frame["correct_client_column"][1:]] == [int, int]
        assert frame["correct_client_column"].tolist()[1:] == [6, 32]
        assert frame["client_column"][0] is None
        assert frame["reviewer_note"].tolist() == [None, None, None]
    assert external["coluna_cliente_correta"].tolist()[1:] == [6, 32]
    assert external["coluna_cliente"].tolist()[1] == "Valor (ocorrência 2)"


def test_conversions_work_with_repeated_column_names() -> None:
    frame = pd.DataFrame([["accepted", "review"]], columns=["status", "status"])
    frame.insert(0, "model_column_number", [1])

    external = rs.to_external_frame(frame, "mappings")

    assert list(external.columns) == ["numero_coluna_modelo", "situacao", "situacao"]
    assert external.iloc[0].tolist() == [1, "aceito", "revisao"]
    back, _ = rs.to_internal_frame(external, "mappings")
    assert back.iloc[0].tolist() == [1, "accepted", "review"]


def test_enum_values_are_converted_only_when_exact_and_only_in_portuguese() -> None:
    # Valor legado num relatório em português e valor em português num
    # relatório legado passam como estão: a comparação exata a jusante decide.
    assert rs.internal_value("mappings", "status", "accepted", rs.PT_BR) == "accepted"
    assert rs.internal_value("mappings", "status", "Aceito", rs.PT_BR) == "Aceito"
    assert rs.internal_value("mappings", "status", "aceito", rs.LEGACY_EN) == "aceito"
    assert rs.internal_value("mappings", "status", "aceito", rs.PT_BR) == "accepted"
    assert rs.internal_value("mappings", "decided_by", "lexico", rs.LEGACY_EN) == "lexico"
    assert rs.internal_value("mappings", "model_column", "Situação", rs.PT_BR) == "Situação"
    assert rs.internal_value("banner", "status", "aceito", rs.PT_BR) == "accepted"
    assert rs.external_value("banner", "decided_by", "catalog") == "catalogo"
    assert rs.external_value("banner", "final_confidence", 0.93) == 0.93


@requires_reports
@pytest.mark.parametrize("report", ARCHIVED_REPORTS, ids=lambda path: path.stem)
def test_archived_reports_have_a_translation_for_every_value(report) -> None:
    book = pd.read_excel(report, sheet_name=None, dtype=object)
    names = list(book)
    schema = rs.detect_schema(book[names[0]].columns, "mappings")
    for table in REPORT_TABLES:
        sheet = rs.find_sheet(names, rs.TABLE_SHEETS[table], schema)
        if sheet is None:
            continue
        internal, found = rs.to_internal_frame(book[sheet], table)
        assert found == schema
        if schema == rs.LEGACY_EN:
            # Um relatório legado é lido sem mudar nada, nem o tipo das
            # células (os inteiros 6 e 32 dos pareceres de sap_procosa).
            for column in book[sheet].columns:
                assert internal[column].equals(book[sheet][column]), (sheet, column)

        for (kind_table, field), kind in rs.FIELD_KIND.items():
            if kind_table != table or field not in internal.columns:
                continue
            values = [v for v in internal[field].tolist() if isinstance(v, str) and v]
            if kind.startswith(("enum:", "phrase:")):
                enum = rs.ENUMS[kind.split(":", 1)[1]]
                assert set(values) <= set(enum) | set(enum.values()), (field, set(values) - set(enum))
            elif kind == "steps":
                tokens = {
                    token.removesuffix(rs.AGREES_SUFFIX_EN)
                    for value in values
                    for token in value.split(rs.STEPS_SEPARATOR)
                }
                assert tokens <= set(rs.ALL_STEP_IDS), tokens - set(rs.ALL_STEP_IDS)
            elif kind == "signals":
                tokens = {token for value in values for token in value.split(rs.SIGNALS_SEPARATOR)}
                assert tokens <= set(rs.ENUMS["signals"]), tokens

        # Ida e volta pelo português devolve o mesmo vocabulário interno.
        back, back_schema = rs.to_internal_frame(rs.to_external_frame(internal, table), table)
        assert back_schema == rs.PT_BR
        for column in internal.columns:
            if rs.FIELD_KIND.get((table, column), "").startswith("phrase:"):
                continue
            assert back[column].equals(internal[column]), (sheet, column)

    metadata_sheet = rs.find_sheet(names, "run_metadata", schema)
    metadata = rs.internal_metadata(book[metadata_sheet], schema)
    assert set(metadata) <= set(rs.METADATA_KEYS)
    exported = rs.to_external_metadata(metadata)
    assert rs.internal_metadata(exported, rs.PT_BR) == metadata


# ---------------------------------------------------------------------------
# Esquema: detecção, abas e cabeçalhos
# ---------------------------------------------------------------------------


def test_the_schema_is_detected_from_the_marker_column() -> None:
    assert rs.detect_schema(["numero_coluna_modelo", "situacao"]) == rs.PT_BR
    assert rs.detect_schema(("model_column_number", "status", None)) == rs.LEGACY_EN
    assert rs.detect_schema(["chave", "valor"], "run_metadata") == rs.PT_BR
    assert rs.detect_schema(["key", "value"], "run_metadata") == rs.LEGACY_EN


@pytest.mark.parametrize(
    "header",
    [
        ["numero_coluna_modelo", "model_column_number"],
        ["numero_coluna_modelo", "status"],
        ["model_column_number", "parecer_revisor"],
    ],
)
def test_a_mixed_header_is_refused(header) -> None:
    with pytest.raises(ValueError, match="Esquema misto"):
        rs.detect_schema(header)


def test_a_header_without_a_marker_is_refused() -> None:
    with pytest.raises(ValueError, match="Não foi possível identificar o esquema"):
        rs.detect_schema(["coluna_modelo", "situacao"])
    with pytest.raises(ValueError, match="Não foi possível identificar o esquema"):
        rs.detect_schema([])


def test_reading_a_half_translated_report_is_refused() -> None:
    frame = pd.DataFrame(columns=["numero_coluna_modelo", "situacao", "reviewer_status"])
    with pytest.raises(
        ValueError, match=r"Esquema misto.*Renomeie 'reviewer_status' para 'parecer_revisor'"
    ):
        rs.to_internal_frame(frame, "mappings")
    with pytest.raises(
        ValueError, match=r"Esquema misto.*Renomeie 'parecer_revisor' para 'reviewer_status'"
    ):
        rs.internal_header_index(("model_column_number", "parecer_revisor"), "mappings")


def test_find_sheet_refuses_the_other_schemas_sheet() -> None:
    assert rs.find_sheet(["Mapeamentos", "Metadados da Execução"], "run_metadata", rs.PT_BR) == (
        "Metadados da Execução"
    )
    assert rs.find_sheet(["Mappings", "Run Metadata"], "run_metadata", rs.LEGACY_EN) == "Run Metadata"
    assert rs.find_sheet(["Mapeamentos"], "run_metadata", rs.PT_BR) is None
    with pytest.raises(ValueError, match="Esquema misto"):
        rs.find_sheet(["Mapeamentos", "Run Metadata"], "run_metadata", rs.PT_BR)
    with pytest.raises(ValueError, match="Esquema misto"):
        rs.find_sheet(["Mappings", "Metadados da Execução"], "run_metadata", rs.LEGACY_EN)
    with pytest.raises(ValueError, match="ao mesmo tempo"):
        rs.find_sheet(["Run Metadata", "Metadados da Execução"], "run_metadata", rs.PT_BR)


@pytest.mark.parametrize(
    "written",
    [
        "metadados da execução",
        "METADADOS DA EXECUÇÃO",
        "Metadados da Execucao",
        # "ç" e "ã" decompostos (letra + sinal combinante), como alguns
        # programas gravam.
        "Metadados da Execução",
        " Metadados da Execução ",
    ],
)
def test_find_sheet_ignores_case_and_accents_and_returns_the_name_as_written(written) -> None:
    # Sem isso a aba não seria encontrada e a verificação de hash seria
    # pulada em silêncio.
    assert rs.find_sheet(["Mapeamentos", written], "run_metadata", rs.PT_BR) == written
    with pytest.raises(ValueError, match="Esquema misto"):
        rs.find_sheet(["Mappings", written], "run_metadata", rs.LEGACY_EN)


def test_find_sheet_refuses_a_legacy_sheet_in_another_case() -> None:
    # O Excel não distingue caixa nos nomes de aba: "run metadata" é a aba legada.
    with pytest.raises(ValueError, match="Esquema misto"):
        rs.find_sheet(["Mapeamentos", "run metadata"], "run_metadata", rs.PT_BR)


def test_find_sheet_refuses_two_sheets_that_differ_only_by_accents() -> None:
    with pytest.raises(ValueError, match="mais de uma aba"):
        rs.find_sheet(
            ["Mapeamentos", "Metadados da Execução", "Metadados da Execucao"],
            "run_metadata",
            rs.PT_BR,
        )


def test_find_sheet_without_a_schema_accepts_either_name() -> None:
    # Para localizar a aba de mapeamentos antes de ler o cabeçalho.
    assert rs.find_sheet(["Mapeamentos", "Metadados da Execução"], "mappings", None) == "Mapeamentos"
    assert rs.find_sheet(["Mappings", "Run Metadata"], "mappings", None) == "Mappings"
    assert rs.find_sheet(["Planilha1"], "mappings", None) is None
    with pytest.raises(ValueError, match="ao mesmo tempo"):
        rs.find_sheet(["Mappings", "Mapeamentos"], "mappings", None)


def test_sheet_and_header_names_in_each_schema() -> None:
    assert rs.sheet_name("mappings") == "Mapeamentos"
    assert rs.sheet_name("queue") == "Fila de Revisão"
    assert rs.sheet_name("unmatched", rs.LEGACY_EN) == "Unmatched Client Columns"
    assert rs.header_label("mappings", "reviewer_status", rs.PT_BR) == "parecer_revisor"
    assert rs.header_label("mappings", "reviewer_status", rs.LEGACY_EN) == "reviewer_status"
    assert rs.header_label("banner", "llm_final_confidence", rs.PT_BR) == "confianca_final_ia"
    assert rs.header_label("mapped_data", "__mapping_info__", rs.PT_BR) == "__info_mapeamento__"
    assert rs.header_label("mappings", "campo_novo", rs.PT_BR) == "campo_novo"
    assert rs.metadata_key_label("name_margin") == "margem_minima_nome"
    assert rs.metadata_key_label("name_margin", rs.LEGACY_EN) == "name_margin"
    with pytest.raises(ValueError, match="Esquema desconhecido"):
        rs.header_label("mappings", "status", "pt")


def test_internal_header_index_reads_openpyxl_rows_in_both_schemas() -> None:
    index, schema = rs.internal_header_index(
        ("numero_coluna_modelo", "situacao", None, "prova_aritmetica"), "mappings"
    )
    assert schema == rs.PT_BR
    assert index == {"model_column_number": 0, "status": 1, "arithmetic_proof": 3}

    index, schema = rs.internal_header_index(("model_column_number", "status"), "mappings")
    assert schema == rs.LEGACY_EN
    assert index == {"model_column_number": 0, "status": 1}


# ---------------------------------------------------------------------------
# Metadados da execução
# ---------------------------------------------------------------------------


def test_metadata_is_written_in_portuguese_and_read_back_as_internal() -> None:
    metadata = {
        "mapping_schema_version": 3,
        "embedding_model": "offline",
        "llm_model": "gpt-4o-mini",
        "name_margin": 0.05,
        "reference_manifest_sha256": "disabled",
        "client_sha256": "abc123",
    }

    frame = rs.to_external_metadata(metadata)

    assert list(frame.columns) == ["chave", "valor"]
    assert frame.to_dict("records") == [
        {"chave": "versao_esquema_mapeamento", "valor": 3},
        {"chave": "modelo_embedding", "valor": "offline"},
        {"chave": "modelo_ia", "valor": "gpt-4o-mini"},
        {"chave": "margem_minima_nome", "valor": 0.05},
        {"chave": "sha256_manifesto_referencia", "valor": "desativado"},
        {"chave": "sha256_cliente", "valor": "abc123"},
    ]
    assert rs.internal_metadata(frame, rs.PT_BR) == {
        "mapping_schema_version": "3",
        "embedding_model": "offline",
        "llm_model": "gpt-4o-mini",
        "name_margin": "0.05",
        "reference_manifest_sha256": "disabled",
        "client_sha256": "abc123",
    }


def test_legacy_metadata_is_read_like_the_apply_reads_it_today() -> None:
    frame = pd.DataFrame(
        {"key": ["client_sha256", " model_sha256 ", None, "embedding_model"],
         "value": [" abc ", "def", "ignorado", "disabled"]}
    )

    assert rs.internal_metadata(frame, rs.LEGACY_EN) == {
        "client_sha256": "abc",
        "model_sha256": "def",
        "embedding_model": "disabled",
    }


def test_metadata_in_the_wrong_schema_or_with_repeated_keys_is_refused() -> None:
    with pytest.raises(ValueError, match="Esquema misto"):
        rs.internal_metadata(pd.DataFrame({"key": ["a"], "value": ["b"]}), rs.PT_BR)
    with pytest.raises(ValueError, match="chaves repetidas"):
        rs.internal_metadata(
            pd.DataFrame({"chave": ["sha256_cliente", "client_sha256"], "valor": ["a", "b"]}),
            rs.PT_BR,
        )
    with pytest.raises(ValueError, match="precisa ter as colunas"):
        rs.internal_metadata(pd.DataFrame({"chave": ["a"]}), rs.PT_BR)
    # Como o apply sempre fez: repetição na coluna crua, inclusive de chaves em branco.
    for keys in (["client_sha256", "client_sha256"], [None, None]):
        with pytest.raises(ValueError, match="chaves repetidas"):
            rs.internal_metadata(pd.DataFrame({"key": keys, "value": ["a", "b"]}), rs.LEGACY_EN)


# ---------------------------------------------------------------------------
# Parecer do revisor, nome de exibição e etapas
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("typed", "internal"),
    [
        ("correct", "correct"),
        ("Correto", "correct"),
        ("  CORRETO ", "correct"),
        ("Corrêto", "correct"),
        ("incorrect", "incorrect"),
        ("INCORRETO", "incorrect"),
        ("no_mapping_exists", "no_mapping_exists"),
        ("Inexistente", "no_mapping_exists"),
        ("", ""),
        ("   ", ""),
        (None, ""),
        (math.nan, ""),
        # Desconhecidos voltam crus, para o apply recusar mostrando o que foi digitado.
        ("talvez", "talvez"),
        ("sem_correspondencia", "sem_correspondencia"),
        (" Talvez ", " Talvez "),
        (1, 1),
    ],
)
def test_reviewer_status_is_normalised_in_both_languages(typed, internal) -> None:
    assert rs.normalize_reviewer_status(typed) == internal


@pytest.mark.parametrize(
    ("internal", "external"),
    [
        ("Valor ICMS (occurrence 2)", "Valor ICMS (ocorrência 2)"),
        ("Unnamed column 7", "Coluna sem nome 7"),
        ("Unnamed column 7 (occurrence 2)", "Coluna sem nome 7 (ocorrência 2)"),
        ("Valor ICMS", "Valor ICMS"),
        ("Nota (occurrence 2) final", "Nota (occurrence 2) final"),
        ("Linha\nquebrada (occurrence 3)", "Linha\nquebrada (ocorrência 3)"),
        ("", ""),
        (None, None),
        (6, 6),
    ],
)
def test_display_names_translate_in_both_directions(internal, external) -> None:
    assert rs.display_name_to_external(internal) == external
    assert rs.display_name_to_external(external) == external
    assert rs.display_name_to_internal(external) == internal
    assert rs.display_name_to_internal(internal) == internal


@pytest.mark.parametrize(
    ("typed", "internal"),
    [
        ("Valor ICMS (ocorrencia 2)", "Valor ICMS (occurrence 2)"),
        ("Valor ICMS (Ocorrência 2)", "Valor ICMS (occurrence 2)"),
        ("Valor ICMS (ocorrência 2)", "Valor ICMS (occurrence 2)"),
        ("coluna sem nome 3", "Unnamed column 3"),
    ],
)
def test_display_names_accept_what_a_reviewer_types(typed, internal) -> None:
    assert rs.display_name_to_internal(typed) == internal


@pytest.mark.parametrize(
    ("internal", "external"),
    [
        ("lexical", "lexico"),
        ("resolver + semantic + profile + identifier", "resolvedor_ia + semantico + perfil + identificador"),
        ("lexical + catalog (agrees) + unrecorded", "lexico + catalogo (concorda) + nao_registrado"),
        ("unrecorded", "nao_registrado"),
        ("catalog (agrees)", "catalogo (concorda)"),
        ("family_conflict", "conflito_tributario"),
        ("underspecified + exclusivity", "cabecalho_generico + coluna_ja_consumida"),
        ("passo_desconhecido", "passo_desconhecido"),
        ("", ""),
        (None, None),
    ],
)
def test_step_cells_translate_in_both_directions(internal, external) -> None:
    assert rs.steps_to_external(internal) == external
    assert rs.steps_to_external(external) == external
    assert rs.steps_to_internal(external) == internal
    assert rs.steps_to_internal(internal) == internal


def test_step_cells_accept_mixed_languages_on_read() -> None:
    assert rs.steps_to_internal("lexico + catalog (concorda)") == "lexical + catalog (agrees)"
    assert rs.steps_to_internal("semantic + aritmetica") == "semantic + arithmetic"


def test_steps_read_back_are_understood_by_mapping_steps() -> None:
    internal = rs.steps_to_internal("lexico + catalogo + aritmetica")
    assert mapping_steps.deciding_step(internal) == "arithmetic"
    assert mapping_steps.format_steps(internal) == "lexical + catalog + arithmetic"


# ---------------------------------------------------------------------------
# Prova aritmética e números
# ---------------------------------------------------------------------------


def test_a_proof_is_decisive_unless_empty_or_marked_in_either_language() -> None:
    english = "valor = base x aliquota holds for pis on 70 rows at fit 0.9857"
    assert rs.proof_is_decisive(english)
    assert not rs.proof_is_decisive(
        english + " - but this client has 2 disjoint pis blocks that each satisfy it, "
        "so this does not single out this column"
    )
    portuguese = "valor = base x alíquota se confirma para PIS em 70 linhas, com aderência 0,9857"
    assert rs.proof_is_decisive(portuguese)
    assert not rs.proof_is_decisive(
        portuguese + f" - mas este cliente tem 2 blocos disjuntos de PIS; ela {rs.PROOF_NOT_DECISIVE_PT}"
    )
    assert not rs.proof_is_decisive("")
    assert not rs.proof_is_decisive(None)
    assert not rs.proof_is_decisive(math.nan)
    assert rs.PROOF_NOT_DECISIVE_EN == "does not single out"
    assert rs.PROOF_NOT_DECISIVE_PT == "não distingue esta coluna"


def test_the_proof_check_matches_the_one_the_mapper_uses_today() -> None:
    for proof in (None, "", "valor = base x aliquota holds", "x does not single out y"):
        expected = bool(proof) and "does not single out" not in proof
        assert rs.proof_is_decisive(proof) == expected


@pytest.mark.parametrize(
    ("value", "places", "text"),
    [
        (0.882, 1, "88,2%"),
        (1.0, 1, "100,0%"),
        (0.6667, 1, "66,7%"),
        (0.0, 1, "0,0%"),
        (0.12345, 2, "12,35%"),
    ],
)
def test_percentages_use_a_decimal_comma(value, places, text) -> None:
    assert rs.format_percent_br(value, places) == text


def test_percentages_keep_the_digits_of_the_english_format() -> None:
    for value in (0.0004, 0.0005, 0.1235, 0.66666, 0.8819, 0.9995, 0.99999):
        assert rs.format_percent_br(value) == f"{value:.1%}".replace(".", ",")


@pytest.mark.parametrize(
    ("value", "places", "text"),
    [
        (0.9996, None, "0,9996"),
        (0.75, None, "0,75"),
        (1.0, None, "1,0"),
        (5501, None, "5501"),
        (0.12345, 2, "0,12"),
        (-0.5, 1, "-0,5"),
        (12.0, 0, "12"),
    ],
)
def test_decimals_use_a_decimal_comma(value, places, text) -> None:
    assert rs.format_decimal_br(value, places) == text


# ---------------------------------------------------------------------------
# Independência do módulo
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("module", ["report_schema.py", "cli_ptbr.py"])
def test_the_boundary_modules_import_nothing_from_the_mapper(module) -> None:
    tree = ast.parse((ROOT / module).read_text(encoding="utf-8"))
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module.split(".")[0])
    assert not imported & {"invoice_mapper", "apply_invoice_mapping", "mapeval", "mapping_steps"}


# ---------------------------------------------------------------------------
# CLI em português
# ---------------------------------------------------------------------------

ENGLISH_CHROME = (
    "usage:",
    "positional arguments",
    "options:",
    "show this help",
    "show program's version",
    "error:",
    "the following arguments",
    "unrecognized arguments",
    "expected one argument",
    "invalid choice",
    "invalid int value",
    "invalid float value",
    "tipo int",
    "tipo float",
    "ignored explicit",
    "argument --",
)


def _sheet(value: str) -> int:
    """Conversão própria, como as que os CLIs passam em ``type=``."""
    return int(value)


def _parser() -> PortugueseArgumentParser:
    parser = PortugueseArgumentParser(prog="ferramenta", description="Mapeia as colunas.")
    parser.add_argument("client", help="planilha exportada pelo cliente")
    parser.add_argument("--output", help="relatório de saída")
    parser.add_argument("--candidates", type=int, default=10, help="quantidade de candidatos")
    parser.add_argument("--final-threshold", type=float, default=0.9)
    parser.add_argument("--sheet", type=_sheet)
    parser.add_argument("--mode", choices=["offline", "online"])
    parser.add_argument("--offline", action="store_true")
    parser.add_argument("--version", action="version", version="ferramenta 1.0")
    extra = parser.add_argument_group("outras opções")
    extra.add_argument("--build", action="version", version="build 7")
    return parser


def test_the_help_is_in_portuguese(capsys) -> None:
    with pytest.raises(SystemExit) as stop:
        _parser().parse_args(["--help"])

    assert stop.value.code == 0
    out = capsys.readouterr().out
    assert out.startswith("uso: ferramenta [-h]")
    assert "\nargumentos posicionais:\n" in out
    assert "\nopções:\n" in out
    assert "-h, --help" in out
    assert "mostra esta mensagem de ajuda e sai" in out
    assert "planilha exportada pelo cliente" in out
    # A ajuda padrão de version, no grupo principal e num grupo próprio.
    assert out.count("mostra a versão do programa e sai") == 2
    for english in ENGLISH_CHROME:
        assert english not in out, english


@pytest.mark.parametrize(
    ("argv", "message"),
    [
        ([], "os seguintes argumentos são obrigatórios: client"),
        (["c.xlsx", "--mode", "nuvem"], "argumento --mode: valor não permitido: 'nuvem'"),
        (["c.xlsx", "--bogus"], "argumentos não reconhecidos: --bogus"),
        (["c.xlsx", "--output"], "argumento --output: esperava um argumento"),
        (
            ["c.xlsx", "--candidates", "dez"],
            "argumento --candidates: valor inválido 'dez': esperava um número inteiro",
        ),
        (
            ["c.xlsx", "--final-threshold", "0,9"],
            "argumento --final-threshold: valor inválido '0,9': esperava um número, "
            "com ponto como separador decimal (por exemplo, 0.9)",
        ),
        (["c.xlsx", "--sheet", "aba"], "valor inválido 'aba': esperava um valor do tipo _sheet"),
        (["c.xlsx", "--offline=sim"], "argumento --offline: esta opção não aceita valor: 'sim'"),
    ],
)
def test_common_errors_are_in_portuguese(capsys, argv, message) -> None:
    with pytest.raises(SystemExit) as stop:
        _parser().parse_args(argv)

    assert stop.value.code == 2
    err = capsys.readouterr().err
    assert err.startswith("uso: ferramenta")
    assert "ferramenta: erro: " in err
    assert message in err
    for english in ENGLISH_CHROME:
        assert english not in err, english


def test_other_parsers_stay_in_english(capsys) -> None:
    _parser().format_help()
    with pytest.raises(SystemExit):
        _parser().parse_args([])
    capsys.readouterr()

    plain = argparse.ArgumentParser(prog="plain")
    plain.add_argument("value")
    plain.print_usage()
    assert capsys.readouterr().out.startswith("usage: plain")
    with pytest.raises(SystemExit):
        plain.parse_args([])
    assert "error: the following arguments are required" in capsys.readouterr().err
