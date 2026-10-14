"""Testes de ``mapper/schema.py``: vocabulário dos arquivos e campos do modelo.

Os testes de introspecção verificam se os campos e valores gravados pelo código
têm correspondência nas tabelas do schema.
"""

from __future__ import annotations

import ast
import dataclasses
import json
import math
from pathlib import Path
import re

import openpyxl
import pandas as pd
import pytest

from mapper import apply, report, schema
from mapper.decide import (
    EVIDENCE_STEPS,
    GUARDS,
    STEPS_BY_ID,
    MappingResult,
    deciding_step,
    format_steps,
    rests_only_on_headers,
)
from mapper.evaluate import read_mappings
from mapper.report import (
    REVIEWER_COLUMNS,
    WHY_HELD_BELOW_THRESHOLD,
    WHY_HELD_NAME_ONLY,
    WHY_HELD_VACUOUS,
    AdjudicationRow,
    UnmatchedColumn,
    summarize_adjudication_queue,
)
from mapper.schema import target_field
from mapper.spreadsheet import ColumnRef, SheetTable

ROOT = Path(__file__).resolve().parent.parent
SNAKE_CASE = re.compile(r"[a-z0-9]+(?:_[a-z0-9]+)*")
ARCHIVED_REPORTS = sorted((ROOT / "data" / "reports").glob("mapping_*_full.xlsx"))
requires_reports = pytest.mark.skipif(
    not ARCHIVED_REPORTS, reason="relatórios arquivados ausentes"
)
REPORT_TABLES = ("mappings", "run_metadata", "unmatched", "queue")


# ---------------------------------------------------------------------------
# Invariantes do vocabulário
# ---------------------------------------------------------------------------


def test_sheet_names_are_legal_in_excel(tmp_path) -> None:
    book = openpyxl.Workbook()
    for portuguese in schema.SHEETS.values():
        assert len(portuguese) <= 31, portuguese
        assert not set(portuguese) & set(":\\/?*[]"), portuguese
        assert not portuguese.startswith("'") and not portuguese.endswith("'")
        book.create_sheet(portuguese)
    path = tmp_path / "abas.xlsx"
    book.save(path)

    reloaded = openpyxl.load_workbook(path).sheetnames
    assert set(schema.SHEETS.values()) <= set(reloaded)
    assert len(set(schema.SHEETS.values())) == len(schema.SHEETS)


def test_headers_are_snake_case_ascii() -> None:
    for table, columns in schema.COLUMNS.items():
        for portuguese in columns.values():
            if table == "mapped_data":
                assert portuguese == "__info_mapeamento__"
                continue
            assert SNAKE_CASE.fullmatch(portuguese), (table, portuguese)
    for portuguese in schema.METADATA_KEYS.values():
        assert SNAKE_CASE.fullmatch(portuguese), portuguese


def test_no_portuguese_name_equals_a_legacy_one() -> None:
    # Os nomes internos são os do formato antigo, pelos quais ele é recusado.
    legacy = (
        {field for columns in schema.COLUMNS.values() for field in columns}
        | set(schema.METADATA_KEYS)
        | set(schema.LEGACY_SHEETS.values())
    )
    portuguese = (
        {name for columns in schema.COLUMNS.values() for name in columns.values()}
        | set(schema.METADATA_KEYS.values())
        | set(schema.SHEETS.values())
    )
    assert not legacy & portuguese, sorted(legacy & portuguese)
    assert set(schema.LEGACY_SHEETS) <= set(schema.SHEETS)


def test_headers_and_keys_are_distinct_within_each_table() -> None:
    for table, columns in schema.COLUMNS.items():
        assert len(set(columns.values())) == len(columns), table
    assert len(set(schema.METADATA_KEYS.values())) == len(schema.METADATA_KEYS)


def test_enum_values_are_distinct_and_translation_is_idempotent() -> None:
    for name, values in schema.ENUMS.items():
        assert len(set(values.values())) == len(values), name
        for internal, portuguese in values.items():
            # Um valor gravado que também é chave do enum precisa ser a
            # tradução de si mesmo, senão traduzir duas vezes o mudaria.
            if portuguese in values:
                assert values[portuguese] == portuguese, (name, internal, portuguese)
        for portuguese in values.values():
            assert portuguese == "" or SNAKE_CASE.fullmatch(portuguese), (name, portuguese)


def test_every_field_kind_names_a_column_and_an_enum() -> None:
    for table, kinds in schema.FIELD_KIND.items():
        assert set(kinds) <= set(schema.COLUMNS[table]), table
        for kind in kinds.values():
            if kind.startswith("enum:"):
                assert kind.split(":", 1)[1] in schema.ENUMS, kind


def test_step_ids_are_snake_case_and_distinct() -> None:
    portuguese = list(schema.ALL_STEP_IDS.values())
    assert len(set(portuguese)) == len(portuguese)
    assert not set(portuguese) & set(schema.ALL_STEP_IDS)
    for step in portuguese:
        assert SNAKE_CASE.fullmatch(step), step


# ---------------------------------------------------------------------------
# Cobertura por introspecção
# ---------------------------------------------------------------------------


def test_every_mapping_result_field_has_a_header() -> None:
    fields = [field.name for field in dataclasses.fields(MappingResult)] + list(REVIEWER_COLUMNS)
    assert fields == list(schema.COLUMNS["mappings"])


def test_every_unmatched_column_field_has_a_header() -> None:
    fields = [field.name for field in dataclasses.fields(UnmatchedColumn)]
    assert fields == list(schema.COLUMNS["unmatched"])


def test_every_adjudication_row_field_has_a_header() -> None:
    fields = [field.name for field in dataclasses.fields(AdjudicationRow)]
    assert fields == list(schema.COLUMNS["queue"])


def test_every_banner_row_has_a_label() -> None:
    fields = [field for _, field in apply.MAPPING_METADATA]
    assert fields == list(schema.COLUMNS["banner"])
    assert set(fields) <= set(schema.COLUMNS["mappings"])


def test_every_portuguese_verdict_reads_back_as_an_internal_one() -> None:
    verdicts = schema.ENUMS["reviewer_status"]
    assert {schema.normalize_reviewer_status(pt) for pt in verdicts.values()} == set(verdicts)
    assert set(verdicts) == {"", "correct", "incorrect", "no_mapping_exists"}


def test_every_analysis_step_and_guard_has_a_translation() -> None:
    assert {step.id for step in EVIDENCE_STEPS} == set(schema.STEP_IDS)
    assert {step.id for step in GUARDS} == set(schema.GUARD_IDS)
    assert set(STEPS_BY_ID) <= set(schema.ALL_STEP_IDS)


def test_every_status_and_method_literal_has_a_translation() -> None:
    source = "\n".join(
        path.read_text(encoding="utf-8") for path in sorted((ROOT / "mapper").glob("*.py"))
    )
    statuses = set(re.findall(r'status="(\w+)"', source))
    methods = set(re.findall(r'method="(\w+)"', source))
    assert statuses and methods
    assert statuses <= set(schema.ENUMS["status"])
    assert methods <= set(schema.ENUMS["method"])


def test_every_queue_reason_has_a_translation() -> None:
    # Uma linha por ramo de why_held: abaixo do threshold, só nome (confiança
    # presa logo abaixo do threshold) e coluna proposta vazia.
    client = SheetTable(
        columns=[
            ColumnRef("column_001", 0, "Cheia", 1),
            ColumnRef("column_002", 1, "Vazia", 1),
        ],
        data=pd.DataFrame([[1, None], [2, None]]),
    )

    def proposal(number: int, client_number: int, final: float) -> MappingResult:
        return MappingResult(
            model_column_number=number, model_column=f"M{number}",
            client_column_number=client_number, client_column="x",
            name_confidence=final, name_runner_up_confidence=0.0, name_margin=0.0,
            data_confidence=None, final_confidence=final, method="data",
            status="review", reason="",
        )

    rows = summarize_adjudication_queue(
        client_table=client,
        results=[proposal(1, 1, 0.5), proposal(2, 1, 0.8999), proposal(3, 2, 0.95)],
        final_threshold=0.9,
        sample_size=8,
    )
    held = {row.why_held for row in rows}
    assert held == {
        WHY_HELD_BELOW_THRESHOLD,
        WHY_HELD_NAME_ONLY,
        WHY_HELD_VACUOUS,
    }


def test_every_run_metadata_key_and_sentinel_has_a_translation(tmp_path, monkeypatch) -> None:
    model_path = tmp_path / "modelo.xlsx"
    client_path = tmp_path / "cliente.xlsx"
    pd.DataFrame([["CFOP", "NCM"], ["5102", "84714900"]]).to_excel(
        model_path, header=False, index=False
    )
    pd.DataFrame([["CFOP"], ["5102"]]).to_excel(client_path, header=False, index=False)
    written: dict[str, object] = {}

    def capture(results, output_path, metadata=None, unmatched=None, queue=None):
        written.update(metadata or {})

    monkeypatch.setattr(report, "write_results", capture)
    report.main(
        [
            str(client_path),
            "--model", str(model_path),
            "--output", str(tmp_path / "relatorio.xlsx"),
            "--offline",
            "--no-reference-data",
        ]
    )

    assert written
    assert set(written) <= set(schema.METADATA_KEYS)
    for key in schema.METADATA_SENTINEL_KEYS:
        assert written[key] in schema.METADATA_SENTINELS, (key, written[key])


# ---------------------------------------------------------------------------
# Ida e volta
# ---------------------------------------------------------------------------


def _internal_mappings() -> pd.DataFrame:
    rows = [
        {
            "model_column_number": 1, "model_column": "CFOP",
            "client_column_number": 3, "client_column": "Cod. Fiscal",
            "name_confidence": 0.61, "status": "accepted", "method": "catalog",
            "reason": "Versioned official cfop codes cover 88.2% ...",
            "catalog_name": "ibge_municipality", "arithmetic_proof": None,
            "supported_by": "lexical + catalog", "decided_by": "catalog", "vetoed_by": "",
            "reviewer_status": "", "correct_client_column": "", "reviewer_note": "",
        },
        {
            "model_column_number": 2, "model_column": "Valor ICMS",
            "client_column_number": 12, "client_column": "Valor ICMS (occurrence 2)",
            "name_confidence": 1.0, "status": "review", "method": "data",
            "reason": "Prosa da IA, que passa intacta.",
            "catalog_name": None,
            "arithmetic_proof": "valor = base x aliquota holds for icms on 40 rows at fit 1.0",
            "supported_by": "resolver + semantic + arithmetic", "decided_by": "arithmetic",
            "vetoed_by": "family_conflict",
            "reviewer_status": "incorrect", "correct_client_column": "Unnamed column 4",
            "reviewer_note": "Nota do revisor, byte a byte.",
        },
        {
            "model_column_number": 3, "model_column": "Base PIS",
            "client_column_number": 4, "client_column": "Unnamed column 4",
            "name_confidence": 0.0, "status": "no_match", "method": "name_only",
            "reason": "", "catalog_name": "cst_icms_taxation", "arithmetic_proof": None,
            "supported_by": "lexical + catalog (agrees) + unrecorded", "decided_by": "unrecorded",
            "vetoed_by": "vacuity",
            "reviewer_status": "no_mapping_exists", "correct_client_column": 6,
            "reviewer_note": None,
        },
        {
            "model_column_number": 4, "model_column": "NCM",
            "client_column_number": None, "client_column": None,
            "name_confidence": 0.2, "status": "review", "method": "data_error",
            "reason": "", "catalog_name": "ibs_cbs_cst", "arithmetic_proof": None,
            "supported_by": "", "decided_by": "", "vetoed_by": "",
            "reviewer_status": "correct", "correct_client_column": None, "reviewer_note": "",
        },
    ]
    return pd.DataFrame(rows)


def _internal_unmatched() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "client_column_number": 7, "client_column": "Direito (occurrence 3)",
                "disposition": "proposed_not_accepted", "review_priority": "shortlist",
                "rows_with_value": 12, "distinct_values": 3,
                "signals": "arithmetic,catalog,cnpj_checkdigit,date",
                "signal_detail": "holds dates", "closest_model_column": "Data de Emissão",
                "closest_name_confidence": 0.5,
            },
            {
                "client_column_number": 9, "client_column": "Unnamed column 9",
                "disposition": "never_referenced", "review_priority": "low",
                "rows_with_value": 0, "distinct_values": 0, "signals": "",
                "signal_detail": "No values in this column.",
                "closest_model_column": "CFOP", "closest_name_confidence": 0.1,
            },
        ]
    )


def _internal_queue() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "rank": 1, "model_column_number": 2, "model_column": "Valor ICMS",
                "proposed_client_column_number": 12,
                "proposed_client_column": "Valor ICMS (occurrence 2)",
                "final_confidence": 0.8999, "gate_distance": 0.0001, "name_margin": 0.0,
                "rows_with_value": 40, "resolver_disagrees": "", "arithmetic_proof": "",
                "client_sample_values": "1 | 2", "why_held": "abaixo do limiar de aceite",
            },
        ]
    )


def _one_way(table: str, field: str) -> bool:
    """O parecer do revisor só é traduzido na gravação; na leitura passa como foi digitado."""
    return schema.FIELD_KIND.get(table, {}).get(field, "") == "enum:reviewer_status"


@pytest.mark.parametrize(
    ("table", "frame"),
    [
        ("mappings", _internal_mappings()),
        ("unmatched", _internal_unmatched()),
        ("queue", _internal_queue()),
    ],
)
def test_internal_to_portuguese_and_back_is_lossless(table, frame) -> None:
    original = frame.copy()

    external = schema.to_external_frame(frame, table)

    assert list(external.columns) == [schema.COLUMNS[table][name] for name in frame.columns]
    back = schema.to_internal_frame(external, table)
    assert list(back.columns) == list(frame.columns)
    for column in frame.columns:
        if _one_way(table, column):
            continue
        assert back[column].equals(frame[column]), column
    if table == "mappings":
        # O parecer volta como foi gravado; quem lê o normaliza.
        assert [schema.normalize_reviewer_status(v) for v in back["reviewer_status"]] == (
            frame["reviewer_status"].tolist()
        )
    # Nenhuma das conversões altera o DataFrame recebido.
    assert frame.equals(original)
    assert schema.to_internal_frame(external, table).equals(back)


def _blank_as_none(values: list[object]) -> list[object]:
    # O Excel não distingue texto vazio de célula vazia.
    return [None if value == "" or pd.isna(value) else value for value in values]


def test_a_report_written_in_portuguese_reads_back_through_excel(tmp_path) -> None:
    path = tmp_path / "relatorio.xlsx"
    tables = {"mappings": _internal_mappings(), "unmatched": _internal_unmatched(), "queue": _internal_queue()}
    metadata = {"mapping_schema_version": schema.SCHEMA_VERSION_PT, "reference_manifest_sha256": "disabled"}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        schema.to_external_frame(tables["mappings"], "mappings").to_excel(
            writer, sheet_name=schema.sheet_name("mappings"), index=False
        )
        schema.to_external_metadata(metadata).to_excel(
            writer, sheet_name=schema.sheet_name("run_metadata"), index=False
        )
        for table in ("unmatched", "queue"):
            schema.to_external_frame(tables[table], table).to_excel(
                writer, sheet_name=schema.sheet_name(table), index=False
            )

    book = pd.read_excel(path, sheet_name=None, dtype=object)
    names = list(book)
    assert names == ["Mapeamentos", "Metadados da Execução", "Colunas do Cliente sem Uso", "Fila de Revisão"]
    for table, frame in tables.items():
        back = schema.to_internal_frame(book[schema.find_sheet(names, table)], table)
        for column in frame.columns:
            if _one_way(table, column):
                continue
            assert _blank_as_none(back[column].tolist()) == _blank_as_none(frame[column].tolist()), column
    read_metadata = schema.internal_metadata(book[schema.find_sheet(names, "run_metadata")])
    assert read_metadata == {"mapping_schema_version": "3", "reference_manifest_sha256": "disabled"}


def test_the_external_mappings_frame_is_in_portuguese() -> None:
    external = schema.to_external_frame(_internal_mappings(), "mappings").to_dict("records")

    assert [row["situacao"] for row in external] == ["aceito", "revisao", "sem_correspondencia", "revisao"]
    assert [row["metodo"] for row in external] == ["catalogo", "dados", "somente_nome", "erro_dados"]
    assert external[0]["nome_catalogo"] == "municipio_ibge"
    assert external[2]["nome_catalogo"] == "cst_icms_tributacao"
    assert external[3]["nome_catalogo"] == "cst_ibs_cbs"
    assert external[1]["coluna_cliente"] == "Valor ICMS (ocorrência 2)"
    assert external[1]["coluna_cliente_correta"] == "Coluna sem nome 4"
    assert external[2]["coluna_cliente"] == "Coluna sem nome 4"
    assert external[2]["coluna_cliente_correta"] == 6
    assert external[1]["apoiado_por"] == "resolvedor_ia + semantico + aritmetica"
    assert external[2]["apoiado_por"] == "lexico + catalogo (concorda) + nao_registrado"
    assert external[2]["decidido_por"] == "nao_registrado"
    assert [row["vetado_por"] for row in external] == ["", "conflito_tributario", "coluna_vazia", ""]
    assert [row["parecer_revisor"] for row in external] == ["", "incorreto", "inexistente", "correto"]
    # Texto livre e dados passam intactos.
    assert external[1]["justificativa"] == "Prosa da IA, que passa intacta."
    assert external[1]["prova_aritmetica"].startswith("valor = base x aliquota holds")
    assert external[1]["observacao_revisor"] == "Nota do revisor, byte a byte."
    assert external[1]["coluna_modelo"] == "Valor ICMS"


def test_the_external_side_tables_are_in_portuguese() -> None:
    unmatched = schema.to_external_frame(_internal_unmatched(), "unmatched").to_dict("records")
    assert unmatched[0]["coluna_cliente"] == "Direito (ocorrência 3)"
    assert unmatched[0]["disposicao"] == "proposta_nao_aceita"
    assert unmatched[0]["prioridade_revisao"] == "prioritaria"
    assert unmatched[0]["sinais"] == "aritmetica,catalogo,dv_cnpj,data"
    assert unmatched[1]["coluna_cliente"] == "Coluna sem nome 9"
    assert unmatched[1]["disposicao"] == "nunca_referenciada"
    assert unmatched[1]["prioridade_revisao"] == "baixa"
    assert unmatched[1]["sinais"] == ""

    queue = schema.to_external_frame(_internal_queue(), "queue").to_dict("records")
    assert queue[0]["coluna_cliente_proposta"] == "Valor ICMS (ocorrência 2)"
    assert queue[0]["motivo_retencao"] == "abaixo do limiar de aceite"
    assert queue[0]["amostra_valores_cliente"] == "1 | 2"


def test_a_hand_written_portuguese_report_reads_as_internal() -> None:
    frame = pd.DataFrame(
        [
            {
                "numero_coluna_modelo": 1, "coluna_modelo": "CFOP",
                "numero_coluna_cliente": 2, "coluna_cliente": "CFOP (ocorrencia 2)",
                "situacao": "aceito", "metodo": "nome", "nome_catalogo": "municipio_ibge",
                "apoiado_por": "lexico + catalogo (concorda)", "decidido_por": "lexico",
                "vetado_por": "", "parecer_revisor": "Incorreto",
                "coluna_cliente_correta": "coluna sem nome 5",
                "colunaExtra": "fica como está",
            }
        ]
    )

    row = schema.to_internal_frame(frame, "mappings").to_dict("records")[0]

    assert row["model_column_number"] == 1
    assert row["client_column"] == "CFOP (occurrence 2)"
    assert row["status"] == "accepted"
    assert row["method"] == "name"
    assert row["catalog_name"] == "ibge_municipality"
    assert row["supported_by"] == "lexical + catalog (agrees)"
    assert row["decided_by"] == "lexical"
    # O parecer passa como foi digitado; quem lê o valida e normaliza.
    assert row["reviewer_status"] == "Incorreto"
    assert schema.normalize_reviewer_status(row["reviewer_status"]) == "incorrect"
    assert row["correct_client_column"] == "Unnamed column 5"
    assert row["colunaExtra"] == "fica como está"


def test_conversions_keep_every_cell_exactly_as_it_was() -> None:
    # Como em sap_procosa lido com dtype=object: os números de coluna
    # digitados pelo revisor são inteiros numa coluna object, ao lado de NaN.
    stored = pd.DataFrame(
        {
            "numero_coluna_modelo": [1, 2, 3],
            "coluna_cliente": [None, "Valor (ocorrência 2)", "CFOP"],
            "parecer_revisor": [math.nan, "incorreto", "incorreto"],
            "coluna_cliente_correta": pd.Series([math.nan, 6, 32], dtype=object),
            "observacao_revisor": pd.Series([None, None, None], dtype=object),
        }
    )

    internal = schema.to_internal_frame(stored, "mappings")
    external = schema.to_external_frame(internal, "mappings")
    back = schema.to_internal_frame(external, "mappings")

    for frame in (internal, back):
        assert frame["correct_client_column"].dtype == object
        assert [type(value) for value in frame["correct_client_column"][1:]] == [int, int]
        assert frame["correct_client_column"].tolist()[1:] == [6, 32]
        assert frame["client_column"][0] is None
        assert frame["client_column"][1] == "Valor (occurrence 2)"
        assert frame["reviewer_note"].tolist() == [None, None, None]
    assert external["coluna_cliente_correta"].tolist()[1:] == [6, 32]
    assert external["coluna_cliente"].tolist()[1] == "Valor (ocorrência 2)"


def test_conversions_work_with_repeated_column_names() -> None:
    frame = pd.DataFrame([["accepted", "review"]], columns=["status", "status"])
    frame.insert(0, "model_column_number", [1])

    external = schema.to_external_frame(frame, "mappings")

    assert list(external.columns) == ["numero_coluna_modelo", "situacao", "situacao"]
    assert external.iloc[0].tolist() == [1, "aceito", "revisao"]
    back = schema.to_internal_frame(external, "mappings")
    assert back.iloc[0].tolist() == [1, "accepted", "review"]


def test_enum_values_are_converted_only_when_exact() -> None:
    # Só converte valores que correspondem exatamente ao vocabulário do schema.
    # Preserva os demais para validação posterior, inclusive o parecer do revisor.
    assert schema.internal_value("mappings", "status", "aceito") == "accepted"
    assert schema.internal_value("mappings", "status", "Aceito") == "Aceito"
    assert schema.internal_value("mappings", "status", "accepted") == "accepted"
    assert schema.internal_value("mappings", "decided_by", "lexico") == "lexical"
    assert schema.internal_value("mappings", "model_column", "Situação") == "Situação"
    assert schema.internal_value("banner", "status", "aceito") == "accepted"
    assert schema.internal_value("mappings", "reviewer_status", " Correto ") == " Correto "
    assert schema.external_value("banner", "decided_by", "catalog") == "catalogo"
    assert schema.external_value("banner", "final_confidence", 0.93) == 0.93


@requires_reports
@pytest.mark.parametrize("report", ARCHIVED_REPORTS, ids=lambda path: path.stem)
def test_archived_reports_have_a_translation_for_every_value(report) -> None:
    book = pd.read_excel(report, sheet_name=None, dtype=object)
    names = list(book)
    for table in REPORT_TABLES:
        sheet = schema.find_sheet(names, table)
        if sheet is None:
            continue
        internal = schema.to_internal_frame(book[sheet], table)

        for field, kind in schema.FIELD_KIND.get(table, {}).items():
            if field not in internal.columns:
                continue
            values = [v for v in internal[field].tolist() if isinstance(v, str) and v]
            if kind.startswith("enum:"):
                enum = schema.ENUMS[kind.split(":", 1)[1]]
                assert set(values) <= set(enum) | set(enum.values()), (field, set(values) - set(enum))
            elif kind == "steps":
                tokens = {
                    token.removesuffix(schema.AGREES_SUFFIX_EN)
                    for value in values
                    for token in value.split(schema.STEPS_SEPARATOR)
                }
                assert tokens <= set(schema.ALL_STEP_IDS), tokens - set(schema.ALL_STEP_IDS)
            elif kind == "signals":
                tokens = {token for value in values for token in value.split(schema.SIGNALS_SEPARATOR)}
                assert tokens <= set(schema.ENUMS["signals"]), tokens

        back = schema.to_internal_frame(schema.to_external_frame(internal, table), table)
        for column in internal.columns:
            assert back[column].equals(internal[column]), (sheet, column)

    metadata_sheet = schema.find_sheet(names, "run_metadata")
    metadata = schema.internal_metadata(book[metadata_sheet])
    assert set(metadata) <= set(schema.METADATA_KEYS)
    exported = schema.to_external_metadata(metadata)
    assert schema.internal_metadata(exported) == metadata


# ---------------------------------------------------------------------------
# Formato antigo: recusa, abas e cabeçalhos
# ---------------------------------------------------------------------------


def _old_english_report(path: Path) -> Path:
    """Cria um relatório mínimo no formato anterior à tradução.

    A aba Mappings usa cabeçalhos e valores internos em inglês; Run Metadata
    usa as colunas key e value.
    """
    mapping = _internal_mappings().iloc[[0]].copy()
    mapping["client_column_number"] = 1
    mapping["client_column"] = "CFOP source"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        mapping.to_excel(writer, sheet_name="Mappings", index=False)
        pd.DataFrame([["mapping_schema_version", 2]], columns=["key", "value"]).to_excel(
            writer, sheet_name="Run Metadata", index=False
        )
    return path


def test_an_old_english_report_is_refused_with_the_conversion_hint(tmp_path) -> None:
    # Pelos dois leitores: o apply e o mapeval.
    report = _old_english_report(tmp_path / "mapping_antigo.xlsx")
    model_path = tmp_path / "model.xlsx"
    client_path = tmp_path / "client.xlsx"
    pd.DataFrame([["CFOP"], ["5102"]]).to_excel(model_path, header=False, index=False)
    pd.DataFrame([["CFOP source"], ["5102"]]).to_excel(client_path, header=False, index=False)
    hint = (
        r"formato antigo, em inglês \(versão 2 do esquema\).*esquema em português \(versão 3\)"
        r".*migrate_reports\.py do commit 5fe5092: git worktree add /tmp/cbs-5fe5092 5fe5092"
    )

    with pytest.raises(SystemExit) as stop:
        apply.main(
            [str(client_path), str(report), "--model", str(model_path),
             "--output", str(tmp_path / "saida.xlsx")]
        )
    assert re.search(hint, str(stop.value)), stop.value
    assert "'model_column_number'" in str(stop.value)
    assert not (tmp_path / "saida.xlsx").exists()

    with pytest.raises(ValueError, match=hint) as raised:
        read_mappings(report)
    assert "'Mappings'" in str(raised.value)


def test_the_header_must_be_in_portuguese() -> None:
    schema.check_header(["numero_coluna_modelo", "situacao", None])
    schema.check_header(["chave", "valor"], "run_metadata")
    with pytest.raises(
        ValueError, match=r"formato antigo.*tem a coluna 'model_column_number'.*commit 5fe5092"
    ):
        schema.check_header(("model_column_number", "status", None))


@pytest.mark.parametrize(
    ("header", "message"),
    [
        (
            ["numero_coluna_modelo", "model_column_number"],
            "Renomeie 'model_column_number' para 'numero_coluna_modelo'",
        ),
        (["numero_coluna_modelo", "status"], "Renomeie 'status' para 'situacao'"),
        # Sem a coluna em português, a aba de mapeamentos está no formato antigo.
        (["model_column_number", "parecer_revisor"], "migrate_reports.py do commit 5fe5092"),
    ],
)
def test_a_header_with_old_names_is_refused(header, message) -> None:
    with pytest.raises(ValueError, match=re.escape(message)):
        schema.check_header(header)


def test_a_header_without_the_required_column_is_refused() -> None:
    for header in (["coluna_modelo", "situacao"], []):
        with pytest.raises(ValueError, match="não tem a coluna 'numero_coluna_modelo'"):
            schema.check_header(header)


def test_reading_a_half_translated_report_is_refused() -> None:
    # Lido pela metade, o relatório ignoraria os pareceres do revisor sem aviso.
    frame = pd.DataFrame(columns=["numero_coluna_modelo", "situacao", "reviewer_status"])
    with pytest.raises(ValueError, match="Renomeie 'reviewer_status' para 'parecer_revisor'"):
        schema.to_internal_frame(frame, "mappings")
    with pytest.raises(ValueError, match="Renomeie 'key' para 'chave', 'value' para 'valor'"):
        schema.to_internal_frame(pd.DataFrame(columns=["key", "value"]), "run_metadata")
    with pytest.raises(ValueError, match="Renomeie 'status' para 'situacao'"):
        schema.internal_header_index(("numero_coluna_modelo", "status"), "mappings")


def test_find_sheet_refuses_a_sheet_with_the_old_name() -> None:
    assert schema.find_sheet(["Mapeamentos", "Metadados da Execução"], "run_metadata") == (
        "Metadados da Execução"
    )
    assert schema.find_sheet(["Mapeamentos"], "run_metadata") is None
    # Ignorada, ela faria a verificação de hash ser pulada em silêncio.
    for names in (["Mapeamentos", "Run Metadata"], ["Run Metadata", "Metadados da Execução"]):
        with pytest.raises(
            ValueError, match="Renomeie 'Run Metadata' para 'Metadados da Execução'"
        ):
            schema.find_sheet(names, "run_metadata")
    # Sem a aba de mapeamentos em português, o relatório está no formato antigo.
    with pytest.raises(ValueError, match="tem a aba 'Mappings'.*commit 5fe5092"):
        schema.find_sheet(["Mappings", "Run Metadata"], "mappings")
    with pytest.raises(ValueError, match="Renomeie 'Mappings' para 'Mapeamentos'"):
        schema.find_sheet(["Mappings", "Mapeamentos"], "mappings")


@pytest.mark.parametrize(
    "written",
    [
        "metadados da execução",
        "METADADOS DA EXECUÇÃO",
        "Metadados da Execucao",
        # "ç" e "ã" decompostos (letra + sinal combinante), como alguns
        # programas gravam.
        "Metadados da Execuc\u0327a\u0303o",
        " Metadados da Execução ",
    ],
)
def test_find_sheet_ignores_case_and_accents_and_returns_the_name_as_written(written) -> None:
    # Senão a verificação de hash seria pulada em silêncio.
    assert schema.find_sheet(["Mapeamentos", written], "run_metadata") == written


def test_find_sheet_refuses_an_old_sheet_name_in_another_case() -> None:
    # O Excel não distingue caixa nos nomes de aba: "run metadata" é a aba antiga.
    with pytest.raises(ValueError, match="formato antigo"):
        schema.find_sheet(["Mapeamentos", "run metadata"], "run_metadata")


def test_find_sheet_refuses_two_sheets_that_differ_only_by_accents() -> None:
    with pytest.raises(ValueError, match="mais de uma aba"):
        schema.find_sheet(
            ["Mapeamentos", "Metadados da Execução", "Metadados da Execucao"],
            "run_metadata",
        )


def test_find_sheet_finds_the_mappings_sheet() -> None:
    assert schema.find_sheet(["Mapeamentos", "Metadados da Execução"], "mappings") == "Mapeamentos"
    assert schema.find_sheet(["Planilha1"], "mappings") is None
    # As abas da planilha mapeada não têm nome antigo a recusar.
    assert schema.find_sheet(["Dados Mapeados", "Cobertura"], "coverage") == "Cobertura"


def test_sheet_and_header_names() -> None:
    assert schema.sheet_name("mappings") == "Mapeamentos"
    assert schema.sheet_name("queue") == "Fila de Revisão"
    assert schema.header_label("mappings", "reviewer_status") == "parecer_revisor"
    assert schema.header_label("banner", "llm_final_confidence") == "confianca_final_ia"
    assert schema.header_label("mapped_data", "__mapping_info__") == "__info_mapeamento__"
    assert schema.header_label("mappings", "campo_novo") == "campo_novo"


def test_internal_header_index_reads_openpyxl_rows() -> None:
    index = schema.internal_header_index(
        ("numero_coluna_modelo", "situacao", None, "prova_aritmetica"), "mappings"
    )
    assert index == {"model_column_number": 0, "status": 1, "arithmetic_proof": 3}


# ---------------------------------------------------------------------------
# Metadados da execução
# ---------------------------------------------------------------------------


def test_metadata_is_written_in_portuguese_and_read_back_as_internal() -> None:
    metadata = {
        "mapping_schema_version": 3,
        "embedding_model": "offline",
        "llm_model": "gpt-4o-mini",
        "name_margin": 0.05,
        "reference_manifest_sha256": "disabled",
        "client_sha256": "abc123",
    }

    frame = schema.to_external_metadata(metadata)

    assert list(frame.columns) == ["chave", "valor"]
    assert frame.to_dict("records") == [
        {"chave": "versao_esquema_mapeamento", "valor": 3},
        {"chave": "modelo_embedding", "valor": "offline"},
        {"chave": "modelo_ia", "valor": "gpt-4o-mini"},
        {"chave": "margem_minima_nome", "valor": 0.05},
        {"chave": "sha256_manifesto_referencia", "valor": "desativado"},
        {"chave": "sha256_cliente", "valor": "abc123"},
    ]
    assert schema.internal_metadata(frame) == {
        "mapping_schema_version": "3",
        "embedding_model": "offline",
        "llm_model": "gpt-4o-mini",
        "name_margin": "0.05",
        "reference_manifest_sha256": "disabled",
        "client_sha256": "abc123",
    }


def test_metadata_is_cleaned_like_the_apply_always_cleaned_it() -> None:
    frame = pd.DataFrame(
        {"chave": ["sha256_cliente", " sha256_modelo ", None, "modelo_embedding"],
         "valor": [" abc ", "def", "ignorado", "desativado"]}
    )

    assert schema.internal_metadata(frame) == {
        "client_sha256": "abc",
        "model_sha256": "def",
        "embedding_model": "disabled",
    }


def test_metadata_with_old_names_or_repeated_keys_is_refused() -> None:
    with pytest.raises(ValueError, match="Renomeie 'key' para 'chave'"):
        schema.internal_metadata(pd.DataFrame({"key": ["a"], "value": ["b"]}))
    # Igual à chave interna, ela valeria no lugar da em português sem aviso.
    with pytest.raises(ValueError, match="Renomeie 'client_sha256' para 'sha256_cliente'"):
        schema.internal_metadata(
            pd.DataFrame({"chave": ["sha256_cliente", " client_sha256 "], "valor": ["a", "b"]})
        )
    with pytest.raises(ValueError, match="precisa ter as colunas"):
        schema.internal_metadata(pd.DataFrame({"chave": ["a"]}))
    # Como o apply sempre fez: repetição na coluna original, inclusive de chaves em branco.
    for keys in (["sha256_cliente", "sha256_cliente"], [None, None]):
        with pytest.raises(ValueError, match="chaves repetidas"):
            schema.internal_metadata(pd.DataFrame({"chave": keys, "valor": ["a", "b"]}))


# ---------------------------------------------------------------------------
# Parecer do revisor, nome de exibição e etapas
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("typed", "internal"),
    [
        ("correto", "correct"),
        ("Correto", "correct"),
        ("  CORRETO ", "correct"),
        ("Corrêto", "correct"),
        ("INCORRETO", "incorrect"),
        ("incorreto", "incorrect"),
        ("Incorreto", "incorrect"),
        (" INCORRETO ", "incorrect"),
        ("Incorréto", "incorrect"),
        (" INCORRÉTO ", "incorrect"),
        ("Inexistente", "no_mapping_exists"),
        ("", ""),
        ("   ", ""),
        (None, ""),
        (math.nan, ""),
        # Os pareceres em inglês do formato antigo coincidem com os valores
        # internos; só a recusa explícita impede que passem calados.
        ("correct", None),
        (" Correct ", None),
        ("incorrect", None),
        ("no_mapping_exists", None),
        # Desconhecidos: quem lê recusa mostrando o que foi digitado.
        ("talvez", None),
        ("sem_correspondencia", None),
        (" Talvez ", None),
        (1, None),
    ],
)
def test_reviewer_status_is_read_only_in_portuguese(typed, internal) -> None:
    assert schema.normalize_reviewer_status(typed) == internal


@pytest.mark.parametrize(
    ("internal", "external"),
    [
        ("Valor ICMS (occurrence 2)", "Valor ICMS (ocorrência 2)"),
        ("Unnamed column 7", "Coluna sem nome 7"),
        ("Unnamed column 7 (occurrence 2)", "Coluna sem nome 7 (ocorrência 2)"),
        ("Valor ICMS", "Valor ICMS"),
        ("Nota (occurrence 2) final", "Nota (occurrence 2) final"),
        ("Linha\nquebrada (occurrence 3)", "Linha\nquebrada (ocorrência 3)"),
        ("", ""),
        (None, None),
        (6, 6),
    ],
)
def test_display_names_translate_in_both_directions(internal, external) -> None:
    assert schema.display_name_to_external(internal) == external
    assert schema.display_name_to_external(external) == external
    assert schema.display_name_to_internal(external) == internal


@pytest.mark.parametrize(
    ("typed", "internal"),
    [
        ("Valor ICMS (ocorrencia 2)", "Valor ICMS (occurrence 2)"),
        ("Valor ICMS (Ocorrência 2)", "Valor ICMS (occurrence 2)"),
        # "ê" decomposto (letra + sinal combinante).
        ("Valor ICMS (ocorre\u0302ncia 2)", "Valor ICMS (occurrence 2)"),
        ("coluna sem nome 3", "Unnamed column 3"),
    ],
)
def test_display_names_accept_what_a_reviewer_types(typed, internal) -> None:
    assert schema.display_name_to_internal(typed) == internal


@pytest.mark.parametrize(
    ("internal", "external"),
    [
        ("lexical", "lexico"),
        ("resolver + semantic + profile + identifier", "resolvedor_ia + semantico + perfil + identificador"),
        ("lexical + catalog (agrees) + unrecorded", "lexico + catalogo (concorda) + nao_registrado"),
        ("unrecorded", "nao_registrado"),
        ("catalog (agrees)", "catalogo (concorda)"),
        ("family_conflict", "conflito_tributario"),
        ("underspecified + exclusivity", "cabecalho_generico + coluna_ja_consumida"),
        ("passo_desconhecido", "passo_desconhecido"),
        ("", ""),
        (None, None),
    ],
)
def test_step_cells_translate_in_both_directions(internal, external) -> None:
    assert schema.steps_to_external(internal) == external
    assert schema.steps_to_external(external) == external
    assert schema.steps_to_internal(external) == internal
    for field in ("supported_by", "decided_by"):
        assert schema.internal_value("mappings", field, external) == internal


def test_steps_read_back_are_understood_by_mapping_steps() -> None:
    external = "resolvedor_ia + lexico + semantico + perfil + identificador + catalogo + aritmetica"
    expected = "resolver + lexical + semantic + profile + identifier + catalog + arithmetic"
    internal = schema.steps_to_internal(external)
    assert internal == expected
    assert schema.steps_to_external(internal) == external
    assert deciding_step(internal) == "arithmetic"
    assert format_steps(internal) == expected
    assert not rests_only_on_headers(internal)


# ---------------------------------------------------------------------------
# Formatação de números
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "places", "text"),
    [
        (0.882, 1, "88,2%"),
        (1.0, 1, "100,0%"),
        (0.6667, 1, "66,7%"),
        (0.0, 1, "0,0%"),
        (0.12345, 2, "12,35%"),
    ],
)
def test_percentages_use_a_decimal_comma(value, places, text) -> None:
    assert schema.format_percent_br(value, places) == text


def test_percentages_keep_the_digits_of_the_english_format() -> None:
    for value in (0.0004, 0.0005, 0.1235, 0.66666, 0.8819, 0.9995, 0.99999):
        assert schema.format_percent_br(value) == f"{value:.1%}".replace(".", ",")


@pytest.mark.parametrize(
    ("value", "places", "text"),
    [
        (0.9996, None, "0,9996"),
        (0.75, None, "0,75"),
        (1.0, None, "1,0"),
        (5501, None, "5501"),
        (0.12345, 2, "0,12"),
        (-0.5, 1, "-0,5"),
        (12.0, 0, "12"),
    ],
)
def test_decimals_use_a_decimal_comma(value, places, text) -> None:
    assert schema.format_decimal_br(value, places) == text


# ---------------------------------------------------------------------------
# Independência do módulo
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("module", ["mapper/schema.py", "mapper/cli.py"])
def test_the_boundary_modules_import_nothing_from_the_mapper(module) -> None:
    tree = ast.parse((ROOT / module).read_text(encoding="utf-8"))
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module == "mapper":
            imported.update(f"mapper.{alias.name}" for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module)
    # As etapas do mapeamento e os comandos ficam acima destes módulos.
    assert not imported & {
        "mapper.headers", "mapper.evidence", "mapper.resolver", "mapper.decide",
        "mapper.report", "mapper.apply", "mapper.evaluate",
    }


def test_explicit_registry_covers_every_canonical_column():
    raw = json.loads(Path("data/target_fields.json").read_text())
    model = SheetTable.from_excel("data/model.xlsx")
    assert [(row[0], row[1]) for row in raw["fields"]] == [
        (c.index + 1, c.name) for c in model.columns
    ]
    assert all(target_field(c.name) is not None for c in model.columns)


def test_target_descriptions_have_distinct_explicit_identities():
    assert target_field("CFOP_descrição").number == 35
    assert target_field("item/material_descrição").number == 195
    assert target_field("Descrição") is None
    assert target_field("CFOP_descrição", number=195) is None
    assert target_field("CFOP_descrição", number=35).role == "text"
