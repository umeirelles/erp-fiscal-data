"""Testes da migração dos relatórios legados (migrate_reports) e do tradutor legacy_reports.

Quatro grupos:

* um relatório legado sintético, com pareceres do revisor, sufixo de
  ocorrência, ``unrecorded``, ``(agrees)``, prosa da IA e uma observação
  humana, migrado e comparado com o português esperado; migrar duas vezes dá
  o mesmo arquivo; a verificação embutida barra uma migração errada sem
  tocar no original; arquivos mistos, que não são planilha ou com a aba de
  mapeamentos fora do lugar são recusados sem gravar; a regeneração
  divergente recusa a migração; o texto vazio continua texto vazio;
* antidesvio: para cada modelo de frase inglês legado, a tradução da saída
  do código antigo (copiado de ``2ae4b6d:invoice_mapper.py``) é exatamente o
  que a função produtora atual grava para as mesmas entradas;
* ``@requires_workbooks``: a saída offline em inglês do código de 2ae4b6d,
  traduzida, é exatamente o texto de uma execução offline nova do código
  atual; e, migrada pelo ``migrate_reports``, é célula a célula (valor, tipo,
  formato e estilo) a execução nova, menos a data e a linha da migração;
* ``@requires_workbooks``: migrar uma cópia do relatório arquivado de
  sap_procosa preserva os 26 pareceres e o ``apply`` continua entregando 62
  colunas; o de oracle passa na prova cruzada por regeneração.

A linha de base em inglês não fica no repositório: o teste extrai o código
de 2ae4b6d com ``git archive`` numa pasta temporária e roda nele o mapeador
offline para os quatro clientes. Para usar uma linha de base já gravada,
aponte a variável ``MIGRACAO_BASELINE_EN`` para a pasta com os arquivos
``mapping_<cliente>_offline.xlsx``.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Sequence

import openpyxl
import pandas as pd
import pytest

import apply_invoice_mapping
import invoice_mapper as im
import legacy_reports
import migrate_reports
import report_schema as rs
from reference_catalog import ReferenceCatalogs

ROOT = Path(__file__).resolve().parent.parent

#: Commit que gravou os relatórios arquivados e a linha de base em inglês.
LEGACY_COMMIT = "2ae4b6d"

ALREADY_PORTUGUESE = migrate_reports.ALREADY_PORTUGUESE

EXPORTS = {
    "datasul": "data/clients/datasul_maio.xlsx",
    "oracle": "data/clients/oracle.xlsx",
    "sap_procosa": "data/clients/sap_procosa.XLSX",
    "sap_terumo": "data/clients/sap_terumo.XLSX",
}

requires_workbooks = pytest.mark.skipif(
    not all((ROOT / path).exists() for path in EXPORTS.values())
    or not (ROOT / "data/model.xlsx").exists(),
    reason="planilhas dos clientes ausentes",
)

#: Linha de base em inglês já gravada (opcional). Sem ela, o teste a gera a
#: partir do código de 2ae4b6d (ver ``baseline_offline_reports``).
_BASELINE_ENV = os.environ.get("MIGRACAO_BASELINE_EN", "")
BASELINE_EN = Path(_BASELINE_ENV) if _BASELINE_ENV else None


# ---------------------------------------------------------------------------
# Relatório legado sintético
# ---------------------------------------------------------------------------

#: Cabeçalho da aba Mappings na ordem dos relatórios arquivados (decided_by
#: antes de supported_by, diferente da ordem que o mapeador grava hoje). A
#: migração mantém essa ordem.
LEGACY_MAPPINGS_HEADER = [
    "model_column_number",
    "model_column",
    "client_column_number",
    "client_column",
    "name_confidence",
    "name_runner_up_confidence",
    "name_margin",
    "data_confidence",
    "final_confidence",
    "method",
    "status",
    "reason",
    "catalog_name",
    "catalog_coverage",
    "catalog_distinct_coverage",
    "catalog_confidence",
    "decided_by",
    "supported_by",
    "vetoed_by",
    "arithmetic_proof",
    "llm_data_confidence",
    "llm_final_confidence",
    "reviewer_status",
    "correct_client_column",
    "reviewer_note",
]

LEGACY_NOT_DECISIVE_PROOF = (
    "valor = base x aliquota holds for icms on 120 rows at fit 0.9996 - but this "
    "client has 2 disjoint icms blocks that each satisfy it, so this does not single "
    "out this column"
)
LLM_PROSE = (
    "Both columns hold 7-digit IBGE municipality codes and the samples overlap exactly."
)
#: Uma observação humana que, por acaso, é igual a uma frase do mapeador:
#: observações nunca são traduzidas.
HUMAN_NOTE_LOOKS_LIKE_TEMPLATE = "No unused client columns remain."
HUMAN_NOTE = "Wrong block; the right one is column 6 (checked against the de_para)."

MAPPING_ROWS: list[dict[str, object]] = [
    {
        "model_column_number": 1, "model_column": "CFOP", "client_column_number": 3,
        "client_column": "CFOP", "name_confidence": 1.0, "name_runner_up_confidence": 0.4,
        "name_margin": 0.6, "final_confidence": 1.0, "method": "name", "status": "accepted",
        "reason": "Accepted from header similarity with a clear lead over the runner-up.",
        "decided_by": "lexical", "supported_by": "lexical + catalog (agrees)",
        "reviewer_status": "correct", "reviewer_note": HUMAN_NOTE_LOOKS_LIKE_TEMPLATE,
    },
    {
        "model_column_number": 2, "model_column": "Valor ICMS", "client_column_number": 5,
        "client_column": "Valor ICMS (occurrence 2)", "name_confidence": 0.95,
        "name_runner_up_confidence": 0.9, "name_margin": 0.05, "data_confidence": 0.0,
        "final_confidence": 0.95, "method": "data", "status": "review",
        "reason": (
            "'Valor ICMS (occurrence 2)' has no populated rows, so it cannot be this "
            "column's source however well the headers match. Not accepted."
        ),
        "decided_by": "", "supported_by": "unrecorded", "vetoed_by": "vacuity",
        "arithmetic_proof": LEGACY_NOT_DECISIVE_PROOF,
        "reviewer_status": "incorrect", "correct_client_column": 6, "reviewer_note": HUMAN_NOTE,
    },
    {
        "model_column_number": 3, "model_column": "Municipio", "client_column_number": 7,
        "client_column": "Cod Mun", "name_confidence": 0.5, "name_runner_up_confidence": 0.4,
        "name_margin": 0.1, "data_confidence": 0.9, "final_confidence": 0.95, "method": "data",
        "status": "accepted", "reason": LLM_PROSE, "catalog_name": "ibge_municipality",
        "catalog_coverage": 0.98, "catalog_distinct_coverage": 0.97, "catalog_confidence": 0.96,
        "decided_by": "resolver", "supported_by": "resolver + semantic + catalog",
        "llm_data_confidence": 0.9, "llm_final_confidence": 0.95,
    },
    {
        "model_column_number": 4, "model_column": "Valor IPI", "name_confidence": 0.3,
        "name_runner_up_confidence": 0.2, "name_margin": 0.1, "final_confidence": 0.0,
        "method": "data", "status": "no_match",
        "reason": (
            "No candidate had enough combined header and distinctive row evidence; "
            "the client export may not contain this model field."
        ),
        "reviewer_status": "no_mapping_exists",
    },
    {
        "model_column_number": 5, "model_column": "Valor PIS", "client_column_number": 8,
        "client_column": "Valor ICMS ST", "name_confidence": 0.7,
        "name_runner_up_confidence": 0.6, "name_margin": 0.1, "data_confidence": 0.9,
        "final_confidence": 0.88, "method": "data", "status": "review",
        "reason": (
            "Explicit tax-family conflict (pis vs icms); keep this proposal for review "
            "because the row content provides distinctive evidence that the client header "
            "may be bogus."
        ),
        "decided_by": "", "supported_by": "semantic", "llm_final_confidence": 0.75,
    },
    {
        "model_column_number": 6, "model_column": "Natureza", "client_column_number": 3,
        "client_column": "CFOP", "name_confidence": 0.2, "name_runner_up_confidence": 0.1,
        "name_margin": 0.1, "data_confidence": 0.99, "final_confidence": 0.92,
        "method": "catalog", "status": "accepted",
        "reason": (
            "Versioned official cfop codes cover 88.2% of the 17 distinct codes used "
            "(66.7% of rows) and uniquely identify this candidate among the available columns."
        ),
        "catalog_name": "cfop", "decided_by": "catalog", "supported_by": "catalog",
    },
    {
        "model_column_number": 7, "model_column": "Valor COFINS", "client_column_number": 2,
        "client_column": "Valor", "name_confidence": 0.6, "name_runner_up_confidence": 0.5,
        "name_margin": 0.1, "final_confidence": 0.6, "method": "data_error", "status": "review",
        "reason": "Data comparison failed: ValueError: could not convert string to float: 'abc'",
    },
    {
        "model_column_number": 8, "model_column": "Situacao", "client_column_number": 9,
        "client_column": "Unnamed column 9", "name_confidence": 0.8999,
        "name_runner_up_confidence": 0.8, "name_margin": 0.0999, "final_confidence": 0.8999,
        "method": "name_only", "status": "review",
        "reason": (
            "Header score or runner-up margin was insufficient; data comparison was disabled."
        ),
    },
]

LEGACY_METADATA: list[tuple[str, object]] = [
    ("mapping_schema_version", 2),
    ("generated_at_utc", "2026-08-25T17:26:27.092955+00:00"),
    ("prompt_version", "2026-08-04.2"),
    ("client_file", "cliente.xlsx"),
    ("client_sha256", "0" * 64),
    ("model_file", "modelo.xlsx"),
    ("model_sha256", "1" * 64),
    ("client_sheet", "0"),
    ("model_sheet", "0"),
    ("client_header_row", 1),
    ("model_header_row", 1),
    ("embedding_model", "offline"),
    ("llm_model", "offline"),
    ("name_threshold", 0.9),
    ("name_margin", 0.05),
    ("final_threshold", 0.9),
    ("candidate_count", 10),
    ("sample_size", 8),
    ("reference_manifest_sha256", "disabled"),
]

LEGACY_UNMATCHED_HEADER = list(rs.COLUMNS["unmatched"])
UNMATCHED_ROWS = [
    [4, "Valor ICMS", "proposed_not_accepted", "shortlist", 120, 100, "arithmetic,catalog",
     "participates in a proven valor = base x aliquota trio as the value; 88.2% of its 17 "
     "distinct codes are valid cfop", "Valor ICMS", 1.0],
    [10, "Doc (occurrence 2)", "never_referenced", "low", 0, 0, None,
     "No values in this column.", "Documento", 0.5],
    [11, "CNPJ Emit", "never_referenced", "shortlist", 120, 30, "cnpj_checkdigit,date",
     "100.0% of values pass the cnpj check digit; holds dates", "CNPJ", 0.8],
]

LEGACY_QUEUE_HEADER = list(rs.COLUMNS["queue"])
QUEUE_ROWS = [
    [1, 2, "Valor ICMS", 5, "Valor ICMS (occurrence 2)", 0.95, -0.05, 0.05, 0, None,
     LEGACY_NOT_DECISIVE_PROOF, None,
     "the proposed column has no populated rows - confirming this would ship a blank column"],
    [2, 5, "Valor PIS", 8, "Valor ICMS ST", 0.88, 0.02, 0.1, 120, "resolver scored it 0.75",
     None, "1 | 2 | 3", "below the accept threshold"],
    [3, 8, "Situacao", 9, "Unnamed column 9", 0.8999, 0.0001, 0.0999, 120, None, None,
     "x | y", "name evidence only - this branch cannot auto-accept however well the headers "
     "match"],
]

#: Português esperado das frases e valores do relatório sintético.
EXPECTED_REASONS = [
    "Aceito pela semelhança do cabeçalho, com vantagem clara sobre a segunda melhor opção.",
    "'Valor ICMS (ocorrência 2)' não tem linhas preenchidas, portanto não pode ser a origem "
    "desta coluna, por mais que os cabeçalhos coincidam. Não aceito.",
    LLM_PROSE,
    "Nenhum candidato reuniu evidência suficiente de cabeçalho e de linhas distintivas; a "
    "exportação do cliente pode não conter este campo do modelo.",
    "Os cabeçalhos apresentam conflito de família de tributo (PIS vs ICMS); esta proposta "
    "fica para revisão porque o conteúdo das linhas traz evidência distintiva de que o "
    "cabeçalho do cliente pode estar errado.",
    "A tabela oficial versionada de CFOP cobre 88,2% dos 17 códigos distintos usados (66,7% "
    "das linhas) e identifica este candidato de forma única entre as colunas disponíveis.",
    "A comparação de dados falhou: ValueError: could not convert string to float: 'abc'",
    "A pontuação do cabeçalho ou a margem sobre a segunda melhor opção foi insuficiente; a "
    "comparação de dados estava desativada.",
]
EXPECTED_NOT_DECISIVE_PROOF = (
    "valor = base x alíquota se confirma para ICMS em 120 linhas, com aderência 0,9996 - mas "
    "este cliente tem 2 blocos disjuntos de ICMS que satisfazem a equação, portanto ela não "
    "distingue esta coluna"
)


def _write_legacy_report(
    path: Path,
    *,
    metadata: list[tuple[str, object]] | None = None,
    metadata_title: str = "Run Metadata",
) -> Path:
    """Grava o relatório legado sintético com o openpyxl, com os tipos de cada célula."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Mappings"
    sheet.append(LEGACY_MAPPINGS_HEADER)
    for row in MAPPING_ROWS:
        sheet.append([row.get(name) for name in LEGACY_MAPPINGS_HEADER])
    sheet["A1"].font = openpyxl.styles.Font(bold=True)
    sheet.column_dimensions["L"].width = 60
    meta = book.create_sheet(metadata_title)
    meta.append(["key", "value"])
    for key, value in metadata or LEGACY_METADATA:
        meta.append([key, value])
    unmatched = book.create_sheet("Unmatched Client Columns")
    unmatched.append(LEGACY_UNMATCHED_HEADER)
    for row in UNMATCHED_ROWS:
        unmatched.append(row)
    queue = book.create_sheet("Adjudication Queue")
    queue.append(LEGACY_QUEUE_HEADER)
    for row in QUEUE_ROWS:
        queue.append(row)
    book.save(path)
    return path


def _write_tables(folder: Path) -> tuple[Path, Path]:
    """Exportação e modelo mínimos: só servem para os hashes e a regeneração."""
    export = folder / "cliente.xlsx"
    pd.DataFrame(
        {
            "Valor": [1.0, 2.0, 3.0],
            "CFOP": ["5102", "6102", "5405"],
            "Valor ICMS": [1.0, 2.0, 3.0],
        }
    ).to_excel(export, index=False)
    model = folder / "modelo.xlsx"
    pd.DataFrame(
        columns=[row["model_column"] for row in MAPPING_ROWS]
    ).to_excel(model, index=False)
    return export, model


@pytest.fixture
def legacy_report(tmp_path: Path) -> tuple[Path, Path, Path]:
    """(relatório legado, exportação, modelo) numa pasta temporária."""
    export, model = _write_tables(tmp_path)
    report = _write_legacy_report(tmp_path / "mapping_cliente_full.xlsx")
    return report, export, model


def _migrate(report: Path, export: Path, model: Path, *extra: str) -> int:
    return migrate_reports.main(
        [str(report), str(export), "--model", str(model), "--reference-data",
         str(ROOT / "reference_data"), *extra]
    )


def _rows(sheet) -> tuple[list[object], list[tuple[object, ...]]]:
    rows = list(sheet.iter_rows(values_only=True))
    return list(rows[0]), rows[1:]


def _column(sheet, header: str) -> list[object]:
    names, rows = _rows(sheet)
    position = names.index(header)
    return [row[position] for row in rows]


def _snapshot(path: Path) -> dict[str, list[list[tuple[object, str, str]]]]:
    """Valor, tipo e formato de número de cada célula de cada aba, na ordem das abas."""
    book = openpyxl.load_workbook(path)
    return {
        sheet.title: [
            [(cell.value, type(cell.value).__name__, cell.number_format) for cell in row]
            for row in sheet.iter_rows()
        ]
        for sheet in book.worksheets
    }


# ---------------------------------------------------------------------------
# Migração do relatório sintético
# ---------------------------------------------------------------------------


def test_synthetic_legacy_report_becomes_the_portuguese_schema(legacy_report, capsys) -> None:
    report, export, model = legacy_report
    assert _migrate(report, export, model) == 0
    out = capsys.readouterr().out
    # Os hashes do relatório sintético não conferem: só a tradução por modelo.
    assert "Aviso: conferência por regeneração não feita: o sha256_cliente gravado" in out
    assert "Verificação: ok (aceito=3 revisao=4 sem_correspondencia=1)." in out
    assert "Preservados linha a linha: 3 pareceres e 2 observações do revisor." in out
    assert f"Arquivo gravado: {report}" in out

    book = openpyxl.load_workbook(report)
    assert book.sheetnames == [
        "Mapeamentos", "Metadados da Execução", "Colunas do Cliente sem Uso", "Fila de Revisão",
    ]
    mappings = book["Mapeamentos"]
    # Mesma ordem de colunas do arquivo legado, só com os nomes em português.
    assert _rows(mappings)[0] == [rs.COLUMNS["mappings"][name] for name in LEGACY_MAPPINGS_HEADER]
    assert mappings["A1"].font.b is True
    assert mappings.column_dimensions["L"].width == 60

    assert _column(mappings, "justificativa") == EXPECTED_REASONS
    assert _column(mappings, "situacao") == [
        "aceito", "revisao", "aceito", "sem_correspondencia", "revisao", "aceito", "revisao",
        "revisao",
    ]
    assert _column(mappings, "metodo") == [
        "nome", "dados", "dados", "dados", "dados", "catalogo", "erro_dados", "somente_nome",
    ]
    assert _column(mappings, "coluna_cliente") == [
        "CFOP", "Valor ICMS (ocorrência 2)", "Cod Mun", None, "Valor ICMS ST", "CFOP", "Valor",
        "Coluna sem nome 9",
    ]
    assert _column(mappings, "apoiado_por") == [
        "lexico + catalogo (concorda)", "nao_registrado", "resolvedor_ia + semantico + catalogo",
        None, "semantico", "catalogo", None, None,
    ]
    assert _column(mappings, "decidido_por")[:3] == ["lexico", None, "resolvedor_ia"]
    assert _column(mappings, "vetado_por")[1] == "coluna_vazia"
    assert _column(mappings, "nome_catalogo")[2] == "municipio_ibge"
    assert _column(mappings, "prova_aritmetica")[1] == EXPECTED_NOT_DECISIVE_PROOF
    assert not rs.proof_is_decisive(_column(mappings, "prova_aritmetica")[1])
    # Números e nomes do modelo intactos, com o tipo.
    assert _column(mappings, "confianca_final") == [row["final_confidence"] for row in MAPPING_ROWS]
    assert _column(mappings, "coluna_modelo") == [row["model_column"] for row in MAPPING_ROWS]

    unmatched = book["Colunas do Cliente sem Uso"]
    assert _column(unmatched, "coluna_cliente") == [
        "Valor ICMS", "Doc (ocorrência 2)", "CNPJ Emit",
    ]
    assert _column(unmatched, "disposicao") == [
        "proposta_nao_aceita", "nunca_referenciada", "nunca_referenciada",
    ]
    assert _column(unmatched, "prioridade_revisao") == ["prioritaria", "baixa", "prioritaria"]
    assert _column(unmatched, "sinais") == ["aritmetica,catalogo", None, "dv_cnpj,data"]
    assert _column(unmatched, "detalhe_sinais") == [
        "participa como valor de um trio comprovado de valor = base x alíquota; 88,2% dos seus "
        "17 códigos distintos são válidos na tabela oficial de CFOP",
        "Esta coluna não tem valores.",
        "100,0% dos valores passam na validação do dígito verificador de CNPJ; contém datas",
    ]

    queue = book["Fila de Revisão"]
    assert _column(queue, "coluna_cliente_proposta") == [
        "Valor ICMS (ocorrência 2)", "Valor ICMS ST", "Coluna sem nome 9",
    ]
    assert _column(queue, "divergencia_ia") == [
        None, "o resolvedor de IA atribuiu confiança 0,75", None,
    ]
    assert _column(queue, "prova_aritmetica") == [EXPECTED_NOT_DECISIVE_PROOF, None, None]
    assert _column(queue, "amostra_valores_cliente") == [None, "1 | 2 | 3", "x | y"]
    assert _column(queue, "motivo_retencao") == [
        "a coluna proposta não tem linhas preenchidas - confirmá-la entregaria uma coluna em "
        "branco",
        "abaixo do limiar de aceite",
        "só há evidência de nome - este caminho não aceita automaticamente, por mais que os "
        "cabeçalhos coincidam",
    ]

    metadata = book["Metadados da Execução"]
    header, rows = _rows(metadata)
    assert header == ["chave", "valor"]
    values = dict(rows)
    assert list(values) == [rs.METADATA_KEYS[key] for key, _ in LEGACY_METADATA] + [
        "migrado_da_versao_esquema"
    ]
    assert values["versao_esquema_mapeamento"] == 3
    assert values["migrado_da_versao_esquema"] == 2
    assert values["gerado_em_utc"] == "2026-08-25T17:26:27.092955+00:00"
    assert values["modelo_ia"] == "offline"
    assert values["sha256_manifesto_referencia"] == "desativado"


def test_reviewer_verdicts_and_notes_survive_byte_for_byte(legacy_report) -> None:
    report, export, model = legacy_report
    assert _migrate(report, export, model) == 0
    mappings = openpyxl.load_workbook(report)["Mapeamentos"]
    assert _column(mappings, "parecer_revisor")[:4] == ["correto", "incorreto", None, "inexistente"]
    corrected = _column(mappings, "coluna_cliente_correta")
    # O inteiro digitado pelo revisor continua inteiro.
    assert corrected[1] == 6 and type(corrected[1]) is int
    notes = _column(mappings, "observacao_revisor")
    # Observação igual a uma frase do mapeador: nunca é traduzida.
    assert notes[0] == HUMAN_NOTE_LOOKS_LIKE_TEMPLATE
    assert notes[1] == HUMAN_NOTE

    # E o apply lê os pareceres migrados como os originais.
    frame, schema = rs.to_internal_frame(pd.read_excel(report, sheet_name=0), "mappings")
    assert schema == rs.PT_BR
    assert [rs.normalize_reviewer_status(value) for value in frame["reviewer_status"]][:4] == [
        "correct", "incorrect", "", "no_mapping_exists",
    ]
    assert frame["status"].tolist() == [row["status"] for row in MAPPING_ROWS]
    assert frame["client_column"].tolist()[1] == "Valor ICMS (occurrence 2)"
    assert frame["supported_by"].tolist()[:2] == ["lexical + catalog (agrees)", "unrecorded"]


def test_migrating_twice_is_identical_cell_by_cell(legacy_report, capsys) -> None:
    report, export, model = legacy_report
    assert _migrate(report, export, model) == 0
    first_bytes = report.read_bytes()
    first = _snapshot(report)
    capsys.readouterr()

    assert _migrate(report, export, model) == 0
    assert capsys.readouterr().out.strip() == ALREADY_PORTUGUESE.format(report=report)
    assert report.read_bytes() == first_bytes
    assert _snapshot(report) == first
    # E as frases migradas, traduzidas de novo, não mudam.
    book = openpyxl.load_workbook(report)
    for table, fields in legacy_reports.TEXT_FIELDS.items():
        if table not in migrate_reports.TABLES:
            continue
        sheet = book[rs.sheet_name(table)]
        for name in fields:
            for value in _column(sheet, rs.COLUMNS[table][name]):
                assert legacy_reports.translate_cell(table, name, value) == value



def test_dry_run_leaves_the_report_untouched(legacy_report, capsys) -> None:
    report, export, model = legacy_report
    before = report.read_bytes()
    assert _migrate(report, export, model, "--dry-run") == 0
    out = capsys.readouterr().out
    assert "Verificação: ok (" in out
    assert "Simulação (--dry-run): o relatório não foi alterado." in out
    assert report.read_bytes() == before
    assert list(report.parent.glob("*.migrando.xlsx")) == []


def test_a_failed_verification_keeps_the_original(legacy_report, monkeypatch) -> None:
    report, export, model = legacy_report
    before = report.read_bytes()
    translate_table = migrate_reports._translate_table

    def corrupting(worksheet, table, summary):
        translate_table(worksheet, table, summary)
        if table == "mappings":
            worksheet.cell(2, LEGACY_MAPPINGS_HEADER.index("final_confidence") + 1).value = 0.5

    monkeypatch.setattr(migrate_reports, "_translate_table", corrupting)
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    assert "A verificação da migração falhou" in str(failure.value.code)
    assert "Mapeamentos!I2" in str(failure.value.code)
    assert report.read_bytes() == before
    assert list(report.parent.glob("*.migrando.xlsx")) == []


def test_a_mixed_schema_report_is_refused_without_writing(tmp_path) -> None:
    export, model = _write_tables(tmp_path)
    report = _write_legacy_report(
        tmp_path / "misto.xlsx", metadata_title=rs.sheet_name("run_metadata")
    )
    before = report.read_bytes()
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    assert "Esquema misto" in str(failure.value.code)
    assert report.read_bytes() == before


def test_a_portuguese_report_with_a_legacy_sheet_is_mixed_not_already_migrated(
    legacy_report, capsys
) -> None:
    report, export, model = legacy_report
    assert _migrate(report, export, model) == 0
    capsys.readouterr()
    book = openpyxl.load_workbook(report)
    book[rs.sheet_name("run_metadata")].title = rs.sheet_name("run_metadata", rs.LEGACY_EN)
    book.save(report)
    before = report.read_bytes()
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    message = str(failure.value.code)
    assert message.startswith("Esquema misto: o relatório está no esquema em português")
    assert "já está no esquema em português" not in capsys.readouterr().out
    assert report.read_bytes() == before


def test_portuguese_machine_values_in_a_legacy_report_are_refused(tmp_path) -> None:
    export, model = _write_tables(tmp_path)
    report = _write_legacy_report(tmp_path / "mapping_cliente_full.xlsx")
    book = openpyxl.load_workbook(report)
    mappings = book["Mappings"]
    mappings.cell(2, LEGACY_MAPPINGS_HEADER.index("status") + 1).value = "aceito"
    mappings.cell(2, LEGACY_MAPPINGS_HEADER.index("decided_by") + 1).value = "lexico"
    # O parecer em português é entrada legítima do revisor em qualquer esquema.
    mappings.cell(3, LEGACY_MAPPINGS_HEADER.index("reviewer_status") + 1).value = "Incorreto"
    book.save(report)
    before = report.read_bytes()
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    message = str(failure.value.code)
    assert message.startswith(
        "Esquema misto: o relatório está no esquema legado em inglês, mas 2 células já têm o "
        "valor em português: Mappings!K2 (status): 'aceito'; Mappings!Q2 (decided_by): 'lexico'."
    )
    assert "reviewer_status" not in message
    assert report.read_bytes() == before


def test_a_file_that_is_not_a_workbook_or_has_mappings_out_of_place_is_refused(
    tmp_path,
) -> None:
    export, model = _write_tables(tmp_path)
    text = tmp_path / "nao_e_planilha.xlsx"
    text.write_text("isto não é uma planilha\n", encoding="utf-8")
    with pytest.raises(SystemExit) as failure:
        _migrate(text, export, model)
    assert str(failure.value.code).startswith(f"{text} não é uma planilha .xlsx válida")

    report = _write_legacy_report(tmp_path / "mapping_cliente_full.xlsx")
    book = openpyxl.load_workbook(report)
    book.move_sheet("Run Metadata", offset=-1)
    book.save(report)
    before = report.read_bytes()
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    assert str(failure.value.code).startswith(
        "A aba 'Mappings' precisa ser a primeira do relatório"
    )
    assert report.read_bytes() == before

    with pytest.raises(SystemExit) as failure:
        _migrate(tmp_path, export, model)
    assert str(failure.value.code) == f"Esperava um arquivo, mas {tmp_path} é uma pasta."


def test_empty_text_cells_keep_their_type_and_the_check_sees_it(
    legacy_report, monkeypatch
) -> None:
    report, export, model = legacy_report
    book = openpyxl.load_workbook(report)
    # O mapeador grava os campos vazios como texto vazio.
    book["Mappings"].cell(2, LEGACY_MAPPINGS_HEADER.index("vetoed_by") + 1).value = ""
    book.save(report)
    coordinate = f"{openpyxl.utils.get_column_letter(LEGACY_MAPPINGS_HEADER.index('vetoed_by') + 1)}2"
    assert openpyxl.load_workbook(report)["Mappings"][coordinate].data_type == "inlineStr"
    before = report.read_bytes()

    # Sem a regravação do texto vazio, a célula viraria uma célula em branco,
    # e a verificação embutida precisa barrar isso.
    monkeypatch.setattr(migrate_reports, "_keep_empty_text_cells", lambda workbook: None)
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    assert f"Mapeamentos!{coordinate}" in str(failure.value.code)
    assert "tipo ou formatação da célula diferente" in str(failure.value.code)
    assert report.read_bytes() == before

    monkeypatch.undo()
    assert _migrate(report, export, model) == 0
    cell = openpyxl.load_workbook(report)["Mapeamentos"][coordinate]
    assert cell.value is None and cell.data_type == "inlineStr"


def test_a_regeneration_mismatch_refuses_without_writing(tmp_path) -> None:
    export, model = _write_tables(tmp_path)
    metadata = [
        (key, im.file_sha256(export) if key == "client_sha256"
         else im.file_sha256(model) if key == "model_sha256" else value)
        for key, value in LEGACY_METADATA
    ]
    report = _write_legacy_report(tmp_path / "mapping_cliente_full.xlsx", metadata=metadata)
    before = report.read_bytes()
    with pytest.raises(SystemExit) as failure:
        _migrate(report, export, model)
    # Os hashes conferem, mas as abas derivadas do relatório sintético não saem
    # desta exportação: com as mesmas entradas, o mapeador atual não reproduz o
    # arquivado, e a migração não grava nada.
    message = str(failure.value.code)
    assert message.startswith("A conferência por regeneração encontrou ")
    assert " com o arquivado; o relatório não foi alterado:" in message
    assert "o relatório não foi alterado" in message
    assert report.read_bytes() == before
    assert list(report.parent.glob("*.migrando.xlsx")) == []


def test_command_line_is_in_portuguese(capsys, tmp_path) -> None:
    with pytest.raises(SystemExit) as done:
        migrate_reports.main(["--help"])
    assert done.value.code == 0
    out = capsys.readouterr().out
    assert out.startswith("uso: ")
    assert "RELATORIO EXPORTACAO" in out
    assert "--dry-run" in out and "argumentos posicionais" in out

    with pytest.raises(SystemExit) as missing:
        migrate_reports.main([str(tmp_path / "nao_existe.xlsx"), str(tmp_path / "x.xlsx")])
    assert str(missing.value.code).startswith("Arquivo não encontrado: ")


# ---------------------------------------------------------------------------
# legacy_reports: o que passa intacto
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "value",
    [
        LLM_PROSE,
        EXPECTED_REASONS[0],
        EXPECTED_NOT_DECISIVE_PROOF,
        "No unused client columns remain. And something else.",
        "  No unused client columns remain.",
        "",
        None,
        6,
        float("nan"),
    ],
)
def test_text_that_matches_no_template_passes_through(value) -> None:
    for kind in legacy_reports.KINDS:
        translated = legacy_reports.translate_text(value, kind)
        assert translated is value or translated == value


def test_a_signal_detail_with_one_unknown_sentence_is_left_whole() -> None:
    text = "holds dates; something a person wrote"
    assert legacy_reports.translate_cell("unmatched", "signal_detail", text) == text


def test_reviewer_note_and_model_columns_are_never_translated() -> None:
    sentence = "below the accept threshold"
    for field in ("reviewer_note", "model_column", "client_sample_values", "closest_model_column"):
        for table in ("mappings", "queue", "unmatched"):
            assert legacy_reports.translate_cell(table, field, sentence) == sentence
            assert migrate_reports.migrate_value(table, field, sentence) == sentence


def test_why_held_agrees_with_the_report_schema_phrases() -> None:
    for english, portuguese in rs.ENUMS["why_held"].items():
        assert legacy_reports.translate_cell("queue", "why_held", english) == portuguese


# ---------------------------------------------------------------------------
# Antidesvio: frase antiga traduzida == frase da função produtora atual
# ---------------------------------------------------------------------------
#
# As funções _legacy_* abaixo são cópias do código de 2ae4b6d:invoice_mapper.py
# (as linhas estão citadas). Elas produzem o inglês que os relatórios
# arquivados contêm; a função produtora atual, chamada com as mesmas
# entradas, produz o português que o mapeador grava hoje.


def _legacy_trio_proof_text(trio):
    # 2ae4b6d:invoice_mapper.py:1904-1916
    if trio is None:
        return None
    text = (
        f"valor = base x aliquota holds for {trio['tax_scope']} on "
        f"{trio['rows_tested']} rows at fit {trio['fit_score']}"
    )
    alternatives = int(trio.get("alternatives", 1))
    if alternatives > 1:
        text += (
            f" - but this client has {alternatives} disjoint {trio['tax_scope']} "
            "blocks that each satisfy it, so this does not single out this column"
        )
    return text


def _legacy_unmatched_signals(*, column, profile, values, reference_catalogs, trio_slots):
    # 2ae4b6d:invoice_mapper.py:2101-2146
    signals: list[str] = []
    details: list[str] = []

    slot = trio_slots.get(column.index)
    if slot is not None:
        signals.append("arithmetic")
        details.append(
            f"participates in a proven valor = base x aliquota trio as the {slot}"
        )

    if reference_catalogs is not None:
        evidence = im.best_catalog_evidence(
            [
                reference_catalogs.analyze(values, name)
                for name in reference_catalogs.catalogs
            ]
        )
        if evidence is not None and im.catalog_evidence_is_distinctive(
            evidence, final_threshold=1.0
        ):
            signals.append("catalog")
            details.append(
                f"{float(evidence['distinct_coverage']):.1%} of its "
                f"{int(evidence['unique_formatted_count'])} distinct codes are valid "
                f"{evidence['catalog']}"
            )

    for kind in ("nfe_key", "cnpj", "cpf"):
        fraction = float(profile.get(f"valid_{kind}_fraction", 0.0) or 0.0)
        if fraction >= 0.80:
            signals.append(f"{kind}_checkdigit")
            details.append(f"{fraction:.1%} of values pass the {kind} check digit")
            break

    if profile.get("inferred_type") == "datetime":
        signals.append("date")
        details.append("holds dates")

    if not signals:
        quality = profile.get("evidence_quality")
        if quality == "empty":
            details.append("No values in this column.")
        elif quality == "all_zero":
            details.append("Every value in this column is zero.")
        else:
            details.append("Carries data, but no signal matched. Listed for completeness.")
    return signals, "; ".join(details)


def _legacy_catalog_reason(evidence):
    # 2ae4b6d:invoice_mapper.py:1597-1602
    return (
        f"Versioned official {evidence['catalog']} codes cover "
        f"{float(evidence.get('distinct_coverage', evidence['coverage'])):.1%} "
        f"of the {int(evidence['unique_formatted_count'])} distinct codes used "
        f"({float(evidence['coverage']):.1%} of rows) and uniquely "
        "identify this candidate among the available columns."
    )


def _legacy_conflict_reason(header_conflict):
    # 2ae4b6d:invoice_mapper.py:2947-2950
    return (
        f"Explicit {header_conflict}; keep this proposal for review because the "
        "row content provides distinctive evidence that the client header may be bogus."
    )


def _legacy_vacuity_reason(chosen):
    # 2ae4b6d:invoice_mapper.py:2972-2976
    return (
        f"{chosen.display_name!r} has no populated rows, so it cannot be "
        "this column's source however well the headers match. "
        "Not accepted."
    )


def _legacy_queue_texts(result, table, column, final_threshold):
    # 2ae4b6d:invoice_mapper.py:3107-3142
    final = float(result.final_confidence or 0.0)
    vacuous = im.client_column_is_vacuous(table, column)
    llm_final = result.llm_final_confidence
    disagrees = (
        f"resolver scored it {llm_final}"
        if llm_final is not None and float(llm_final) < final - 0.10
        else ""
    )
    clamped = abs(final - (final_threshold - 0.0001)) < 1e-9
    why_held = (
        "the proposed column has no populated rows - confirming "
        "this would ship a blank column"
        if vacuous
        else "name evidence only - this branch cannot auto-accept "
        "however well the headers match"
        if clamped
        else "below the accept threshold"
    )
    return disagrees, why_held


#: Justificativas fixas: inglês de 2ae4b6d -> constante atual. As funções que
#: as gravam (dentro de map_columns) não são chamáveis isoladamente.
FIXED_REASONS = {
    "No unused client columns remain.": im.REASON_NO_UNUSED,
    "Accepted from header similarity with a clear lead over the runner-up.": (
        im.REASON_PASS_ONE_ACCEPTED
    ),
    "Header score or runner-up margin was insufficient; data comparison was disabled.": (
        im.REASON_OFFLINE
    ),
    "A unique high-coverage fiscal relationship supports this column "
    "over alternatives (for example, base times rate reproduces value).": (
        im.REASON_UNIQUE_RELATIONSHIP
    ),
    "Distinctive non-zero sample values overlap exactly with this "
    "candidate and not with the competing columns.": im.REASON_OVERLAP,
    "The header and a row-wise fiscal relationship jointly "
    "support this column with a clear lead over alternatives.": im.REASON_JOINT,
    "The resolver found supporting evidence but did not commit to a column; "
    "retain the strongest header candidate for review.": im.REASON_NULL_FALLBACK,
    "The strongest header candidate has weak row evidence; preserve the "
    "name-based mapping for review instead of replacing it with a merely "
    "similar value shape.": im.REASON_PRESERVE_BEST,
    "No candidate had enough combined header and distinctive row evidence; "
    "the client export may not contain this model field.": im.REASON_UNSUPPORTED_NO_MATCH,
    "A high-coverage fiscal equation strongly contradicts this proposal; "
    "keep it for review instead of accepting it automatically.": im.REASON_CONTRADICTION,
    "The header match is ambiguous and the row samples are empty or all-zero; "
    "keep this mapping for review.": im.REASON_AMBIGUOUS_EMPTY,
}


def _queue_table() -> im.SheetTable:
    columns = [
        im.ColumnRef(id=f"column_{i + 1:03d}", index=i, name=name, occurrence=1)
        for i, name in enumerate(["Cheia", "Vazia", "Travada"])
    ]
    return im.SheetTable(columns=columns, data=pd.DataFrame([["x", None, "y"], ["z", None, "w"]]))


def _queue_result(number, client_number, final, llm_final=None):
    return im.MappingResult(
        model_column_number=number, model_column=f"Modelo {number}",
        client_column_number=client_number, client_column=f"Coluna {client_number}",
        name_confidence=final, name_runner_up_confidence=final, name_margin=0.0,
        data_confidence=None, final_confidence=final, method="data", status="review",
        reason="", llm_final_confidence=llm_final,
    )


def _drift_cases() -> list[tuple[str, str, str, str]]:
    """(campo, inglês antigo, português da produtora atual, rótulo) para cada modelo."""
    cases: list[tuple[str, str, str, str]] = []

    for english, portuguese in FIXED_REASONS.items():
        cases.append(("reason", english, portuguese, "justificativa fixa"))

    # Conflito de cabeçalho: o inglês continua sendo o de semantic_header_conflict.
    for left, right in (
        ("Valor ICMS", "Valor PIS"),
        ("Valor PIS", "Valor ICMS e COFINS"),
        ("Valor ICMS ST", "Valor ICMS"),
        ("Base ICMS", "Aliquota ICMS"),
    ):
        english = im.semantic_header_conflict(left, right)
        assert english is not None
        cases.append((
            "reason",
            _legacy_conflict_reason(english),
            im.REASON_CONFLICT_REVIEW.format(
                header_conflict=im.describe_header_conflict_pt(left, right)
            ),
            f"conflito {left} x {right}",
        ))

    # Vacuidade: a mesma expressão de map_columns, com o nome de exibição interno.
    for name, occurrence in (("Nota Fiscal", 1), ("Aliq. Cofins", 2), ("Unnamed column 7", 3),
                             ("it's \"quoted\"", 1)):
        chosen = im.ColumnRef(id="column_001", index=0, name=name, occurrence=occurrence)
        cases.append((
            "reason",
            _legacy_vacuity_reason(chosen),
            im.REASON_VACUITY.format(
                client_column=rs.display_name_to_external(chosen.display_name)
            ),
            f"vacuidade {chosen.display_name}",
        ))

    for exc in (ValueError("could not convert string to float: 'abc'"), KeyError("x")):
        cases.append((
            "reason",
            f"Data comparison failed: {type(exc).__name__}: {exc}",
            im.REASON_DATA_ERROR.format(exc_type=type(exc).__name__, exc=exc),
            f"erro {type(exc).__name__}",
        ))

    # Decisão por catálogo: a produtora atual, com a mesma evidência.
    column = im.ColumnRef(id="column_001", index=0, name="CFOP", occurrence=1)
    for catalog in ("cfop", "ibge_municipality", "cst_icms_taxation"):
        for distinct, coverage, count in ((1.0, 1.0, 62), (0.882352, 0.6667, 17)):
            evidence = {
                "catalog": catalog, "specificity": 1.0, "score": 0.99, "all_zero_only": False,
                "nonempty_count": 100, "format_fraction": 1.0, "unique_valid_count": 10,
                "valid_count": 10, "distinct_coverage": distinct, "coverage": coverage,
                "unique_formatted_count": count,
            }
            decision = im._unique_distinctive_catalog_decision(
                candidates=[(column, pd.Series([1]), 0.5)], catalog_evidence={0: [evidence]}
            )
            cases.append((
                "reason", _legacy_catalog_reason(evidence), decision.reason, f"catálogo {catalog}"
            ))

    # Prova aritmética, com uma e com várias alternativas.
    for scope, rows, fit, alternatives in (
        ("pis", 5501, 0.9996, 1), ("icms", 120, 1.0, 2), ("cofins", 30, 0.97123, 3),
    ):
        trio = {"tax_scope": scope, "rows_tested": rows, "fit_score": fit,
                "alternatives": alternatives}
        cases.append((
            "arithmetic_proof", _legacy_trio_proof_text(trio), im._trio_proof_text(trio),
            f"prova {scope} x{alternatives}",
        ))

    # Detalhe dos sinais: as duas versões de _unmatched_signals, mesmas entradas.
    catalogs = ReferenceCatalogs.from_directory(ROOT / "reference_data")
    cfop = pd.Series(["5102", "6102", "1102", "5405", "6108", "5949", "2102", "5101"] * 5)
    empty = pd.Series([], dtype=object)
    for slot in ("base", "rate", "value"):
        for values, profile, references in (
            (cfop, im.summarize_values(cfop, 8), catalogs),
            (empty, {"valid_cnpj_fraction": 0.9, "inferred_type": "datetime"}, None),
            (empty, {"valid_nfe_key_fraction": 0.8333333}, None),
            (empty, {"valid_cpf_fraction": 1.0}, None),
        ):
            arguments = dict(column=column, profile=profile, values=values,
                             reference_catalogs=references, trio_slots={0: slot})
            english = _legacy_unmatched_signals(**arguments)
            portuguese = im._unmatched_signals(**arguments)
            assert english[0] == portuguese[0]
            cases.append(("signal_detail", english[1], portuguese[1], f"sinais {slot}"))
    for profile in ({"evidence_quality": "empty"}, {"evidence_quality": "all_zero"},
                    {"evidence_quality": "informative"}, {"inferred_type": "datetime"}):
        arguments = dict(column=column, profile=profile, values=empty,
                         reference_catalogs=None, trio_slots={})
        cases.append((
            "signal_detail", _legacy_unmatched_signals(**arguments)[1],
            im._unmatched_signals(**arguments)[1], f"sinais {profile}",
        ))

    # Fila de revisão: motivo e divergência da produtora atual.
    table = _queue_table()
    results = [
        _queue_result(1, 1, 0.88, llm_final=0.75),
        _queue_result(2, 2, 0.95),
        _queue_result(3, 3, 0.8999),
    ]
    rows = im.summarize_adjudication_queue(
        client_table=table, results=results, final_threshold=0.9, sample_size=8
    )
    by_model = {row.model_column_number: row for row in rows}
    for result in results:
        row = by_model[result.model_column_number]
        disagrees, why_held = _legacy_queue_texts(
            result, table, table.columns[result.client_column_number - 1], 0.9
        )
        cases.append(("why_held", why_held, row.why_held, f"fila {result.model_column_number}"))
        if disagrees:
            cases.append((
                "resolver_disagrees", disagrees, row.resolver_disagrees,
                f"divergência {result.model_column_number}",
            ))
    return cases


DRIFT_CASES = _drift_cases()


@pytest.mark.parametrize(
    "kind, english, expected",
    [case[:3] for case in DRIFT_CASES],
    ids=[case[3] for case in DRIFT_CASES],
)
def test_legacy_sentence_translates_to_what_the_current_producer_writes(
    kind, english, expected
) -> None:
    assert english != expected
    assert legacy_reports.translate_text(english, kind) == expected
    # Sem dizer o campo, o resultado é o mesmo; e traduzir de novo não muda.
    assert legacy_reports.translate_text(english) == expected
    assert legacy_reports.translate_text(expected, kind) == expected


def test_every_legacy_template_has_a_drift_case() -> None:
    covered = set()
    for kind, english, _, _ in DRIFT_CASES:
        if kind == "signal_detail":
            pieces = english.split(legacy_reports.SIGNAL_SEPARATOR)
            covered |= {legacy_reports.match_template(piece, "signal").id for piece in pieces}
        else:
            covered.add(legacy_reports.match_template(english, kind).id)
    assert covered == {template.id for template in legacy_reports.TEMPLATES}


def test_the_header_conflict_twin_covers_the_three_kinds() -> None:
    kinds = set()
    for left, right in (("Valor ICMS", "Valor PIS"), ("Valor ICMS ST", "Valor ICMS"),
                        ("Base ICMS", "Aliquota ICMS")):
        english = im.semantic_header_conflict(left, right)
        kinds.add(english.split(" ")[0])
        assert legacy_reports.translate_header_conflict(english) == (
            im.describe_header_conflict_pt(left, right)
        )
    assert kinds == {"tax-family", "tax-qualifier", "semantic-role"}


# ---------------------------------------------------------------------------
# Planilhas reais
# ---------------------------------------------------------------------------


def _clean_environment() -> dict[str, str]:
    environment = {key: value for key, value in os.environ.items() if key != "OPENAI_API_KEY"}
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    return environment


def _start_offline_runs(
    mapper: Path, folder: Path, cwd: Path, extra: Sequence[str] = ()
) -> dict[str, tuple[Path, subprocess.Popen]]:
    """Dispara ``mapper --offline`` para os quatro clientes, sem esperar."""
    return {
        client: (
            folder / f"mapping_{client}_offline.xlsx",
            subprocess.Popen(
                [sys.executable, str(mapper), str(ROOT / export), "--offline",
                 "--output", str(folder / f"mapping_{client}_offline.xlsx"), *extra],
                cwd=cwd, env=_clean_environment(),
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
            ),
        )
        for client, export in EXPORTS.items()
    }


def _finish_offline_runs(runs: dict[str, tuple[Path, subprocess.Popen]]) -> dict[str, Path]:
    reports = {}
    for client, (output, process) in runs.items():
        _, errors = process.communicate(timeout=1800)
        assert process.returncode == 0, errors
        reports[client] = output
    return reports


def _legacy_code(folder: Path) -> Path | None:
    """Extrai com ``git archive`` o código Python e os catálogos de 2ae4b6d."""
    archive = subprocess.run(
        ["git", "archive", "--format=tar", LEGACY_COMMIT, "--", "*.py", "reference_data"],
        cwd=ROOT, capture_output=True,
    )
    if archive.returncode != 0:
        return None
    extracted = subprocess.run(["tar", "-x", "-C", str(folder)], input=archive.stdout)
    return folder if extracted.returncode == 0 else None


@pytest.fixture(scope="module")
def offline_runs(tmp_path_factory) -> dict[str, dict[str, Path]]:
    """Execuções offline do código atual e do de 2ae4b6d, todas em paralelo.

    ``"atual"``: o mapeador de hoje, em português. ``"legado"``: a saída
    offline em inglês do código que gravou os relatórios legados. Usa
    ``MIGRACAO_BASELINE_EN`` se a variável aponta para os quatro arquivos;
    senão roda o mapeador antigo extraído do git, com as exportações e o
    modelo atuais (que não mudaram desde então).
    """
    fresh = _start_offline_runs(
        ROOT / "invoice_mapper.py", tmp_path_factory.mktemp("offline_atual"), ROOT
    )
    legacy: dict[str, tuple[Path, subprocess.Popen]] | None = None
    baseline: dict[str, Path] = {}
    if BASELINE_EN is not None:
        baseline = {client: BASELINE_EN / f"mapping_{client}_offline.xlsx" for client in EXPORTS}
    if not baseline or not all(path.exists() for path in baseline.values()):
        code = _legacy_code(tmp_path_factory.mktemp(f"codigo_{LEGACY_COMMIT}"))
        if code is not None:
            legacy = _start_offline_runs(
                code / "invoice_mapper.py", tmp_path_factory.mktemp("offline_legado"), code,
                ["--model", str(ROOT / "data/model.xlsx")],
            )
        baseline = {}
    runs = {"atual": _finish_offline_runs(fresh)}
    if legacy is not None:
        baseline = _finish_offline_runs(legacy)
    if not baseline:
        pytest.skip(f"código de {LEGACY_COMMIT} indisponível no git")
    runs["legado"] = baseline
    return runs


@pytest.fixture(scope="module")
def fresh_offline_reports(offline_runs) -> dict[str, Path]:
    """Uma execução offline nova do código atual para cada cliente."""
    return offline_runs["atual"]


@pytest.fixture(scope="module")
def baseline_offline_reports(offline_runs) -> dict[str, Path]:
    """A saída offline em inglês do código de 2ae4b6d para cada cliente."""
    return offline_runs["legado"]


@pytest.mark.lento
@requires_workbooks
@pytest.mark.parametrize("client", list(EXPORTS))
def test_old_english_output_translated_is_the_new_portuguese_output(
    client, fresh_offline_reports, baseline_offline_reports
) -> None:
    old_book = openpyxl.load_workbook(baseline_offline_reports[client])
    new_book = openpyxl.load_workbook(fresh_offline_reports[client])
    compared = 0
    for table in migrate_reports.TABLES:
        old_header, old_rows = _rows(old_book[rs.sheet_name(table, rs.LEGACY_EN)])
        new_header, new_rows = _rows(new_book[rs.sheet_name(table)])
        assert len(new_rows) == len(old_rows), table
        for name in sorted(legacy_reports.TEXT_FIELDS[table]):
            old_position = old_header.index(name)
            new_position = new_header.index(rs.COLUMNS[table][name])
            for line, (old_row, new_row) in enumerate(zip(old_rows, new_rows), start=2):
                old, new = old_row[old_position], new_row[new_position]
                if old in (None, ""):
                    assert new in (None, ""), (table, name, line)
                    continue
                translated = legacy_reports.translate_cell(table, name, old)
                # Offline não há prosa da IA: toda frase casa com um modelo.
                assert translated != old, (table, name, line, old)
                assert translated == new, (table, name, line)
                compared += 1
    assert compared > 0


_CELL_STYLES = ("font", "fill", "alignment", "border", "protection")


def _cells(path: Path) -> dict[str, list[list[tuple[object, ...]]]]:
    """Valor, tipo, tipo da célula no arquivo, formato e estilo de cada célula."""
    book = openpyxl.load_workbook(path)
    return {
        sheet.title: [
            [
                (cell.value, type(cell.value).__name__, cell.data_type, cell.number_format)
                + tuple(repr(getattr(cell, name)) for name in _CELL_STYLES)
                for cell in row
            ]
            for row in sheet.iter_rows()
        ]
        for sheet in book.worksheets
    }


@pytest.mark.lento
@requires_workbooks
@pytest.mark.parametrize("client", list(EXPORTS))
def test_a_migrated_old_offline_report_is_a_fresh_run_cell_by_cell(
    client, tmp_path, fresh_offline_reports, baseline_offline_reports
) -> None:
    report = tmp_path / f"mapping_{client}_offline.xlsx"
    shutil.copyfile(baseline_offline_reports[client], report)
    assert migrate_reports.migrate_report(
        report, ROOT / EXPORTS[client], ROOT / "data/model.xlsx", ROOT / "reference_data",
        log=lambda message: None,
    ).written
    migrated, fresh = _cells(report), _cells(fresh_offline_reports[client])
    assert list(migrated) == list(fresh)
    metadata = rs.sheet_name("run_metadata")
    *rows, marker = migrated.pop(metadata)
    fresh_rows = fresh.pop(metadata)
    # Só mudam a data da execução e a linha que registra a migração.
    assert (marker[0][0], marker[1][0]) == ("migrado_da_versao_esquema", 2)
    differing = [
        row[0][0] for row, other in zip(rows, fresh_rows, strict=True) if row != other
    ]
    assert differing == ["gerado_em_utc"]
    assert migrated == fresh


def _legacy_archived_report(client: str, folder: Path) -> Path:
    """Cópia do relatório arquivado no esquema legado.

    Usa o da árvore de trabalho enquanto ele ainda está em inglês; depois da
    migração dos arquivados, busca a versão legada no commit 2ae4b6d.
    """
    source = ROOT / f"data/reports/mapping_{client}_full.xlsx"
    target = folder / source.name
    if source.exists():
        shutil.copyfile(source, target)
        book = openpyxl.load_workbook(target, read_only=True)
        header = next(book.worksheets[0].iter_rows(max_row=1, values_only=True))
        book.close()
        if rs.detect_schema(header) == rs.LEGACY_EN:
            return target
    shown = subprocess.run(
        ["git", "show", f"{LEGACY_COMMIT}:data/reports/{source.name}"],
        cwd=ROOT, capture_output=True,
    )
    if shown.returncode != 0:
        pytest.skip(f"relatório legado de {client} indisponível")
    target.write_bytes(shown.stdout)
    return target


def _verdicts(sheet, schema: str) -> dict[int, tuple[object, object, object]]:
    header, rows = _rows(sheet)
    number = header.index(rs.header_label("mappings", "model_column_number", schema))
    positions = [header.index(rs.header_label("mappings", name, schema))
                 for name in im.REVIEWER_COLUMNS]
    return {
        row[number]: tuple(row[position] for position in positions)
        for row in rows
        if any(row[position] not in (None, "") for position in positions)
    }


@pytest.mark.lento
@requires_workbooks
def test_migrating_archived_sap_procosa_keeps_the_26_verdicts_and_apply_delivers_62(
    tmp_path, capsys
) -> None:
    report = _legacy_archived_report("sap_procosa", tmp_path)
    before = _verdicts(openpyxl.load_workbook(report).worksheets[0], rs.LEGACY_EN)
    export = ROOT / EXPORTS["sap_procosa"]
    assert _migrate(report, export, ROOT / "data/model.xlsx") == 0
    out = capsys.readouterr().out
    assert "Preservados linha a linha: 26 pareceres e 26 observações do revisor." in out

    book = openpyxl.load_workbook(report)
    # Layout antigo mantido: nenhuma aba nem coluna acrescentada.
    assert book.sheetnames == ["Mapeamentos", "Metadados da Execução"]
    assert book.worksheets[0].max_column == 23
    after = _verdicts(book.worksheets[0], rs.PT_BR)
    assert len(after) == 26 and set(after) == set(before)
    assert Counter(status for status, _, _ in after.values()) == {
        "correto": 15, "incorreto": 2, "inexistente": 9,
    }
    assert {number for number, (status, _, _) in after.items() if status == "inexistente"} == {
        1, 8, 15, 22, 23, 24, 25, 26, 27,
    }
    corrections = {number: column for number, (_, column, _) in after.items() if column is not None}
    assert corrections == {37: 6, 38: 32}
    assert all(type(column) is int for column in corrections.values())
    for number, (status, column, note) in before.items():
        new_status, new_column, new_note = after[number]
        assert rs.normalize_reviewer_status(new_status) == rs.normalize_reviewer_status(status)
        assert new_column == column and type(new_column) is type(column)
        assert new_note == note

    apply_invoice_mapping.main(
        [str(export), str(report), "--output", str(tmp_path / "sap_procosa_mapped.xlsx"),
         "--model", str(ROOT / "data/model.xlsx")]
    )
    captured = capsys.readouterr()
    assert "colunas_mapeadas=62" in captured.out
    # Relatório migrado: nenhum aviso de esquema legado.
    assert "esquema legado" not in captured.err


@pytest.mark.lento
@requires_workbooks
def test_archived_oracle_report_passes_the_regeneration_cross_check(tmp_path) -> None:
    report = _legacy_archived_report("oracle", tmp_path)
    before = report.read_bytes()
    messages: list[str] = []
    summary = migrate_reports.migrate_report(
        report, ROOT / EXPORTS["oracle"], ROOT / "data/model.xlsx",
        ROOT / "reference_data", dry_run=True, log=messages.append,
    )
    assert summary.regeneration is not None and summary.regeneration.ran
    assert summary.regeneration.mismatches == []
    assert summary.regeneration.compared == {
        migrate_reports.PROOFS_LABEL: 227, "Colunas do Cliente sem Uso": 30, "Fila de Revisão": 37,
    }
    assert summary.status_counts == {"accepted": 18, "review": 37, "no_match": 172}
    assert summary.unexpected_english == []
    assert sum(summary.sentences_kept.values()) == 160  # a prosa da IA
    assert report.read_bytes() == before
