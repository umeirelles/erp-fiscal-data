"""Testes de ``mapper/evaluate.py``: gabaritos, classificação e comando evaluate."""

from __future__ import annotations

from dataclasses import asdict
import json
from pathlib import Path
import re

import openpyxl
import pytest

from mapper import evaluate, schema
from mapper.evaluate import (
    KEY_CORRECTIONS,
    AnswerKey,
    Outcome,
    Row,
    Score,
    classify,
    column_letter_to_index,
    decision_census,
    index_to_column_letter,
    load_answer_key,
    parse_coluna_annotation,
    read_mappings,
    score_report,
)

from support import report_files

ROOT = Path(__file__).resolve().parent.parent
KEYED = {
    "datasul": (
        "data/clients/datasul_maio.xlsx",
        "data/reports/mapping_datasul_full.xlsx",
        "data/keys/de_para_datasul.xlsx",
    ),
    "oracle": (
        "data/clients/oracle.xlsx",
        "data/reports/mapping_oracle_full.xlsx",
        "data/keys/de_para_oracle.xlsx",
    ),
}
requires_workbooks = pytest.mark.skipif(
    not (ROOT / "data/keys/de_para_oracle.xlsx").exists(),
    reason="planilhas dos clientes ausentes",
)

#: Exemplo de explicação; sua redação não determina a métrica.
PORTUGUESE_NON_DECISIVE_PROOF = (
    "valor = base x alíquota se confirma para PIS em 5501 linhas, com aderência "
    "0,9996 - mas este cliente tem 2 blocos disjuntos de PIS que satisfazem a "
    "equação, portanto ela não distingue esta coluna"
)


@pytest.fixture(scope="module")
def baseline_reports(tmp_path_factory):
    """Reconstrói somente os inputs históricos necessários à regressão do evaluator."""
    fixture = json.loads((ROOT / "tests/fixtures/evaluation_baseline.json").read_text())
    directory = tmp_path_factory.mktemp("evaluation_baseline")
    paths = {}
    for name, rows in fixture["clients"].items():
        book = openpyxl.Workbook()
        book.active.title = "Mapeamentos"
        book.active.append(fixture["headers"])
        for row in rows:
            book.active.append(row)
        paths[name] = directory / f"{name}.xlsx"
        book.save(paths[name])
        book.close()
    return paths


@pytest.mark.parametrize(
    ("letters", "index"),
    [("A", 1), ("Z", 26), ("AA", 27), ("AZ", 52), ("BA", 53), ("BV", 74), ("FY", 181), ("GZ", 208)],
)
def test_column_letters_round_trip(letters: str, index: int) -> None:
    assert column_letter_to_index(letters) == index
    assert index_to_column_letter(index) == letters


def test_annotation_tolerates_the_trailing_space_in_the_real_key() -> None:
    # Está assim na linha 1 de de_para_datasul.xlsx.
    assert parse_coluna_annotation("Coluna BV ") == 74


def test_annotation_ignores_cells_that_are_not_annotations() -> None:
    assert parse_coluna_annotation(None) is None
    assert parse_coluna_annotation("Base CBS") is None


def test_a_delivered_vacuous_target_is_fabrication_even_when_the_key_agrees() -> None:
    assert classify(
        chosen_client_column=5,
        expected_client_column=5,
        chosen_is_vacuous=True,
        delivered=True,
    ) is Outcome.FABRICATED


def test_an_undelivered_vacuous_proposal_is_not_a_fabrication() -> None:
    # O apply não preenche uma linha em revisão, então ela não chega ao cliente.
    assert classify(
        chosen_client_column=5,
        expected_client_column=5,
        chosen_is_vacuous=True,
        delivered=False,
    ) is Outcome.CORRECT


def test_unkeyed_columns_are_never_credited() -> None:
    for chosen in (None, 5):
        assert classify(
            chosen_client_column=chosen,
            expected_client_column=None,
            chosen_is_vacuous=False,
        ) is Outcome.UNDETERMINED


def test_keyed_columns_are_correct_missed_or_wrong() -> None:
    def outcome(chosen: int | None) -> Outcome:
        return classify(
            chosen_client_column=chosen, expected_client_column=5, chosen_is_vacuous=False
        )

    assert outcome(5) is Outcome.CORRECT
    assert outcome(None) is Outcome.MISSED
    assert outcome(6) is Outcome.WRONG_COLUMN


@requires_workbooks
@pytest.mark.parametrize(("client", "expected"), [("datasul", 30), ("oracle", 26)])
def test_recall_denominator_is_the_present_set_never_227(client: str, expected: int) -> None:
    key = load_answer_key(ROOT / KEYED[client][2], client)
    assert key.denominator == expected != 227


@requires_workbooks
def test_harness_reproduces_the_corrected_baseline(baseline_reports) -> None:
    total = 0
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(baseline_reports[client], ROOT / client_file, key)
        total += score.count(Outcome.CORRECT)
    assert total == 43


@requires_workbooks
def test_the_four_key_corrections_are_worth_exactly_four_columns(baseline_reports) -> None:
    strict = corrected = 0
    for client, (client_file, report_file, key_file) in KEYED.items():
        for apply_corrections, bucket in ((False, "strict"), (True, "corrected")):
            key = load_answer_key(ROOT / key_file, client, apply_corrections=apply_corrections)
            score = score_report(baseline_reports[client], ROOT / client_file, key)
            if bucket == "strict":
                strict += score.count(Outcome.CORRECT)
            else:
                corrected += score.count(Outcome.CORRECT)
    # Esses totais pertencem à fixture histórica, independente das novas execuções.
    assert corrected - strict == 4
    assert (strict, corrected) == (39, 43)
    assert sum(len(v) for v in KEY_CORRECTIONS.values()) == 4


@requires_workbooks
@pytest.mark.parametrize(("client", "report"), [
    ("datasul", "data/reports/mapping_datasul_full.xlsx"),
    ("oracle", "data/reports/mapping_oracle_full.xlsx"),
])
def test_no_report_delivers_a_column_with_no_data(client: str, report: str) -> None:
    key = load_answer_key(ROOT / KEYED[client][2], client)
    score = score_report(ROOT / report, ROOT / KEYED[client][0], key)
    assert [
        (row.model_column, row.client_column) for row in score.vacuous_targets
    ] == []


@requires_workbooks
def test_delivered_recall_is_never_above_picking_recall() -> None:
    # A entrega é um subconjunto da escolha.
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        assert score.delivered_correct <= score.count(Outcome.CORRECT)
        assert len(score.withheld_correct) == (
            score.count(Outcome.CORRECT) - score.delivered_correct
        )


@requires_workbooks
def test_the_accept_gate_withholds_eleven_correct_mappings(baseline_reports) -> None:
    # A maior perda isolada do pipeline: corretos encontrados e não entregues.
    found = delivered = 0
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(baseline_reports[client], ROOT / client_file, key)
        found += score.count(Outcome.CORRECT)
        delivered += score.delivered_correct
    assert found - delivered == 11
    assert (found, delivered) == (43, 32)


@requires_workbooks
def test_every_withheld_correct_mapping_sits_below_the_accept_threshold(baseline_reports) -> None:
    # Se uma linha retida passar de 0,90, é outra recusa, e não o threshold, que a segura.
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(baseline_reports[client], ROOT / client_file, key)
        for row in score.withheld_correct:
            assert (row.final_confidence or 0.0) < 0.90, (
                f"{client} m{row.model_column} retida com {row.final_confidence}"
            )


def _synthetic_client(path: Path) -> Path:
    """Exportação mínima: c1 e c3 com dados, c2 sem nenhuma linha preenchida."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Base PIS", "Coluna vazia", "CFOP"])
    sheet.append([100.0, None, "5102"])
    sheet.append([200.0, None, "6102"])
    book.save(path)
    return path


#: Linhas de mapeamento no vocabulário interno.
_SYNTHETIC_ROWS = [
    {"model_column_number": 1, "model_column": "Base PIS", "client_column_number": 1,
     "client_column": "Base PIS", "final_confidence": 0.95, "method": "name",
     "status": "accepted", "arithmetic_proof": None, "decided_by": "semantic"},
    {"model_column_number": 2, "model_column": "CFOP", "client_column_number": 3,
     "client_column": "CFOP", "final_confidence": 0.93, "method": "catalog",
     "status": "accepted", "arithmetic_proof": None, "decided_by": "catalog"},
    {"model_column_number": 3, "model_column": "Isentas PIS", "client_column_number": 2,
     "client_column": "Coluna vazia", "final_confidence": 0.5, "method": "data",
     "status": "review", "arithmetic_proof": None, "decided_by": ""},
    {"model_column_number": 4, "model_column": "Base COFINS", "client_column_number": None,
     "client_column": None, "final_confidence": 0.2, "method": "data",
     "status": "no_match", "arithmetic_proof": None, "decided_by": ""},
]


def _synthetic_report(path: Path, proof: str) -> Path:
    """Relatório mínimo em português; a linha m1 leva a prova recebida."""
    fields = list(_SYNTHETIC_ROWS[0])
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = schema.sheet_name("mappings")
    sheet.append([schema.header_label("mappings", field) for field in fields])
    for number, row in enumerate(_SYNTHETIC_ROWS, start=1):
        values = dict(row, arithmetic_proof=proof if number == 1 else None)
        values = {f: schema.external_value("mappings", f, v) for f, v in values.items()}
        sheet.append([values[field] for field in fields])
    book.save(path)
    return path


def test_a_report_scores_without_the_client_workbooks(tmp_path: Path) -> None:
    # Relatório em português; a pontuação volta ao vocabulário interno.
    client = _synthetic_client(tmp_path / "cliente.xlsx")
    key = AnswerKey(client="sintetico", present={1: 1, 2: 3})
    portuguese = score_report(
        _synthetic_report(tmp_path / "pt.xlsx", PORTUGUESE_NON_DECISIVE_PROOF), client, key
    )
    assert [r.status for r in portuguese.rows] == ["accepted", "accepted", "review", "no_match"]
    assert [r.method for r in portuguese.rows] == ["name", "catalog", "data", "data"]
    assert [r.outcome for r in portuguese.rows] == [
        Outcome.CORRECT, Outcome.CORRECT, Outcome.UNDETERMINED, Outcome.UNDETERMINED,
    ]
    assert portuguese.delivered_correct == 2
    assert portuguese.correct_pairs == {(1, 1), (2, 3)}
    # m1 foi decidido pelo cabeçalho; m2, pelo catálogo.
    assert portuguese.objective_decision_share == 0.5
    # m3 propõe a coluna vazia, mas está em revisão: não é fabricação.
    assert [(r.model_column, r.client_column) for r in portuguese.vacuous_proposals] == [(3, 2)]
    assert portuguese.vacuous_targets == []


def test_a_half_translated_report_is_refused(tmp_path: Path) -> None:
    # Cabeçalhos em português numa aba com o nome antigo, ou um cabeçalho em
    # inglês numa aba em português.
    path = _synthetic_report(tmp_path / "pt.xlsx", "")
    book = openpyxl.load_workbook(path)
    book.active.title = schema.LEGACY_SHEETS["mappings"]
    book.save(path)
    with pytest.raises(ValueError, match="formato antigo"):
        read_mappings(path)

    path = _synthetic_report(tmp_path / "pt.xlsx", "")
    book = openpyxl.load_workbook(path)
    book.active["K1"] = "status"
    book.save(path)
    with pytest.raises(ValueError, match="Renomeie 'status' para 'situacao'"):
        read_mappings(path)


def test_a_report_without_the_mappings_sheet_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "sem_aba.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Outra"
    book.save(path)
    with pytest.raises(ValueError, match="não tem a aba 'Mapeamentos'"):
        read_mappings(path)


@requires_workbooks
def test_datasul_counts_recorded_winners_even_when_the_audit_proof_is_ambiguous():
    client_file, report_file, key_file = KEYED["datasul"]
    score = score_report(ROOT / report_file, ROOT / client_file,
                         load_answer_key(ROOT / key_file, "datasul"))
    rows = {r.model_column: r for r in score.accepted}
    assert rows[59].decided_by == "identifier"
    assert rows[184].decided_by == "arithmetic"
    assert rows[59].objective_decision is True
    assert rows[184].objective_decision is True


@requires_workbooks
def test_cli_prints_portuguese_figures(capsys: pytest.CaptureFixture[str]) -> None:
    assert evaluate.main(["--root", str(ROOT), "--client", "oracle", "--show-losses"]) == 0
    text = capsys.readouterr().out
    assert "TOTAL COM GABARITO  encontrados" in text
    assert "aceites decididos por evidência objetiva" in text
    # Porcentagens e confianças com vírgula decimal, nunca com ponto.
    assert re.search(r"\d,\d%", text)
    assert not re.search(r"\d\.\d+%|final=\d\.\d", text)


def test_percentages_use_a_decimal_comma_and_keep_their_width() -> None:
    assert evaluate._pct(0.8) == " 80,0%"
    assert evaluate._pct(1.0) == "100,0%"
    assert evaluate._pct(0.0) == "  0,0%"
    assert evaluate._pct(None) == "  n/d "
    assert evaluate._decimal(0.8999) == "0,8999"
    assert evaluate._decimal(None) == "n/d"


def test_every_outcome_has_a_portuguese_label_that_fits_its_column() -> None:
    assert set(evaluate.OUTCOME_LABELS) == set(Outcome)
    assert evaluate.OUTCOME_LABELS[Outcome.WRONG_COLUMN] == "coluna_errada"
    assert evaluate.OUTCOME_LABELS[Outcome.MISSED] == "perdido"
    # --show-losses só lista perdidas e colunas erradas, numa coluna de 14
    # caracteres alinhada com "retido".
    assert all(
        len(evaluate.OUTCOME_LABELS[outcome]) <= 14
        for outcome in (Outcome.MISSED, Outcome.WRONG_COLUMN)
    )


def test_help_is_in_portuguese(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exit_info:
        evaluate.main(["--help"])
    assert exit_info.value.code == 0
    text = capsys.readouterr().out
    assert text.startswith("uso: run.py evaluate")
    assert "opções:" in text
    assert "mostra esta mensagem de ajuda e sai" in text
    assert "--show-losses" in text and "--strict-key-only" in text


@requires_workbooks
def test_accepts_the_key_cannot_check_are_counted_and_visible() -> None:
    # O exit code lê só as fabricações; um aceite fora do gabarito sai
    # com 0 e só aparece nesta contagem.
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        assert len(score.undetermined_accepts) == (
            len(score.accepted)
            - sum(1 for r in score.accepted if r.outcome is not Outcome.UNDETERMINED)
        )
        good = sum(1 for r in score.accepted if r.outcome is Outcome.CORRECT)
        assert good + len(score.undetermined_accepts) <= len(score.accepted)


@pytest.mark.parametrize("winner, expected", [("arithmetic", True), ("identifier", True),
    ("profile", True), ("catalog", True), ("lexical", False), ("semantic", False), (None, None)])
def test_objective_decision_uses_the_winner_and_never_the_explanation(winner, expected):
    for proof in ["ela não distingue esta coluna", "ela não permite escolher esta coluna", ""]:
        row = Row(1, "Target", 1, "accepted", "data", .95, Outcome.CORRECT,
                  arithmetic_proof=proof, decided_by=winner)
        assert row.objective_decision is expected


def test_missing_historical_winners_remain_unknown_in_metrics():
    score = Score("synthetic", rows=[
        Row(1, "Unknown", 1, "accepted", "catalog", .99, Outcome.CORRECT),
        Row(2, "Known", 2, "accepted", "data", .92, Outcome.CORRECT, decided_by="identifier"),
    ])
    assert score.objective_decision_share == 1.0
    assert decision_census(score) == {"unrecorded": 1, "identifier": 1}
    score.rows = score.rows[:1]
    assert score.objective_decision_share is None


def test_a_human_correction_does_not_inherit_the_old_winning_method(tmp_path):
    path = _synthetic_report(tmp_path / "mapping.xlsx", "")
    book = openpyxl.load_workbook(path)
    sheet = book.active
    col = sheet.max_column + 1
    sheet.cell(1, col, "parecer_revisor")
    sheet.cell(1, col + 1, "coluna_cliente_correta")
    sheet.cell(2, col, "incorreto")
    sheet.cell(2, col + 1, 2)
    book.save(path)
    score = score_report(path, _synthetic_client(tmp_path / "client.xlsx"), AnswerKey("synthetic", {}))
    assert score.rows[0].decided_by is None
    assert score.rows[0].objective_decision is None


def test_evaluation_is_invariant_to_report_row_order(tmp_path):
    rows = [dict(model_column_number=1, model_column="Target A", status="accepted"),
            dict(model_column_number=2, model_column="Target B", client_column_number=2,
                 client_column="Source B", status="accepted")]
    client, report = report_files(tmp_path, rows)
    key = AnswerKey("synthetic", {1: 1, 2: 2})
    before = score_report(report, client, key)
    report_files(tmp_path, rows[::-1])
    after = score_report(report, client, key)
    assert asdict(before) == asdict(after)
    assert after.delivered_correct == 2


@pytest.mark.parametrize("numbers", [[1, 1], [1, 0], [1, 1.5], [1, None], [1, "invalid"], [1, 3]])
def test_evaluation_rejects_invalid_field_identity(tmp_path, numbers):
    client, report = report_files(tmp_path, [
        dict(model_column_number=n, model_column=f"Target {i}") for i, n in enumerate(numbers)
    ])
    with pytest.raises(ValueError, match="numero_coluna_modelo"):
        score_report(report, client, AnswerKey("synthetic", {}))


def test_reviewer_cannot_confirm_the_same_source_for_two_fields(tmp_path):
    client, report = report_files(tmp_path, [
        dict(model_column_number=1, model_column="Target A", reviewer_status="correct"),
        dict(model_column_number=2, model_column="Target B", reviewer_status="correct"),
    ])
    with pytest.raises(ValueError, match="mais de um destino"):
        score_report(report, client, AnswerKey("synthetic", {}))


def test_answer_key_reuse_is_flagged_without_changing_the_key():
    for name, source, targets in [("datasul", 8, [39, 147]), ("oracle", 23, [100, 190])]:
        key = load_answer_key(f"data/keys/de_para_{name}.xlsx", name)
        assert key.source_conflicts == {source: targets}
        assert all(key.present[target] == source for target in targets)
    assert AnswerKey("synthetic", {1: 1, 2: 2}).source_conflicts == {}
