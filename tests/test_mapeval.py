"""Testes do mapeval: as invariantes que tornam os números confiáveis e a linha de base."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pytest

import report_schema
from mapeval import cli
from mapeval.excel import (
    column_letter_to_index,
    index_to_column_letter,
    parse_coluna_annotation,
)
from mapeval.outcomes import Outcome, classify
from mapeval.score import Row, read_mappings, score_report
from mapeval.truth import KEY_CORRECTIONS, AnswerKey, load_answer_key

ROOT = Path(__file__).resolve().parent.parent
KEYED = {
    "datasul": (
        "data/clients/datasul_maio.xlsx",
        "data/reports/mapping_datasul_full.xlsx",
        "data/keys/de_para_datasul.xlsx",
    ),
    "oracle": (
        "data/clients/oracle.xlsx",
        "data/reports/mapping_oracle_full.xlsx",
        "data/keys/de_para_oracle.xlsx",
    ),
}
requires_workbooks = pytest.mark.skipif(
    not (ROOT / "data/keys/de_para_oracle.xlsx").exists(),
    reason="planilhas dos clientes ausentes",
)

#: Prova não decisiva, com o marcador PROOF_NOT_DECISIVE_PT.
PORTUGUESE_NON_DECISIVE_PROOF = (
    "valor = base x alíquota se confirma para PIS em 5501 linhas, com aderência "
    "0,9996 - mas este cliente tem 2 blocos disjuntos de PIS que satisfazem a "
    "equação, portanto ela não distingue esta coluna"
)


def _accepted_row(proof: str, method: str = "data") -> Row:
    return Row(model_column=213, model_name="Aliquota PIS", client_column=96,
               status="accepted", method=method, final_confidence=0.98,
               outcome=Outcome.CORRECT, arithmetic_proof=proof)


@pytest.mark.parametrize(
    ("letters", "index"),
    [("A", 1), ("Z", 26), ("AA", 27), ("AZ", 52), ("BA", 53), ("BV", 74), ("FY", 181), ("GZ", 208)],
)
def test_column_letters_round_trip(letters: str, index: int) -> None:
    assert column_letter_to_index(letters) == index
    assert index_to_column_letter(index) == letters


def test_annotation_tolerates_the_trailing_space_in_the_real_key() -> None:
    # Está assim na linha 1 de de_para_datasul.xlsx.
    assert parse_coluna_annotation("Coluna BV ") == 74


def test_annotation_ignores_cells_that_are_not_annotations() -> None:
    assert parse_coluna_annotation(None) is None
    assert parse_coluna_annotation("Base CBS") is None


def test_a_delivered_vacuous_target_is_fabrication_even_when_the_key_agrees() -> None:
    assert classify(
        chosen_client_column=5,
        expected_client_column=5,
        chosen_is_vacuous=True,
        delivered=True,
    ) is Outcome.FABRICATED


def test_an_undelivered_vacuous_proposal_is_not_a_fabrication() -> None:
    # O apply não preenche uma linha em revisão, então ela não chega ao cliente.
    assert classify(
        chosen_client_column=5,
        expected_client_column=5,
        chosen_is_vacuous=True,
        delivered=False,
    ) is Outcome.CORRECT


def test_unkeyed_columns_are_never_credited() -> None:
    for chosen in (None, 5):
        assert classify(
            chosen_client_column=chosen,
            expected_client_column=None,
            chosen_is_vacuous=False,
        ) is Outcome.UNDETERMINED


def test_keyed_columns_are_correct_missed_or_wrong() -> None:
    def outcome(chosen: int | None) -> Outcome:
        return classify(
            chosen_client_column=chosen, expected_client_column=5, chosen_is_vacuous=False
        )

    assert outcome(5) is Outcome.CORRECT
    assert outcome(None) is Outcome.MISSED
    assert outcome(6) is Outcome.WRONG_COLUMN


@requires_workbooks
@pytest.mark.parametrize(("client", "expected"), [("datasul", 30), ("oracle", 26)])
def test_recall_denominator_is_the_present_set_never_227(client: str, expected: int) -> None:
    key = load_answer_key(ROOT / KEYED[client][2], client)
    assert key.denominator == expected != 227


@requires_workbooks
def test_harness_reproduces_the_corrected_baseline() -> None:
    total = 0
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        total += score.count(Outcome.CORRECT)
    assert total == 43


@requires_workbooks
def test_the_four_key_corrections_are_worth_exactly_four_columns() -> None:
    strict = corrected = 0
    for client, (client_file, report_file, key_file) in KEYED.items():
        for apply_corrections, bucket in ((False, "strict"), (True, "corrected")):
            key = load_answer_key(ROOT / key_file, client, apply_corrections=apply_corrections)
            score = score_report(ROOT / report_file, ROOT / client_file, key)
            if bucket == "strict":
                strict += score.count(Outcome.CORRECT)
            else:
                corrected += score.count(Outcome.CORRECT)
    # A diferença é a invariante; os totais mudam quando os relatórios mudam.
    assert corrected - strict == 4
    assert (strict, corrected) == (39, 43)
    assert sum(len(v) for v in KEY_CORRECTIONS.values()) == 4


@requires_workbooks
@pytest.mark.parametrize(("client", "report"), [
    ("datasul", "data/reports/mapping_datasul_full.xlsx"),
    ("oracle", "data/reports/mapping_oracle_full.xlsx"),
])
def test_no_report_delivers_a_column_with_no_data(client: str, report: str) -> None:
    key = load_answer_key(ROOT / KEYED[client][2], client)
    score = score_report(ROOT / report, ROOT / KEYED[client][0], key)
    assert [
        (row.model_column, row.client_column) for row in score.vacuous_targets
    ] == []


@requires_workbooks
def test_delivered_recall_is_never_above_picking_recall() -> None:
    # A entrega é um subconjunto da escolha.
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        assert score.delivered_correct <= score.count(Outcome.CORRECT)
        assert len(score.withheld_correct) == (
            score.count(Outcome.CORRECT) - score.delivered_correct
        )


@requires_workbooks
def test_the_accept_gate_withholds_eleven_correct_mappings() -> None:
    # A maior perda isolada do pipeline: corretos encontrados e não entregues.
    found = delivered = 0
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        found += score.count(Outcome.CORRECT)
        delivered += score.delivered_correct
    assert found - delivered == 11
    assert (found, delivered) == (43, 32)


@requires_workbooks
def test_every_withheld_correct_mapping_sits_below_the_accept_threshold() -> None:
    # Se uma linha retida passar de 0,90, é outra recusa, e não o limiar, que a segura.
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        for row in score.withheld_correct:
            assert (row.final_confidence or 0.0) < 0.90, (
                f"{client} m{row.model_column} retida com {row.final_confidence}"
            )


def test_a_non_discriminating_proof_is_not_independent_evidence() -> None:
    # Uma prova que dois blocos disjuntos do mesmo tributo satisfazem traz o
    # marcador "não distingue esta coluna" e não conta como evidência.
    assert _accepted_row(
        "valor = base x alíquota se confirma para PIS em 500 linhas, com aderência 0,99"
    ).header_independent
    assert report_schema.PROOF_NOT_DECISIVE_PT in PORTUGUESE_NON_DECISIVE_PROOF
    assert not _accepted_row(PORTUGUESE_NON_DECISIVE_PROOF).header_independent
    # Um aceite por catálogo continua independente, com ou sem prova decisiva.
    assert _accepted_row(PORTUGUESE_NON_DECISIVE_PROOF, method="catalog").header_independent


def _synthetic_client(path: Path) -> Path:
    """Exportação mínima: c1 e c3 com dados, c2 sem nenhuma linha preenchida."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Base PIS", "Coluna vazia", "CFOP"])
    sheet.append([100.0, None, "5102"])
    sheet.append([200.0, None, "6102"])
    book.save(path)
    return path


#: Linhas de mapeamento no vocabulário interno.
_SYNTHETIC_ROWS = [
    {"model_column_number": 1, "model_column": "Base PIS", "client_column_number": 1,
     "client_column": "Base PIS", "final_confidence": 0.95, "method": "name",
     "status": "accepted", "arithmetic_proof": None},
    {"model_column_number": 2, "model_column": "CFOP", "client_column_number": 3,
     "client_column": "CFOP", "final_confidence": 0.93, "method": "catalog",
     "status": "accepted", "arithmetic_proof": None},
    {"model_column_number": 3, "model_column": "Isentas PIS", "client_column_number": 2,
     "client_column": "Coluna vazia", "final_confidence": 0.5, "method": "data",
     "status": "review", "arithmetic_proof": None},
    {"model_column_number": 4, "model_column": "Base COFINS", "client_column_number": None,
     "client_column": None, "final_confidence": 0.2, "method": "data",
     "status": "no_match", "arithmetic_proof": None},
]


def _synthetic_report(path: Path, proof: str) -> Path:
    """Relatório mínimo em português; a linha m1 leva a prova recebida."""
    fields = list(_SYNTHETIC_ROWS[0])
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = report_schema.sheet_name("mappings")
    sheet.append([report_schema.header_label("mappings", field) for field in fields])
    for number, row in enumerate(_SYNTHETIC_ROWS, start=1):
        values = dict(row, arithmetic_proof=proof if number == 1 else None)
        values = {f: report_schema.external_value("mappings", f, v) for f, v in values.items()}
        sheet.append([values[field] for field in fields])
    book.save(path)
    return path


def test_a_report_scores_without_the_client_workbooks(tmp_path: Path) -> None:
    # Relatório em português; a pontuação volta ao vocabulário interno.
    client = _synthetic_client(tmp_path / "cliente.xlsx")
    key = AnswerKey(client="sintetico", present={1: 1, 2: 3})
    portuguese = score_report(
        _synthetic_report(tmp_path / "pt.xlsx", PORTUGUESE_NON_DECISIVE_PROOF), client, key
    )
    assert [r.status for r in portuguese.rows] == ["accepted", "accepted", "review", "no_match"]
    assert [r.method for r in portuguese.rows] == ["name", "catalog", "data", "data"]
    assert [r.outcome for r in portuguese.rows] == [
        Outcome.CORRECT, Outcome.CORRECT, Outcome.UNDETERMINED, Outcome.UNDETERMINED,
    ]
    assert portuguese.delivered_correct == 2
    assert portuguese.correct_pairs == {(1, 1), (2, 3)}
    # m1 tem só prova não decisiva; m2 é aceite por catálogo.
    assert portuguese.header_independent_share == 0.5
    # m3 propõe a coluna vazia, mas está em revisão: não é fabricação.
    assert [(r.model_column, r.client_column) for r in portuguese.vacuous_proposals] == [(3, 2)]
    assert portuguese.vacuous_targets == []


def test_a_half_translated_report_is_refused(tmp_path: Path) -> None:
    # Cabeçalhos em português numa aba com o nome antigo, ou um cabeçalho em
    # inglês numa aba em português.
    path = _synthetic_report(tmp_path / "pt.xlsx", "")
    book = openpyxl.load_workbook(path)
    book.active.title = report_schema.LEGACY_SHEETS["mappings"]
    book.save(path)
    with pytest.raises(ValueError, match="formato antigo"):
        read_mappings(path)

    path = _synthetic_report(tmp_path / "pt.xlsx", "")
    book = openpyxl.load_workbook(path)
    book.active["K1"] = "status"
    book.save(path)
    with pytest.raises(ValueError, match="Renomeie 'status' para 'situacao'"):
        read_mappings(path)


def test_a_report_without_the_mappings_sheet_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "sem_aba.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Outra"
    book.save(path)
    with pytest.raises(ValueError, match="não tem a aba 'Mapeamentos'"):
        read_mappings(path)


@requires_workbooks
def test_the_marker_keeps_datasul_non_decisive_proofs_out_of_the_metric() -> None:
    client_file, report_file, key_file = KEYED["datasul"]
    score = score_report(
        ROOT / report_file, ROOT / client_file, load_answer_key(ROOT / key_file, "datasul")
    )
    marked = [
        r for r in score.accepted
        if r.arithmetic_proof and report_schema.PROOF_NOT_DECISIVE_PT in r.arithmetic_proof
    ]
    assert marked
    assert all(r.method == "catalog" or not r.header_independent for r in marked)


@requires_workbooks
def test_cli_prints_portuguese_figures(capsys: pytest.CaptureFixture[str]) -> None:
    assert cli.main(["--root", str(ROOT), "--client", "oracle", "--show-losses"]) == 0
    text = capsys.readouterr().out
    assert "TOTAL COM GABARITO  encontrados" in text
    assert "aceites independentes de cabeçalho" in text
    # Porcentagens e confianças com vírgula decimal, nunca com ponto.
    assert re.search(r"\d,\d%", text)
    assert not re.search(r"\d\.\d+%|final=\d\.\d", text)


def test_percentages_use_a_decimal_comma_and_keep_their_width() -> None:
    assert cli._pct(0.8) == " 80,0%"
    assert cli._pct(1.0) == "100,0%"
    assert cli._pct(0.0) == "  0,0%"
    assert cli._pct(None) == "  n/d "
    assert cli._decimal(0.8999) == "0,8999"
    assert cli._decimal(None) == "n/d"


def test_every_outcome_has_a_portuguese_label_that_fits_its_column() -> None:
    assert set(cli.OUTCOME_LABELS) == set(Outcome)
    assert cli.OUTCOME_LABELS[Outcome.WRONG_COLUMN] == "coluna_errada"
    assert cli.OUTCOME_LABELS[Outcome.MISSED] == "perdido"
    # --show-losses só lista perdidas e colunas erradas, numa coluna de 14
    # caracteres alinhada com "retido".
    assert all(
        len(cli.OUTCOME_LABELS[outcome]) <= 14
        for outcome in (Outcome.MISSED, Outcome.WRONG_COLUMN)
    )


def test_help_is_in_portuguese(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exit_info:
        cli.main(["--help"])
    assert exit_info.value.code == 0
    text = capsys.readouterr().out
    assert text.startswith("uso: python -m mapeval")
    assert "opções:" in text
    assert "mostra esta mensagem de ajuda e sai" in text
    assert "--show-losses" in text and "--strict-key-only" in text


@requires_workbooks
def test_accepts_the_key_cannot_check_are_counted_and_visible() -> None:
    # O código de saída lê só as fabricações; um aceite fora do gabarito sai
    # com 0 e só aparece nesta contagem.
    for client, (client_file, report_file, key_file) in KEYED.items():
        key = load_answer_key(ROOT / key_file, client)
        score = score_report(ROOT / report_file, ROOT / client_file, key)
        assert len(score.undetermined_accepts) == (
            len(score.accepted)
            - sum(1 for r in score.accepted if r.outcome is not Outcome.UNDETERMINED)
        )
        good = sum(1 for r in score.accepted if r.outcome is Outcome.CORRECT)
        assert good + len(score.undetermined_accepts) <= len(score.accepted)
