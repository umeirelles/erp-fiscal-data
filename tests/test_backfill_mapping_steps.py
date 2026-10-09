"""Testes de backfill_mapping_steps.py, o preenchimento retroativo das etapas.

A ferramenta só se aplica a relatórios legados em inglês, gravados antes das
colunas de etapas (``decided_by``, ``supported_by``, ``vetoed_by``). Estes
testes fixam três coisas:

- ela recusa, sem gravar nada, um relatório no esquema em português ou de
  esquema misto;
- ela continua completando um relatório legado exatamente como antes da
  tradução, com os valores legados em inglês;
- ela continua recusando um relatório legado que já foi completado.

Os valores esperados do relatório legado mínimo foram obtidos rodando a versão
anterior da ferramenta (commit 2ae4b6d) sobre este mesmo fixture.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

import backfill_mapping_steps
import report_schema
from backfill_mapping_steps import NEW_COLUMNS, UNRECORDED, backfill, main

#: Cabeçalho da aba Mappings de um relatório legado anterior às colunas de
#: etapas, no layout de sap_procosa: sem ``arithmetic_proof``, então as
#: colunas novas entram antes de ``reviewer_status`` e todo campo que a
#: ferramenta lê fica à esquerda da âncora (o defeito conhecido 1.1 de
#: docs/KNOWN_ISSUES.md não interfere aqui).
LEGACY_HEADER = [
    "model_column_number",
    "model_column",
    "client_column_number",
    "client_column",
    "name_confidence",
    "name_runner_up_confidence",
    "name_margin",
    "data_confidence",
    "final_confidence",
    "method",
    "status",
    "reason",
    "catalog_name",
    "catalog_coverage",
    "catalog_confidence",
    "llm_data_confidence",
    "llm_final_confidence",
    "reviewer_status",
    "correct_client_column",
    "reviewer_note",
]

#: Onde as colunas de etapas entram: antes de reviewer_status.
ANCHOR = LEGACY_HEADER.index("reviewer_status")

#: Exportação do cliente: um trio de ICMS que fecha valor = base x alíquota
#: em 60 linhas, um CFOP, uma coluna totalmente vazia, um valor de IPI e um
#: campo sem papel fiscal reconhecido.
CLIENT_HEADER = ["ICMS Base", "ICMS Taxa", "ICMS Valor", "CFOP", "Vazia", "Valor IPI", "Campo Livre"]


def _row(
    number: int,
    model: str,
    client: int | None,
    client_name: str | None,
    name: float,
    data: float,
    method: str,
    status: str,
    *,
    catalog: str | None = None,
    catalog_confidence: float | None = None,
    llm: float | None = None,
    reviewer: str | None = None,
    note: str | None = None,
) -> dict[str, object]:
    """Uma linha da aba Mappings legada, com os nomes de campo em inglês."""
    return {
        "model_column_number": number,
        "model_column": model,
        "client_column_number": client,
        "client_column": client_name,
        "name_confidence": name,
        "name_runner_up_confidence": 0.1,
        "name_margin": 0.1,
        "data_confidence": data,
        "final_confidence": 0.9,
        "method": method,
        "status": status,
        "reason": "justificativa qualquer",
        "catalog_name": catalog,
        "catalog_coverage": catalog_confidence,
        "catalog_confidence": catalog_confidence,
        "llm_data_confidence": llm,
        "llm_final_confidence": llm,
        "reviewer_status": reviewer,
        "correct_client_column": None,
        "reviewer_note": note,
    }


#: Linha do relatório legado -> (decided_by, supported_by, vetoed_by) esperados.
#: O relatório é de uma execução com embeddings, então o cabeçalho conta como
#: ``semantic``. ``None`` nos três: a linha não tem proposta válida e não é tocada.
CASES: list[tuple[dict[str, object], tuple[str | None, str | None, str | None]]] = [
    # O trio do cliente prova a base: a aritmética decide.
    (
        _row(1, "Base ICMS", 1, "ICMS Base", 0.95, 0.80, "name", "accepted"),
        ("arithmetic", "semantic + arithmetic", None),
    ),
    # O resolvedor de IA rodou (llm_final_confidence > 0) e entra como apoio.
    (
        _row(
            2, "Valor ICMS", 3, "ICMS Valor", 0.90, 0.80, "llm", "accepted",
            llm=0.92, reviewer="correct", note="conferido pelo analista",
        ),
        ("arithmetic", "resolver + semantic + arithmetic", None),
    ),
    # method == "catalog": o catálogo decidiu.
    (
        _row(
            3, "CFOP", 4, "CFOP", 0.95, 0.97, "catalog", "accepted",
            catalog="cfop", catalog_confidence=0.97,
        ),
        ("catalog", "semantic + catalog", None),
    ),
    # Catálogo abaixo da barra de distinção: só concorda.
    (
        _row(
            4, "CST ICMS", 7, "Campo Livre", 0.50, 0.50, "name", "review",
            catalog="cst_icms", catalog_confidence=0.60,
        ),
        ("semantic", "semantic + catalog (agrees)", None),
    ),
    # Pontuação de dados acima do teto do perfil genérico sem explicação.
    (
        _row(5, "Outro", 7, "Campo Livre", 0.0, 0.95, "data", "review"),
        (UNRECORDED, UNRECORDED, None),
    ),
    # Em revisão, sobre uma coluna sem nenhum valor: trava de coluna vazia.
    (
        _row(6, "Frete", 5, "Vazia", 0.30, 0.10, "name", "review"),
        ("semantic", "semantic", "vacuity"),
    ),
    # Em revisão, ICMS proposto numa coluna de IPI: trava de conflito tributário.
    (
        _row(7, "Valor do ICMS", 6, "Valor IPI", 0.60, 0.70, "name", "review"),
        ("semantic", "semantic", "family_conflict"),
    ),
    # Sem coluna proposta.
    (
        _row(8, "Desconto", None, None, 0.0, 0.0, "none", "no_match"),
        (None, None, None),
    ),
    # Número de coluna fora da exportação.
    (
        _row(9, "Seguro", 99, "Fora", 0.4, 0.4, "name", "review"),
        (None, None, None),
    ),
]


@pytest.fixture
def client_path(tmp_path: Path) -> Path:
    path = tmp_path / "exportacao_cliente.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(CLIENT_HEADER)
    for i in range(60):
        base = 1000.0 + i
        sheet.append([base, 18.0, round(base * 0.18, 2), 5102, None, 10.0 + i, f"item {i}"])
    book.save(path)
    return path


def _write_legacy_report(
    path: Path,
    header: list[str] = LEGACY_HEADER,
    *,
    metadata_sheet: str = "Run Metadata",
) -> Path:
    """Grava um relatório legado mínimo: abas Mappings e Run Metadata."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Mappings"
    sheet.append(header)
    for row, _ in CASES:
        sheet.append([row.get(name) for name in header])
    metadata = book.create_sheet(metadata_sheet)
    metadata.append(["key", "value"])
    metadata.append(["mapping_schema_version", 2])
    metadata.append(["embedding_model", "text-embedding-3-small"])
    book.save(path)
    return path


def _write_portuguese_report(path: Path, *, with_step_columns: bool) -> Path:
    """Grava um relatório no esquema em português, como o mapeador grava hoje."""
    fields = [
        field
        for field in report_schema.COLUMNS["mappings"]
        if with_step_columns or field not in NEW_COLUMNS
    ]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = report_schema.sheet_name("mappings")
    sheet.append([report_schema.header_label("mappings", field, report_schema.PT_BR) for field in fields])
    row = {
        **CASES[0][0],
        "catalog_distinct_coverage": None,
        "arithmetic_proof": "valor = base x alíquota vale para ICMS em 60 linhas",
        "supported_by": "semantic + arithmetic",
        "decided_by": "arithmetic",
        "vetoed_by": "",
    }
    sheet.append([report_schema.external_value("mappings", field, row.get(field)) for field in fields])
    metadata = book.create_sheet(report_schema.sheet_name("run_metadata"))
    metadata.append(["chave", "valor"])
    metadata.append([report_schema.metadata_key_label("mapping_schema_version"), 3])
    metadata.append([report_schema.metadata_key_label("embedding_model"), "text-embedding-3-small"])
    book.save(path)
    return path


def _sheet_rows(path: Path, sheet_name: str) -> list[tuple[object, ...]]:
    book = openpyxl.load_workbook(path)
    try:
        return [tuple(row) for row in book[sheet_name].values]
    finally:
        book.close()


@pytest.mark.parametrize("with_step_columns", [True, False])
def test_a_portuguese_report_is_refused_and_left_byte_for_byte(
    tmp_path: Path, client_path: Path, capsys: pytest.CaptureFixture[str], with_step_columns: bool
) -> None:
    # Sem a checagem de esquema, um relatório em português não tem as colunas
    # legadas decided_by/supported_by/vetoed_by e ganharia três colunas
    # legadas vazias, gravadas no próprio arquivo, com saída 0.
    report = _write_portuguese_report(tmp_path / "relatorio_pt.xlsx", with_step_columns=with_step_columns)
    original = report.read_bytes()

    with pytest.raises(SystemExit) as refused:
        main([str(report), str(client_path)])
    message = str(refused.value.code)
    assert "relatorio_pt.xlsx está no esquema em português" in message
    assert "relatórios legados em inglês" in message
    assert "Nada foi gravado." in message
    assert report.read_bytes() == original
    # Recusou antes de imprimir o resumo.
    assert capsys.readouterr().out == ""

    # A recusa vem antes de ler a exportação do cliente: um caminho que não
    # existe não chega a ser aberto.
    with pytest.raises(SystemExit, match="esquema em português"):
        backfill(report, tmp_path / "exportacao_que_nao_existe.xlsx")
    assert report.read_bytes() == original


@pytest.mark.parametrize(
    "header, metadata_sheet",
    [
        # Cabeçalho legado com uma coluna já traduzida.
        (
            [("situacao" if name == "status" else name) for name in LEGACY_HEADER],
            "Run Metadata",
        ),
        # Cabeçalho legado ao lado da aba de metadados em português: sem a
        # recusa, a ferramenta não acharia "Run Metadata" e rotularia o
        # cabeçalho como lexical sem aviso.
        (LEGACY_HEADER, "Metadados da Execução"),
    ],
    ids=["coluna_traduzida", "aba_metadados_traduzida"],
)
def test_a_mixed_schema_report_is_refused_and_left_byte_for_byte(
    tmp_path: Path, client_path: Path, header: list[str], metadata_sheet: str
) -> None:
    report = _write_legacy_report(tmp_path / "relatorio_misto.xlsx", header, metadata_sheet=metadata_sheet)
    original = report.read_bytes()

    with pytest.raises(SystemExit) as refused:
        backfill(report, client_path)
    message = str(refused.value.code)
    assert message.startswith("relatorio_misto.xlsx: Esquema misto")
    assert message.endswith("Nada foi gravado.")
    assert report.read_bytes() == original


def test_a_minimal_legacy_report_is_backfilled_as_before(
    tmp_path: Path, client_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    report = _write_legacy_report(tmp_path / "relatorio_legado.xlsx")
    metadata_before = _sheet_rows(report, "Run Metadata")

    assert main([str(report), str(client_path)]) == 0
    assert capsys.readouterr().out.strip() == (
        f"{report}: 7 linhas com proposta, 2 com prova aritmética, "
        "1 com etapa não registrada (unrecorded), 2 com veto"
    )

    rows = _sheet_rows(report, "Mappings")
    # As colunas de etapas entram com os nomes legados, antes de reviewer_status.
    assert list(rows[0]) == LEGACY_HEADER[:ANCHOR] + list(NEW_COLUMNS) + LEGACY_HEADER[ANCHOR:]
    assert len(rows) == 1 + len(CASES)
    for written, (row, expected) in zip(rows[1:], CASES):
        cells = dict(zip(rows[0], written))
        # Uma trava vazia é gravada como texto vazio, que volta como None.
        assert tuple(cells[name] for name in NEW_COLUMNS) == expected, row["model_column"]
        # Nenhuma outra célula muda, nem as do revisor, deslocadas para a direita.
        assert {name: cells[name] for name in LEGACY_HEADER} == row
    assert _sheet_rows(report, "Run Metadata") == metadata_before


def test_a_dry_run_counts_the_same_rows_and_writes_nothing(
    tmp_path: Path, client_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    report = _write_legacy_report(tmp_path / "relatorio_legado.xlsx")
    original = report.read_bytes()

    assert main([str(report), str(client_path), "--dry-run"]) == 0
    assert capsys.readouterr().out.strip() == (
        f"{report}: 7 linhas com proposta, 2 com prova aritmética, "
        "1 com etapa não registrada (unrecorded), 2 com veto  (simulação, nada foi gravado)"
    )
    assert report.read_bytes() == original


def test_an_already_backfilled_legacy_report_is_still_refused(
    tmp_path: Path, client_path: Path
) -> None:
    report = _write_legacy_report(tmp_path / "relatorio_legado.xlsx")
    backfill(report, client_path)
    completed = report.read_bytes()

    with pytest.raises(SystemExit) as refused:
        main([str(report), str(client_path)])
    assert str(refused.value.code) == (
        "relatorio_legado.xlsx já tem as colunas de etapas. Nada foi gravado."
    )
    assert report.read_bytes() == completed


@pytest.mark.parametrize(
    "proof, decisive",
    [
        ("valor = base x alíquota vale para ICMS em 60 linhas", True),
        (
            "valor = base x aliquota holds for icms on 60 rows at fit 1.0 - but this client "
            f"has 2 disjoint icms blocks that each satisfy it, so this "
            f"{report_schema.PROOF_NOT_DECISIVE_EN} this column",
            False,
        ),
        (
            "valor = base x alíquota vale para ICMS em 60 linhas, mas o cliente tem 2 blocos "
            f"disjuntos de ICMS que a satisfazem, então a prova {report_schema.PROOF_NOT_DECISIVE_PT}",
            False,
        ),
    ],
    ids=["decisiva", "nao_decisiva_ingles", "nao_decisiva_portugues"],
)
def test_a_non_decisive_proof_is_not_credited_in_either_language(
    tmp_path: Path,
    client_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    proof: str,
    decisive: bool,
) -> None:
    # A frase da prova passa a ser escrita em português na origem; a prova que
    # não distingue a coluna continua sem crédito de aritmética, qualquer que
    # seja a língua da frase.
    monkeypatch.setattr(
        backfill_mapping_steps, "_trio_proof_text", lambda trio: None if trio is None else proof
    )
    report = _write_legacy_report(tmp_path / "relatorio_legado.xlsx")

    tally = backfill(report, client_path)

    rows = _sheet_rows(report, "Mappings")
    base_icms = dict(zip(rows[0], rows[1]))
    assert base_icms["model_column"] == "Base ICMS"
    if decisive:
        assert (base_icms["decided_by"], base_icms["supported_by"]) == (
            "arithmetic",
            "semantic + arithmetic",
        )
        assert tally["arithmetic"] == 2
    else:
        assert (base_icms["decided_by"], base_icms["supported_by"]) == ("semantic", "semantic")
        assert tally["arithmetic"] == 0


def test_the_help_is_in_portuguese(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as finished:
        main(["--help"])
    assert finished.value.code == 0
    help_text = capsys.readouterr().out
    assert help_text.startswith("uso: ")
    assert "RELATORIO" in help_text and "EXPORTACAO" in help_text
    assert "relatório de mapeamento legado" in help_text
    assert "mostra a contagem sem gravar nada" in help_text
    # A ajuda é a docstring do módulo, com os parágrafos preservados.
    assert "Acrescenta as etapas nomeadas da análise" in help_text
    assert "\n\nIsso importa por mais do que o custo." in help_text
