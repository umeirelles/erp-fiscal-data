"""Testes de ``mapper/report.py``: abas do relatório, colunas sem uso, fila e comando map."""

from __future__ import annotations

from dataclasses import asdict
import math
from unittest.mock import patch

import pandas as pd
import pytest

from mapper import report, schema
from mapper.decide import (
    REASON_CATALOG_DECISION,
    REASON_CONFLICT_REVIEW,
    REASON_OFFLINE,
    REASON_PASS_ONE_ACCEPTED,
    REASON_UNSUPPORTED_NO_MATCH,
    REASON_VACUITY,
    MapperConfig,
    MappingResult,
    _map_run,
    _start_run,
    map_columns,
)
from mapper.evidence import (
    _trio_proof_text,
    discover_client_tax_trios,
    summarize_values,
)
from mapper.headers import LexicalHeaderScorer, describe_header_conflict_pt
from mapper.report import (
    MAPPING_SCHEMA_VERSION,
    RESOLVER_DISAGREES,
    SIGNAL_CATALOG,
    SIGNAL_CHECKDIGIT,
    SIGNAL_EMPTY,
    SIGNAL_TRIO,
    WHY_HELD_BELOW_THRESHOLD,
    WHY_HELD_NAME_ONLY,
    WHY_HELD_VACUOUS,
    AdjudicationRow,
    UnmatchedColumn,
    _unmatched_signals,
    build_parser,
    main,
    summarize_adjudication_queue,
    summarize_unmatched_columns,
    write_results,
)
from mapper.resolver import DataMappingDecision
from mapper.spreadsheet import ColumnRef, SheetTable

from support import Scores, table, table_from_headers


def test_mapping_report_contains_reviewer_columns(tmp_path) -> None:
    model = table(["CFOP"], [["5102"]])
    client = table(["CFOP"], [["5102"]])
    result = map_columns(model, client, LexicalHeaderScorer(), data_resolver=None)
    output = tmp_path / "mapping.xlsx"

    write_results(result, output)

    columns = pd.read_excel(output, nrows=0).columns.tolist()
    assert columns[-3:] == [
        "parecer_revisor",
        "coluna_cliente_correta",
        "observacao_revisor",
    ]


def test_mapping_report_writes_run_metadata(tmp_path) -> None:
    model = table(["CFOP"], [["5102"]])
    client = table(["CFOP"], [["5102"]])
    result = map_columns(model, client, LexicalHeaderScorer(), data_resolver=None)
    output = tmp_path / "mapping.xlsx"

    write_results(
        result,
        output,
        metadata={
            "mapping_schema_version": MAPPING_SCHEMA_VERSION,
            "embedding_model": "offline",
            "reference_manifest_sha256": "disabled",
        },
    )

    workbook = pd.ExcelFile(output)
    assert workbook.sheet_names == ["Mapeamentos", "Metadados da Execução"]
    metadata = pd.read_excel(workbook, sheet_name="Metadados da Execução")
    assert metadata.to_dict("records") == [
        {"chave": "versao_esquema_mapeamento", "valor": 3},
        {"chave": "modelo_embedding", "valor": "offline"},
        {"chave": "sha256_manifesto_referencia", "valor": "desativado"},
    ]


def test_the_report_schema_version_is_the_portuguese_one() -> None:
    assert MAPPING_SCHEMA_VERSION == schema.SCHEMA_VERSION_PT == 3


def _unmatched(results: dict[int, MappingResult], client: SheetTable, model: SheetTable):
    """Colunas do cliente sem uso, sem catálogos e sem trios."""
    return summarize_unmatched_columns(
        model_table=model,
        client_table=client,
        results=results,
        similarities=LexicalHeaderScorer().scores(
            [c.name for c in model.columns], [c.name for c in client.columns]
        ),
        reference_catalogs=None,
        client_tax_trios={},
        sample_size=8,
    )


def test_unmatched_means_not_consumed_not_merely_unreferenced() -> None:
    # Uma coluna só proposta em revisão não chega ao cliente, então conta como sem uso.
    client = table(["Base ICMS", "Vlr Frete", "Ruido"], [[10, 20, 30]] * 40)
    model = table(["Base ICMS", "Frete"], [[10, 20]] * 40)
    results = {
        0: MappingResult(
            model_column_number=1, model_column="Base ICMS",
            client_column_number=1, client_column="Base ICMS",
            name_confidence=1.0, name_runner_up_confidence=0.0, name_margin=1.0,
            data_confidence=None, final_confidence=1.0, method="name",
            status="accepted", reason="",
        ),
        1: MappingResult(
            model_column_number=2, model_column="Frete",
            client_column_number=2, client_column="Vlr Frete",
            name_confidence=0.8, name_runner_up_confidence=0.0, name_margin=0.8,
            data_confidence=None, final_confidence=0.8, method="name",
            status="review", reason="",
        ),
    }
    rows = {row.client_column_number: row for row in _unmatched(results, client, model)}
    assert 1 not in rows                                   # consumida
    assert rows[2].disposition == "proposed_not_accepted"  # proposta, não entregue
    assert rows[3].disposition == "never_referenced"


def test_nothing_is_omitted_only_demoted() -> None:
    # Uma coluna que não apresenta correspondência continua listada, com o motivo.
    client = table(["Vazio", "Ruido"], [[None, 7]] * 40)
    model = table(["Frete"], [[1]] * 40)
    rows = _unmatched({}, client, model)
    assert {r.client_column_number for r in rows} == {1, 2}
    empty = next(r for r in rows if r.client_column_number == 1)
    assert empty.review_priority == "low"
    assert empty.signal_detail == SIGNAL_EMPTY == "Esta coluna não tem valores."


def test_signal_details_are_portuguese_with_a_decimal_comma() -> None:
    # Os ids de sinal ficam internos (traduzidos na gravação); o detalhe sai
    # em português, com rótulos fiscais legíveis e vírgula decimal.
    base = [1000.0 + index for index in range(60)]
    sheet = table(
        ["ICMS Base", "ICMS Taxa", "ICMS Valor"],
        [[base[i], 18.0, round(base[i] * 0.18, 2)] for i in range(60)],
    )
    trios = discover_client_tax_trios(sheet)
    slots = {
        int(trio[slot]): slot for trio in trios[("icms", ())] for slot in ("base", "rate", "value")
    }
    column = sheet.columns[1]
    values = sheet.series(column)
    signals, detail = _unmatched_signals(
        column=column,
        profile=summarize_values(values),
        values=values,
        reference_catalogs=None,
        trio_slots=slots,
    )
    assert signals == ["arithmetic"]
    assert detail == SIGNAL_TRIO.format(slot="alíquota")
    assert detail == "participa como alíquota de um trio comprovado de valor = base x alíquota"

    cnpjs = pd.Series(["11.222.333/0001-81"] * 7 + ["123"])
    signals, detail = _unmatched_signals(
        column=ColumnRef("column_001", 0, "CNPJ", 1),
        profile=summarize_values(cnpjs),
        values=cnpjs,
        reference_catalogs=None,
        trio_slots={},
    )
    assert signals == ["cnpj_checkdigit"]
    assert detail == SIGNAL_CHECKDIGIT.format(fraction="87,5%", kind="CNPJ")


def test_shortlist_leads_and_the_rest_follow_in_column_order() -> None:
    client = table(["A", "B", "C"], [[None, 5, None]] * 40)
    model = table(["Frete"], [[1]] * 40)
    rows = _unmatched({}, client, model)
    priorities = [r.review_priority for r in rows]
    assert priorities == sorted(priorities, key=lambda p: p != "shortlist")


def _queue_result(model_number, model_name, client_number, client_name,
                  final, margin=0.0, status="review", **kw):
    """Uma proposta mínima para a fila de revisão."""
    return MappingResult(
        model_column_number=model_number, model_column=model_name,
        client_column_number=client_number, client_column=client_name,
        name_confidence=final, name_runner_up_confidence=final - margin,
        name_margin=margin, data_confidence=None, final_confidence=final,
        method="data", status=status, reason="", **kw)


def test_the_queue_holds_only_proposals_the_gate_did_not_deliver() -> None:
    sheet = table(["A", "B"], [["x", "y"], ["x", "y"]])
    rows = summarize_adjudication_queue(
        client_table=sheet,
        results=[
            _queue_result(1, "One", 1, "A", 0.95, status="accepted"),
            _queue_result(2, "Two", 2, "B", 0.88),
            _queue_result(3, "Three", None, None, 0.0, status="no_match"),
        ],
        final_threshold=0.90, sample_size=8,
    )
    assert [r.model_column_number for r in rows] == [2]
    assert rows[0].gate_distance == pytest.approx(0.02)
    assert rows[0].why_held == WHY_HELD_BELOW_THRESHOLD == "confiança abaixo do threshold de aceitação"


def test_a_proposal_onto_an_empty_column_can_never_be_confirmed() -> None:
    # Confirmá-la entregaria uma coluna em branco, então ela vai para o fim da
    # fila, por melhor que seja o cabeçalho.
    sheet = table(["Full", "Empty"], [["x", None], ["y", None]])
    rows = summarize_adjudication_queue(
        client_table=sheet,
        results=[
            _queue_result(1, "Weak but real", 1, "Full", 0.70),
            _queue_result(2, "Strong but empty", 2, "Empty", 0.98),
        ],
        final_threshold=0.90, sample_size=8,
    )
    assert [r.model_column for r in rows] == ["Weak but real", "Strong but empty"]
    assert rows[-1].rows_with_value == 0
    assert rows[-1].why_held == WHY_HELD_VACUOUS


def test_margin_breaks_ties_only_inside_the_clamped_band() -> None:
    # No teto, a confiança deixa de ordenar as linhas e a margem desempata;
    # fora dele, a ordenação usa a confiança.
    sheet = table(["A", "B", "C"], [["x", "y", "z"], ["x", "y", "z"]])
    rows = summarize_adjudication_queue(
        client_table=sheet,
        results=[
            _queue_result(1, "Clamped, thin margin", 1, "A", 0.8999, margin=0.01),
            _queue_result(2, "Clamped, wide margin", 2, "B", 0.8999, margin=0.40),
            _queue_result(3, "Higher confidence", 3, "C", 0.8999 + 0.01, margin=0.0),
        ],
        final_threshold=0.90, sample_size=8,
    )
    assert [r.model_column_number for r in rows] == [3, 2, 1]


def test_the_queue_says_when_the_branch_rather_than_the_evidence_held_a_row() -> None:
    sheet = table(["A"], [["x"], ["x"]])
    rows = summarize_adjudication_queue(
        client_table=sheet,
        results=[_queue_result(1, "One", 1, "A", 0.8999)],
        final_threshold=0.90, sample_size=8,
    )
    assert rows[0].why_held == WHY_HELD_NAME_ONLY


def test_the_resolver_disagreement_uses_a_decimal_comma() -> None:
    sheet = table(["A"], [["x"], ["y"]])
    rows = summarize_adjudication_queue(
        client_table=sheet,
        results=[_queue_result(1, "One", 1, "A", 0.88, llm_final_confidence=0.75)],
        final_threshold=0.90, sample_size=8,
    )
    assert rows[0].resolver_disagrees == RESOLVER_DISAGREES.format(llm_final="0,75")
    assert rows[0].resolver_disagrees == "confiança informada pelo resolver de IA: 0,75"


# ---------------------------------------------------------------------------
# Gravação no schema em português e ida e volta pelo report_schema
# ---------------------------------------------------------------------------

#: Colunas do revisor, acrescentadas em branco ao fim da aba Mapeamentos.
REVIEWER_FIELDS = ("reviewer_status", "correct_client_column", "reviewer_note")


def _result(**fields) -> MappingResult:
    base = dict(
        name_runner_up_confidence=0.0, name_margin=0.0, data_confidence=None,
        final_confidence=0.0, method="data", status="review", reason="",
    )
    base.update(fields)
    return MappingResult(**base)


def _varied_results() -> list[MappingResult]:
    """Linhas que cobrem cada situação, método, veto, catálogo, etapa e forma de nome de exibição."""
    decisive = _trio_proof_text(
        {"tax_scope": "icms", "rows_tested": 5570, "fit_score": 0.998, "alternatives": 1}
    )
    rival = _trio_proof_text(
        {"tax_scope": "pis", "rows_tested": 70, "fit_score": 0.9857, "alternatives": 2}
    )
    return [
        _result(
            model_column_number=1, model_column="CFOP", client_column_number=3,
            client_column="Cod. Fiscal", name_confidence=0.61, final_confidence=0.9,
            data_confidence=0.8571, method="catalog", status="accepted",
            reason=REASON_CATALOG_DECISION.format(
                catalog="CFOP", distinct_coverage="85,7%", unique_formatted_count=7,
                coverage="85,7%",
            ),
            catalog_name="cfop", catalog_coverage=0.8571, catalog_distinct_coverage=0.8571,
            catalog_confidence=0.8571, supported_by="lexical + catalog", decided_by="catalog",
        ),
        _result(
            model_column_number=2, model_column="Valor ICMS", client_column_number=12,
            client_column="Valor ICMS (occurrence 2)", name_confidence=1.0,
            name_runner_up_confidence=0.5, name_margin=0.5, final_confidence=1.0,
            method="name", status="accepted", reason=REASON_PASS_ONE_ACCEPTED,
            arithmetic_proof=decisive, supported_by="lexical + arithmetic",
            decided_by="arithmetic",
        ),
        _result(
            model_column_number=3, model_column="Base PIS", client_column_number=4,
            client_column="Unnamed column 4", name_confidence=0.88,
            name_runner_up_confidence=0.8, name_margin=0.08, final_confidence=0.88,
            method="name_only", reason=REASON_OFFLINE, arithmetic_proof=rival,
            supported_by="lexical", decided_by="lexical",
        ),
        _result(
            model_column_number=4, model_column="Município", client_column_number=7,
            client_column="Cidade", name_confidence=0.7, final_confidence=0.89,
            data_confidence=0.95, reason="The resolver's own prose stays in English.",
            catalog_name="ibge_municipality", catalog_coverage=0.97, catalog_confidence=0.97,
            supported_by="resolver + semantic + profile + identifier", decided_by="identifier",
            llm_data_confidence=0.9, llm_final_confidence=0.8,
        ),
        _result(
            model_column_number=5, model_column="Nº FCI", client_column_number=9,
            client_column="Nº FCI (occurrence 2)", name_confidence=0.98, name_margin=0.48,
            final_confidence=0.98, data_confidence=0.0,
            reason=REASON_VACUITY.format(client_column="Nº FCI (ocorrência 2)"),
            supported_by="resolver + semantic", decided_by="semantic", vetoed_by="vacuity",
            llm_data_confidence=0.9, llm_final_confidence=0.95,
        ),
        _result(
            model_column_number=6, model_column="Tipo de Imposto CBS", client_column_number=2,
            client_column="CBS Base", name_confidence=0.89, final_confidence=0.95,
            data_confidence=0.95,
            reason=REASON_CONFLICT_REVIEW.format(
                header_conflict=describe_header_conflict_pt("Tipo de Imposto CBS", "CBS Base")
            ),
            supported_by="semantic + profile", decided_by="profile",
            vetoed_by="family_conflict", catalog_name="ibs_cbs_classification",
        ),
        _result(
            model_column_number=7, model_column="cClassTrib", client_column_number=5,
            client_column="Class", name_confidence=0.5, final_confidence=0.5,
            method="data_error", reason="A comparação de dados falhou: RuntimeError: falhou",
        ),
        _result(
            model_column_number=8, model_column="Campo Inexistente", client_column_number=None,
            client_column=None, name_confidence=0.0, method="data", status="no_match",
            reason=REASON_UNSUPPORTED_NO_MATCH,
        ),
    ]


def _varied_unmatched() -> list[UnmatchedColumn]:
    def row(number, name, disposition, priority, signals, detail):
        return UnmatchedColumn(
            client_column_number=number, client_column=name, disposition=disposition,
            review_priority=priority, rows_with_value=40, distinct_values=12,
            signals=signals, signal_detail=detail, closest_model_column="Base ICMS",
            closest_name_confidence=0.4567,
        )

    return [
        row(1, "Base ICMS (occurrence 2)", "never_referenced", "shortlist", "arithmetic,catalog",
            SIGNAL_TRIO.format(slot="base de cálculo") + "; "
            + SIGNAL_CATALOG.format(distinct_coverage="88,2%", unique_formatted_count=17,
                                    catalog="CFOP")),
        row(2, "CNPJ", "proposed_not_accepted", "shortlist", "cnpj_checkdigit",
            SIGNAL_CHECKDIGIT.format(fraction="97,5%", kind="CNPJ")),
        row(3, "Chave", "never_referenced", "shortlist", "nfe_key_checkdigit",
            SIGNAL_CHECKDIGIT.format(fraction="100,0%", kind="chave de acesso da NF-e")),
        row(4, "CPF", "never_referenced", "shortlist", "cpf_checkdigit",
            SIGNAL_CHECKDIGIT.format(fraction="90,0%", kind="CPF")),
        row(5, "Unnamed column 5", "never_referenced", "low", "", SIGNAL_EMPTY),
        row(6, "Data", "never_referenced", "shortlist", "date", "contém datas"),
    ]


def _varied_queue() -> list[AdjudicationRow]:
    def row(rank, number, name, client_number, client_name, why, disagrees="", proof=""):
        return AdjudicationRow(
            rank=rank, model_column_number=number, model_column=name,
            proposed_client_column_number=client_number, proposed_client_column=client_name,
            final_confidence=0.8999, gate_distance=0.0001, name_margin=0.05,
            rows_with_value=0 if why == WHY_HELD_VACUOUS else 40,
            resolver_disagrees=disagrees, arithmetic_proof=proof,
            client_sample_values="1 | 2 | 3", why_held=why,
        )

    return [
        row(1, 3, "Base PIS", 4, "Unnamed column 4", WHY_HELD_NAME_ONLY,
            proof=_varied_results()[2].arithmetic_proof),
        row(2, 4, "Município", 7, "Cidade", WHY_HELD_BELOW_THRESHOLD,
            disagrees=RESOLVER_DISAGREES.format(llm_final="0,75")),
        row(3, 5, "Nº FCI", 9, "Nº FCI (occurrence 2)", WHY_HELD_VACUOUS),
    ]


def _is_missing(value) -> bool:
    return value is None or value == "" or (isinstance(value, float) and math.isnan(value))


def _assert_same_cell(expected, actual, where) -> None:
    if _is_missing(expected) or _is_missing(actual):
        assert _is_missing(expected) and _is_missing(actual), (where, expected, actual)
    elif isinstance(expected, (int, float)) and not isinstance(expected, bool):
        assert math.isclose(float(expected), float(actual)), (where, expected, actual)
    else:
        assert expected == actual, (where, expected, actual)


def _assert_round_trip(rows, frame: pd.DataFrame, table_name: str, extra=()) -> None:
    """Verifica se gravar e ler os resultados preserva os valores de ``asdict``.

    Campos estruturados retornam à representação interna. Mensagens e valores
    numéricos devem ser preservados durante a conversão pelo ``report_schema``.
    """
    internal = schema.to_internal_frame(frame, table_name)
    expected = [asdict(row) for row in rows]
    fields = list(expected[0]) + list(extra)
    assert list(internal.columns) == fields
    assert len(internal) == len(expected)
    for position, record in enumerate(expected):
        for field in fields:
            _assert_same_cell(
                record.get(field, ""), internal[field].iloc[position], (table_name, position, field)
            )


def test_write_results_round_trips_through_report_schema(tmp_path) -> None:
    results = _varied_results()
    unmatched = _varied_unmatched()
    queue = _varied_queue()
    output = tmp_path / "relatorio.xlsx"

    write_results(
        results, output, metadata={"mapping_schema_version": MAPPING_SCHEMA_VERSION},
        unmatched=unmatched, queue=queue,
    )

    workbook = pd.ExcelFile(output)
    assert workbook.sheet_names == [
        "Mapeamentos", "Metadados da Execução", "Colunas do Cliente sem Uso", "Fila de Revisão",
    ]
    _assert_round_trip(
        results, pd.read_excel(workbook, sheet_name="Mapeamentos"), "mappings",
        extra=REVIEWER_FIELDS,
    )
    _assert_round_trip(
        unmatched, pd.read_excel(workbook, sheet_name="Colunas do Cliente sem Uso"), "unmatched"
    )
    _assert_round_trip(queue, pd.read_excel(workbook, sheet_name="Fila de Revisão"), "queue")


def test_the_written_cells_are_in_portuguese(tmp_path) -> None:
    output = tmp_path / "relatorio.xlsx"
    write_results(
        _varied_results(), output, unmatched=_varied_unmatched(), queue=_varied_queue()
    )
    mappings = pd.read_excel(output, sheet_name="Mapeamentos")
    assert list(mappings.columns) == list(schema.COLUMNS["mappings"].values())
    assert list(mappings["situacao"]) == [
        "aceito", "aceito", "revisao", "revisao", "revisao", "revisao", "revisao",
        "sem_correspondencia",
    ]
    assert list(mappings["metodo"]) == [
        "catalogo", "nome", "somente_nome", "dados", "dados", "dados", "erro_dados", "dados",
    ]
    assert list(mappings["coluna_cliente"].fillna("")) == [
        "Cod. Fiscal", "Valor ICMS (ocorrência 2)", "Coluna sem nome 4", "Cidade",
        "Nº FCI (ocorrência 2)", "CBS Base", "Class", "",
    ]
    assert list(mappings["apoiado_por"].fillna("")) == [
        "lexico + catalogo", "lexico + aritmetica", "lexico",
        "resolvedor_ia + semantico + perfil + identificador", "resolvedor_ia + semantico",
        "semantico + perfil", "", "",
    ]
    assert list(mappings["decidido_por"].fillna("")) == [
        "catalogo", "aritmetica", "lexico", "identificador", "semantico", "perfil", "", "",
    ]
    assert list(mappings["vetado_por"].fillna("")) == [
        "", "", "", "", "coluna_vazia", "conflito_tributario", "", "",
    ]
    assert list(mappings["nome_catalogo"].fillna("")) == [
        "cfop", "", "", "municipio_ibge", "", "cclasstrib", "", "",
    ]
    # O texto da prova é preservado; leitores usam os campos de evidência.
    proofs = list(mappings["prova_aritmetica"].fillna(""))
    assert proofs == [result.arithmetic_proof or "" for result in _varied_results()]

    unmatched = pd.read_excel(output, sheet_name="Colunas do Cliente sem Uso")
    assert list(unmatched.columns) == list(schema.COLUMNS["unmatched"].values())
    assert list(unmatched["sinais"].fillna("")) == [
        "aritmetica,catalogo", "dv_cnpj", "dv_chave_nfe", "dv_cpf", "", "data",
    ]
    assert set(unmatched["disposicao"]) == {"nunca_referenciada", "proposta_nao_aceita"}
    assert set(unmatched["prioridade_revisao"]) == {"prioritaria", "baixa"}
    assert list(unmatched["coluna_cliente"])[0] == "Base ICMS (ocorrência 2)"
    assert list(unmatched["coluna_cliente"])[4] == "Coluna sem nome 5"

    queue = pd.read_excel(output, sheet_name="Fila de Revisão")
    assert list(queue.columns) == list(schema.COLUMNS["queue"].values())
    assert list(queue["coluna_cliente_proposta"]) == [
        "Coluna sem nome 4", "Cidade", "Nº FCI (ocorrência 2)",
    ]
    assert list(queue["motivo_retencao"]) == [
        WHY_HELD_NAME_ONLY, WHY_HELD_BELOW_THRESHOLD, WHY_HELD_VACUOUS,
    ]


def test_empty_auxiliary_sheets_keep_their_portuguese_headers(tmp_path) -> None:
    # Uma lista vazia ainda grava a aba: ela afirma que nada ficou de fora. A
    # linha em branco que leva os cabeçalhos sai com células vazias, então a
    # leitura devolve só o cabeçalho.
    output = tmp_path / "relatorio.xlsx"
    write_results(_varied_results(), output, unmatched=[], queue=[])

    workbook = pd.ExcelFile(output)
    assert workbook.sheet_names == [
        "Mapeamentos", "Colunas do Cliente sem Uso", "Fila de Revisão",
    ]
    for sheet, table_name in (
        ("Colunas do Cliente sem Uso", "unmatched"),
        ("Fila de Revisão", "queue"),
    ):
        frame = pd.read_excel(workbook, sheet_name=sheet)
        assert list(frame.columns) == list(schema.COLUMNS[table_name].values())
        assert frame.empty
        internal = schema.to_internal_frame(frame, table_name)
        assert list(internal.columns) == list(schema.COLUMNS[table_name])


def test_a_csv_report_uses_the_portuguese_schema_too(tmp_path) -> None:
    results = _varied_results()
    output = tmp_path / "relatorio.csv"

    write_results(results, output)

    frame = pd.read_csv(output)
    assert list(frame.columns) == list(schema.COLUMNS["mappings"].values())
    assert frame["situacao"].iloc[0] == "aceito"
    _assert_round_trip(results, frame, "mappings", extra=REVIEWER_FIELDS)


def test_an_unknown_output_extension_is_refused_in_portuguese(tmp_path) -> None:
    with pytest.raises(ValueError, match="A saída deve terminar em .xlsx ou .csv"):
        write_results(_varied_results(), tmp_path / "relatorio.txt")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _write_sheet(path, rows) -> None:
    pd.DataFrame(rows).to_excel(path, header=False, index=False)


def test_the_cli_help_is_in_portuguese() -> None:
    parser = build_parser()
    text = parser.format_help()
    assert text.startswith("uso: ")
    assert "argumentos posicionais:" in text
    assert "opções:" in text
    assert "mostra esta mensagem de ajuda e sai" in text
    for english in ("usage:", "positional arguments", "options:", "show this help"):
        assert english not in text, english
    # Toda opção tem ajuda, inclusive as que não tinham nenhuma.
    assert all(action.help for action in parser._actions)
    # As flags continuam em inglês.
    for flag in ("--offline", "--output", "--name-threshold", "--no-reference-data"):
        assert flag in text


def test_main_writes_a_portuguese_report(tmp_path, capsys) -> None:
    model = tmp_path / "modelo.xlsx"
    client = tmp_path / "cliente.xlsx"
    output = tmp_path / "relatorio.xlsx"
    _write_sheet(model, [["CFOP", "Valor Frete"], ["5102", 10.0]])
    _write_sheet(client, [["CFOP", "Frete"], ["5102", 10.0]])

    code = main(
        [str(client), "--model", str(model), "--output", str(output),
         "--offline", "--no-reference-data"]
    )

    assert code == 0
    out = capsys.readouterr().out
    assert f"Arquivo gravado: {output}" in out
    assert "aceito=1 revisao=1 sem_correspondencia=0" in out
    workbook = pd.ExcelFile(output)
    assert workbook.sheet_names == [
        "Mapeamentos", "Metadados da Execução", "Colunas do Cliente sem Uso", "Fila de Revisão",
    ]
    metadata = dict(
        pd.read_excel(workbook, sheet_name="Metadados da Execução")[["chave", "valor"]].values
    )
    assert metadata["versao_esquema_mapeamento"] == 3
    assert metadata["modelo_embedding"] == "offline"
    assert metadata["modelo_ia"] == "offline"
    assert metadata["sha256_manifesto_referencia"] == "desativado"
    assert set(metadata) == set(schema.METADATA_KEYS.values()) - {"migrado_da_versao_esquema"}


def test_main_rejects_a_threshold_out_of_range_in_portuguese(tmp_path) -> None:
    with pytest.raises(SystemExit) as stop:
        main([str(tmp_path / "cliente.xlsx"), "--name-threshold", "1.5", "--offline"])
    assert stop.value.code == "--name-threshold deve estar entre 0 e 1"


def test_main_reports_a_missing_file_in_portuguese_without_a_traceback(tmp_path) -> None:
    model = tmp_path / "modelo.xlsx"
    _write_sheet(model, [["CFOP"], ["5102"]])
    missing = tmp_path / "nao_existe.xlsx"

    with pytest.raises(SystemExit) as stop:
        main([str(missing), "--model", str(model), "--offline", "--no-reference-data"])

    # SystemExit com texto: o Python imprime a mensagem e sai com código 1.
    assert stop.value.code == f"Arquivo não encontrado: {missing}"


def test_main_reports_a_bad_header_row_in_portuguese(tmp_path) -> None:
    model = tmp_path / "modelo.xlsx"
    client = tmp_path / "cliente.xlsx"
    _write_sheet(model, [["CFOP"], ["5102"]])
    _write_sheet(client, [["CFOP"], ["5102"]])

    with pytest.raises(SystemExit) as stop:
        main(
            [str(client), "--model", str(model), "--client-header-row", "5",
             "--offline", "--no-reference-data", "--output", str(tmp_path / "r.xlsx")]
        )

    assert stop.value.code == (
        "Não foi possível gerar o relatório de mapeamento: A linha de cabeçalho 5 "
        f"está fora das linhas da planilha: {client}"
    )


def test_report_accounts_for_every_client_column_and_formats_confidence(tmp_path):
    import openpyxl
    model = table_from_headers(["Target A", "Target B", "Target C", "Target D"], [[1, 2, 3, 4]])
    client = table_from_headers(["Source A", "Source B", "Unused empty source"], [[11, 22, None]])
    scorer = Scores({"Target A": [.99, .4, 0], "Target B": [.4, .85, 0],
                     "Target C": [.4, .75, 0], "Target D": [.4, .7, 0]})
    results = map_columns(model, client, scorer, None)
    unused = summarize_unmatched_columns(
        model_table=model, client_table=client, results=dict(enumerate(results)),
        similarities=scorer.scores([c.name for c in model.columns], []),
        reference_catalogs=None, client_tax_trios={}, sample_size=8,
    )
    accepted = {r.client_column_number for r in results if r.status == "accepted"}
    unassigned = {r.client_column_number for r in unused}
    assert accepted.isdisjoint(unassigned)
    assert accepted | unassigned == {1, 2, 3}
    assert {r.client_column_number: r.disposition for r in unused} == {
        2: "proposed_not_accepted", 3: "never_referenced",
    }
    output = tmp_path / "report.xlsx"
    write_results(results, output, unmatched=unused, queue=[])
    workbook = openpyxl.load_workbook(output)
    try:
        sheet = workbook["Mapeamentos"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        confidence = sheet.cell(2, headers["confianca_final"])
        assert confidence.value == .99
        assert confidence.number_format == "0.00%"
        assert sheet.cell(2, headers["decidido_por"]).value == "semantico"
        assert workbook["Colunas do Cliente sem Uso"].max_row == 3
    finally:
        workbook.close()


def test_report_generation_reuses_prepared_client_profiles():
    from dataclasses import asdict
    model = table_from_headers(["Target A", "Target B"], [["model A", "model B"]])
    client = table_from_headers(["Source X", "Source Y"], [["source X", "source Y"]])
    class Resolver:
        def resolve(self, **kwargs):
            return DataMappingDecision(client_column_id="column_001", data_confidence=.5,
                                               final_confidence=.5, reason="Proposal")
    run = _start_run(model, client, Scores({"Target A": [.86, .7], "Target B": [.86, .7]}),
                            MapperConfig(), None, online=True)
    results = _map_run(run, Resolver())
    def reports(profiles=None):
        unused = summarize_unmatched_columns(
            model_table=model, client_table=client, results=run.results, similarities=run.similarities,
            reference_catalogs=None, client_tax_trios=run.client_tax_trios, sample_size=8,
            client_profiles=profiles,
        )
        queue = summarize_adjudication_queue(client_table=client, results=results,
                    final_threshold=.9, sample_size=8, client_profiles=profiles)
        return [asdict(row) for row in unused], [asdict(row) for row in queue]
    expected = reports()
    with patch.object(report, "summarize_values", side_effect=AssertionError("Profile recomputed")):
        assert reports(run.client_profiles) == expected
