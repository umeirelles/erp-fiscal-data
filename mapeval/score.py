"""Calcula as métricas de um relatório de mapeamento usando o gabarito do cliente.

O relatório está no schema em português. A leitura passa por
``report_schema``, que devolve os cabeçalhos com os nomes internos; a
situação e o método voltam ao vocabulário interno (``accepted``,
``catalog``...) antes de qualquer comparação.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import openpyxl

import report_schema
from mapping_review import normalize_review, resolve_mapping
from mapping_steps import TRANSFERRING_STEP_IDS
from spreadsheet import SheetTable, positive_column_number
from mapeval.outcomes import Outcome, classify
from mapeval.truth import AnswerKey

@dataclass
class Row:
    model_column: int
    model_name: str
    client_column: int | None
    #: Situação no vocabulário interno (``accepted``, ``review``, ``no_match``).
    status: str
    #: Método no vocabulário interno (``name``, ``data``, ``catalog``...).
    method: str
    final_confidence: float | None
    outcome: Outcome
    #: Relação aritmética nos dados do cliente que sustenta a escolha da coluna.
    arithmetic_proof: str | None = None

    #: Evidência que decidiu o aceite automático; ausente em relatórios antigos ou correções humanas.
    decided_by: str | None = None

    @property
    def objective_decision(self) -> bool | None:
        """Informa se o método vencedor usa dados; None significa método não registrado."""
        if self.decided_by in TRANSFERRING_STEP_IDS:
            return True
        if self.decided_by in {"lexical", "semantic"}:
            return False
        return None


@dataclass
class Score:
    client: str
    rows: list[Row] = field(default_factory=list)
    present: int = 0
    #: Mapeamentos aceitos para colunas vazias; fazem a avaliação retornar exit code 1.
    vacuous_targets: list[Row] = field(default_factory=list)
    #: Propostas para colunas vazias ainda não entregues; não alteram o exit code.
    vacuous_proposals: list[Row] = field(default_factory=list)

    def count(self, outcome: Outcome) -> int:
        return sum(1 for row in self.rows if row.outcome is outcome)

    @property
    def recall_on_present(self) -> float | None:
        """Calcula o recall usando apenas as colunas presentes no gabarito como denominador."""
        return self.count(Outcome.CORRECT) / self.present if self.present else None

    @property
    def accepted(self) -> list[Row]:
        return [r for r in self.rows if r.status == "accepted"]

    @property
    def delivered_correct(self) -> int:
        """Conta os acertos entregues após aplicar os pareceres do revisor."""
        return sum(
            1
            for r in self.rows
            if r.outcome is Outcome.CORRECT and r.status == "accepted"
        )

    @property
    def delivered_recall(self) -> float | None:
        """Calcula o recall dos mapeamentos corretos aceitos; não ultrapassa ``recall_on_present``."""
        return self.delivered_correct / self.present if self.present else None

    @property
    def withheld_correct(self) -> list[Row]:
        """Retorna as propostas corretas ainda não aceitas, por confiança decrescente.

        Essa ordenação ajuda o revisor a priorizar a análise.
        """
        held = [
            r
            for r in self.rows
            if r.outcome is Outcome.CORRECT and r.status != "accepted"
        ]
        return sorted(held, key=lambda r: -(r.final_confidence or 0.0))

    @property
    def accept_precision_lower(self) -> float | None:
        """Calcula o limite inferior de precisão considerando apenas os acertos confirmados.

        Mapeamentos fora do gabarito entram no denominador, mas não no numerador.
        Assim, uma nova escolha sem validação reduz o limite inferior."""
        accepts = self.accepted
        if not accepts:
            return None
        good = sum(1 for r in accepts if r.outcome is Outcome.CORRECT)
        return good / len(accepts)

    @property
    def undetermined_accepts(self) -> list[Row]:
        """Aceites em colunas do modelo fora do gabarito, que ele não consegue julgar.

        O exit code considera apenas fabricações. Uma escolha errada em coluna
        preenchida fora do gabarito pode passar pela avaliação com exit code 0;
        esta contagem permite acompanhar os casos sem validação.
        """
        return [r for r in self.accepted if r.outcome is Outcome.UNDETERMINED]

    @property
    def objective_decision_share(self) -> float | None:
        """Proporção de evidência objetiva entre aceites com método automático conhecido."""
        known = [r.objective_decision for r in self.accepted if r.objective_decision is not None]
        return sum(known) / len(known) if known else None

    @property
    def correct_pairs(self) -> set[tuple[int, int]]:
        """Os pares (coluna do modelo, coluna do cliente) corretos.

        Permite comparar os pares entre execuções. A mesma contagem pode esconder
        a perda de um acerto compensada por outro. O procedimento de
        docs/EVALUATION.md usa essa comparação.
        """
        return {
            (r.model_column, r.client_column)
            for r in self.rows
            if r.outcome is Outcome.CORRECT and r.client_column is not None
        }


@lru_cache(maxsize=None)
def _client_structure(client_path: str | Path) -> tuple[SheetTable, frozenset[int]]:
    """Lê os cabeçalhos e identifica as colunas vazias (a partir de 1) na primeira aba.

    Valores zero contam como dados. O resultado fica em cache porque a leitura
    completa da exportação concentra a maior parte do custo da avaliação.
    """
    workbook = openpyxl.load_workbook(client_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows)
        populated = [0] * len(header)
        for row in rows:
            for index, value in enumerate(row[: len(header)]):
                if value is not None and str(value).strip() != "":
                    populated[index] += 1
    finally:
        workbook.close()
    return SheetTable.from_headers(header), frozenset(
        i + 1 for i, count in enumerate(populated) if count == 0
    )


def read_mappings(report_path: str | Path) -> tuple[list[tuple], dict[str, int]]:
    """Lê a aba de mapeamentos de um relatório.

    Retorna as linhas sem conversão (a primeira contém o cabeçalho) e o índice {nome
    interno do campo: posição da coluna}. Os valores das células não são
    convertidos; quem compara situação ou método usa
    ``report_schema.internal_value``.

    Levanta ``ValueError`` se falta a aba ``Mapeamentos`` ou se o relatório
    está, inteiro ou em parte, no formato antigo em inglês; a mensagem diz
    como convertê-lo.
    """
    workbook = openpyxl.load_workbook(report_path, read_only=True, data_only=True)
    try:
        name = report_schema.find_sheet(list(workbook.sheetnames), "mappings")
        if name is None:
            raise ValueError(
                f"{report_path}: o relatório não tem a aba "
                f"{report_schema.sheet_name('mappings')!r}."
            )
        raw = list(workbook[name].iter_rows(values_only=True))
    finally:
        workbook.close()
    return raw, report_schema.internal_header_index(raw[0], "mappings")


def score_report(
    report_path: str | Path,
    client_path: str | Path,
    key: AnswerKey,
) -> Score:
    """Avalia um relatório já gravado, sem executar o mapeador."""
    client, vacuous = _client_structure(Path(client_path))
    raw, column = read_mappings(report_path)

    result = Score(client=key.client, present=key.denominator)
    numbers = [positive_column_number(cells[column["model_column_number"]],
                                      "numero_coluna_modelo") for cells in raw[1:]]
    if len(set(numbers)) != len(numbers):
        raise ValueError("O relatório tem numero_coluna_modelo repetido")
    if sorted(numbers) != list(range(1, len(numbers) + 1)):
        raise ValueError("O relatório deve conter cada numero_coluna_modelo de 1 a N exatamente uma vez")
    delivered_sources: set[int] = set()
    for position, cells in sorted(zip(numbers, raw[1:])):
        row = _score_row(position, cells, column, key, vacuous, client)
        if row.status == "accepted" and row.client_column is not None:
            if row.client_column in delivered_sources:
                raise ValueError(f"A coluna do cliente {row.client_column} foi confirmada para mais de um destino")
            delivered_sources.add(row.client_column)
        result.rows.append(row)
        if row.outcome is Outcome.FABRICATED:
            result.vacuous_targets.append(row)
        elif row.client_column in vacuous:
            result.vacuous_proposals.append(row)
    return result


def _score_row(
    position: int,
    cells: tuple,
    column: dict[str, int],
    key: AnswerKey,
    vacuous: frozenset[int],
    client: SheetTable,
) -> Row:
    """Lê a proposta de uma linha, converte seus valores internos e classifica o resultado."""
    record = {field: report_schema.internal_value("mappings", field, cells[index])
              for field, index in column.items()}
    record.setdefault("reviewer_status", "")
    record.setdefault("correct_client_column", "")
    record["reviewer_status"] = normalize_review(record)
    decision = resolve_mapping(record, client)
    chosen = decision.column.index + 1 if decision.column is not None else None
    status = decision.status
    method = str(
        report_schema.internal_value("mappings", "method", cells[column["method"]]) or ""
    )
    confidence = cells[column["final_confidence"]]
    # Relatórios antigos não têm a coluna prova_aritmetica.
    proof = cells[column["arithmetic_proof"]] if "arithmetic_proof" in column else None
    return Row(
        model_column=position,
        model_name=str(cells[column["model_column"]]),
        client_column=chosen,
        status=status,
        method=method,
        final_confidence=float(confidence) if isinstance(confidence, (int, float)) else None,
        arithmetic_proof=None if proof in (None, "") else str(proof),
        decided_by=(
            str(record.get("decided_by") or "")
            if record["reviewer_status"] != "incorrect" else None
        ),
        outcome=classify(
            chosen_client_column=chosen,
            expected_client_column=key.present.get(position),
            chosen_is_vacuous=chosen in vacuous,
            delivered=status == "accepted",
        ),
    )


def decision_census(score: Score) -> Counter[str]:
    """Conta o método vencedor; não atribui a uma correção humana o método da proposta antiga."""
    return Counter(r.decided_by if r.objective_decision is not None else "unrecorded"
                   for r in score.accepted)
