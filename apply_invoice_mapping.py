"""Gera a planilha no layout do modelo, usando o relatório e os pareceres do revisor.

``report_schema`` lê o relatório em português e converte nomes e valores
para os identificadores internos usados nas validações. Relatórios no
formato antigo recebem uma mensagem com instruções de conversão.

A saída contém Dados Mapeados, Cobertura e Colunas do Cliente sem Uso. Valores do cliente,
justificativas, provas aritméticas e observações do revisor são copiados
sem alteração.
"""

from __future__ import annotations

import argparse
import sys
from datetime import time
from pathlib import Path
from typing import Sequence

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import report_schema
from cli_ptbr import PortugueseArgumentParser
from spreadsheet import ColumnRef, SheetTable, positive_column_number
from mapping_review import effective_status, normalize_review, resolve_mapping, _column_from_number
from invoice_mapper import (
    _error_message,
    describe_header_conflict_pt,
    file_sha256,
    semantic_header_conflict,
)


#: Linhas do banner de Dados Mapeados (linhas 2 a 22 do Excel), na ordem
#: gravada: (rótulo da coluna A, campo de Mapeamentos de onde sai o valor).
MAPPING_METADATA = [
    ("numero_coluna_cliente", "client_column_number"),
    ("coluna_cliente", "client_column"),
    ("situacao", "status"),
    ("metodo", "method"),
    ("confianca_nome", "name_confidence"),
    ("confianca_nome_segunda_melhor", "name_runner_up_confidence"),
    ("margem_nome", "name_margin"),
    ("confianca_dados", "data_confidence"),
    ("confianca_dados_ia", "llm_data_confidence"),
    ("confianca_final_ia", "llm_final_confidence"),
    ("nome_catalogo", "catalog_name"),
    ("cobertura_catalogo", "catalog_coverage"),
    ("confianca_catalogo", "catalog_confidence"),
    ("confianca_final", "final_confidence"),
    ("decidido_por", "decided_by"),
    ("apoiado_por", "supported_by"),
    ("vetado_por", "vetoed_by"),
    ("prova_aritmetica", "arithmetic_proof"),
    ("justificativa", "reason"),
    ("parecer_revisor", "reviewer_status"),
    ("observacao_revisor", "reviewer_note"),
]

#: Primeira linha do Excel com dados do cliente, abaixo do cabeçalho e do
#: banner. Já mudou mais de uma vez: importe-a em vez de fixar o número.
FIRST_DATA_EXCEL_ROW = len(MAPPING_METADATA) + 2

#: Campos do banner gravados como número e exibidos como porcentagem.
CONFIDENCE_ROWS = {
    "name_confidence",
    "name_runner_up_confidence",
    "name_margin",
    "data_confidence",
    "llm_data_confidence",
    "llm_final_confidence",
    "catalog_coverage",
    "catalog_confidence",
    "final_confidence",
}

#: Colunas sem as quais o relatório não pode ser aplicado.
REQUIRED_MAPPING_COLUMNS = {
    "model_column_number",
    "model_column",
    "client_column_number",
    "client_column",
    "name_confidence",
    "name_runner_up_confidence",
    "name_margin",
    "data_confidence",
    "final_confidence",
    "method",
    "status",
    "reason",
}

#: Colunas que os relatórios mais antigos não têm; entram em branco, para que
#: esses relatórios continuem aplicáveis.
OPTIONAL_MAPPING_COLUMNS = (
    "decided_by",
    "supported_by",
    "vetoed_by",
    "arithmetic_proof",
    "catalog_name",
    "catalog_coverage",
    "catalog_confidence",
    "llm_data_confidence",
    "llm_final_confidence",
    "reviewer_status",
    "correct_client_column",
    "reviewer_note",
)

#: Situações exibidas na aba Cobertura, na ordem de apresentação. Apenas
#: "aguardando decisão" indica trabalho pendente; um campo ausente na
#: exportação não exige ação do revisor.
STATE_DELIVERED = "entregue"
STATE_PENDING = "aguardando decisão"
STATE_ABSENT = "não consta na exportação do cliente"
COVERAGE_STATE_ORDER = (STATE_DELIVERED, STATE_PENDING, STATE_ABSENT)

#: Colunas da aba Cobertura (nome interno -> largura). O cabeçalho gravado é o
#: português de ``report_schema.COLUMNS["coverage"]``.
COVERAGE_FIELDS = {
    "state": 26,
    "model_column_number": 10,
    "model_column": 34,
    "excel_column": 8,
    "client_column_number": 10,
    "client_column": 30,
    "final_confidence": 12,
    "reason": 78,
}

PERCENT_FORMAT = "0.00%"
TIME_FORMAT = "hh:mm:ss"
#: Fundo azul do banner (e do cabeçalho da Cobertura).
BANNER_COLOR = "E8F0FE"
#: Fundo verde do cabeçalho das colunas entregues.
DELIVERED_COLOR = "D7E9DC"
INFO_COLUMN_WIDTH = 34
MODEL_COLUMN_WIDTH = 18
#: Até esta coluna do Excel (AD) a largura é fixada mesmo nas colunas ocultas.
LAST_FIXED_WIDTH_EXCEL_COLUMN = 30


def _sheet_argument(value: str) -> str | int:
    """Interpreta dígitos como índice de aba e os demais valores como nome."""
    return int(value) if value.isdigit() else value


def _clean_text(value: object) -> str:
    """Remove espaços nas extremidades do texto; retorna ``""`` para células vazias."""
    return "" if pd.isna(value) else str(value).strip()


def _blank_if_missing(value: object) -> object:
    """Substitui valores ausentes por ``""`` e preserva os demais."""
    return "" if pd.isna(value) else value


def _column_label(field: str) -> str:
    """Retorna o cabeçalho usado para o campo na aba Mapeamentos."""
    return report_schema.header_label("mappings", field)


def _display_name(text: object) -> object:
    """Converte o nome de exibição para o relatório, incluindo o sufixo ``(ocorrência N)``."""
    return report_schema.display_name_to_external(text)


# ---------------------------------------------------------------------------
# Leitura e conferência do relatório
# ---------------------------------------------------------------------------


def _read_mapping_report(mapping_path: str | Path) -> pd.DataFrame:
    """Lê a primeira aba do relatório e a converte para nomes e valores internos.

    A aba é lida pela posição, não pelo nome, então Mapeamentos precisa
    continuar sendo a primeira. ``to_internal_frame`` recusa o formato antigo.
    """
    mapping = pd.read_excel(mapping_path, dtype=object)
    return report_schema.to_internal_frame(mapping, "mappings")


def _run_metadata(mapping_path: str | Path) -> dict[str, str]:
    """Lê os metadados da execução e converte as chaves para os nomes internos.

    Se a aba estiver ausente, retorna ``{}`` e o relatório é aplicado sem
    verificação de hash. Uma aba com nome do formato antigo gera erro para
    evitar que a verificação seja ignorada por uma diferença de idioma.
    """
    workbook = pd.ExcelFile(mapping_path)
    name = report_schema.find_sheet(workbook.sheet_names, "run_metadata")
    if name is None:
        return {}
    frame = pd.read_excel(workbook, sheet_name=name, dtype=object)
    return report_schema.internal_metadata(frame)


def _validate_input_hashes(
    metadata: dict[str, str],
    *,
    client_path: str | Path,
    model_path: str | Path,
) -> None:
    """Recusa o relatório se a exportação ou o modelo mudaram desde a execução que o gerou."""
    for key, path in (("client_sha256", client_path), ("model_sha256", model_path)):
        expected = metadata.get(key)
        if expected and expected != file_sha256(path):
            raise ValueError(
                f"{Path(path).name} difere do arquivo usado para gerar o relatório de mapeamento"
            )


def _ordered_mapping(mapping: pd.DataFrame, model: SheetTable) -> pd.DataFrame:
    """Valida o relatório e organiza suas linhas na ordem das colunas do modelo.

    Preenche as colunas opcionais ausentes e converte os pareceres do revisor
    para os valores internos.
    """
    missing = REQUIRED_MAPPING_COLUMNS - set(mapping.columns)
    if missing:
        labels = sorted(_column_label(field) for field in missing)
        raise ValueError(f"Faltam colunas no relatório de mapeamento: {labels}")

    mapping = mapping.copy()
    mapping["model_column_number"] = [
        positive_column_number(value, _column_label("model_column_number"))
        for value in mapping["model_column_number"]
    ]
    if mapping["model_column_number"].duplicated().any():
        raise ValueError(
            "O relatório de mapeamento tem valores repetidos em "
            f"{_column_label('model_column_number')}"
        )
    expected_numbers = list(range(1, len(model.columns) + 1))
    if sorted(mapping["model_column_number"].tolist()) != expected_numbers:
        raise ValueError(
            f"O relatório de mapeamento deve conter cada coluna do modelo "
            f"1..{len(model.columns)} exatamente uma vez"
        )

    mapping = mapping.set_index("model_column_number").loc[expected_numbers].reset_index()
    for row, model_column in zip(mapping.to_dict("records"), model.columns, strict=True):
        written = _clean_text(row["model_column"])
        if written != model_column.name:
            raise ValueError(
                f"A coluna do modelo {model_column.index + 1} difere entre o modelo e o "
                f"relatório de mapeamento: {model_column.name!r} != {written!r}"
            )

    for field in OPTIONAL_MAPPING_COLUMNS:
        if field not in mapping.columns:
            mapping[field] = ""
    mapping["reviewer_status"] = [normalize_review(row) for row in mapping.to_dict("records")]
    return mapping


# ---------------------------------------------------------------------------
# Escolha da coluna do cliente de cada coluna do modelo
# ---------------------------------------------------------------------------


def _refuse_header_conflict(model_column: ColumnRef, column: ColumnRef) -> None:
    """Rejeita associações entre cabeçalhos incompatíveis segundo as regras atuais."""
    conflict = semantic_header_conflict(model_column.name, column.name)
    if conflict is None:
        return
    description = describe_header_conflict_pt(model_column.name, column.name) or conflict
    raise ValueError(
        f"O mapeamento aceito {model_column.name!r} -> "
        f"{_display_name(column.display_name)!r} viola as regras atuais de "
        f"coerência entre cabeçalhos: {description}. Revise-o explicitamente "
        f"(preencha {_column_label('reviewer_status')}) antes de aplicar."
    )


def _refuse_proposal_from_another_export(row: dict, client: SheetTable) -> None:
    """Verifica se o nome da coluna proposta corresponde ao da exportação informada."""
    proposed = _column_from_number(client, row["client_column_number"])
    proposed_name = _clean_text(row["client_column"])
    if proposed is None or not proposed_name or proposed.display_name == proposed_name:
        return
    raise ValueError(
        f"O relatório de mapeamento não confere com a planilha do cliente na coluna "
        f"{proposed.index + 1}: {_display_name(proposed_name)!r} != "
        f"{_display_name(proposed.display_name)!r}"
    )


def _effective_columns(
    mapping: pd.DataFrame,
    client: SheetTable,
    model: SheetTable,
    accepted_only: bool,
) -> list[ColumnRef | None]:
    """Define a origem de cada coluna do modelo após aplicar os pareceres do revisor.

    Cada coluna do cliente pode ser entregue apenas uma vez. Propostas ainda
    em revisão podem compartilhar a mesma origem.
    """
    selected: list[ColumnRef | None] = []
    confirmed_indexes: set[int] = set()
    for row, model_column in zip(mapping.to_dict("records"), model.columns, strict=True):
        decision = resolve_mapping(row, client)
        confirmed = decision.delivered
        column = decision.column if confirmed or not accepted_only else None
        if column is not None and row["status"] == "accepted" and row["reviewer_status"] == "":
            _refuse_header_conflict(model_column, column)
        _refuse_proposal_from_another_export(row, client)

        if column is not None and confirmed:
            if column.index in confirmed_indexes:
                raise ValueError(
                    f"A coluna do cliente {_display_name(column.display_name)!r} foi mapeada "
                    "mais de uma vez após as correções do revisor"
                )
            confirmed_indexes.add(column.index)
        selected.append(column)
    return selected


# ---------------------------------------------------------------------------
# Gravação da planilha
# ---------------------------------------------------------------------------


def _banner_rows(
    mapping: pd.DataFrame, selected: Sequence[ColumnRef | None]
) -> list[list[object]]:
    """Monta os rótulos e valores do banner, com uma célula para cada coluna do modelo.

    ``numero_coluna_cliente`` e ``coluna_cliente`` identificam a origem efetiva,
    considerando as correções e confirmações do revisor.
    """
    delivered = {
        "client_column_number": ["" if column is None else column.index + 1 for column in selected],
        "client_column": ["" if column is None else column.display_name for column in selected],
    }
    rows: list[list[object]] = []
    for label, field in MAPPING_METADATA:
        if field in delivered:
            values = delivered[field]
        else:
            values = [_blank_if_missing(value) for value in mapping[field]]
        rows.append(
            [label, *(report_schema.external_value("banner", field, value) for value in values)]
        )
    return rows


def _client_data(
    client: SheetTable, model: SheetTable, selected: Sequence[ColumnRef | None], info_header: str
) -> pd.DataFrame:
    """Organiza os valores do cliente no layout do modelo, deixando vazias as colunas sem origem."""
    row_count = len(client.data)
    data = pd.concat(
        [
            pd.Series([pd.NA] * row_count, dtype=object)
            if column is None
            else client.series(column).reset_index(drop=True)
            for column in selected
        ],
        axis=1,
    )
    data.columns = [column.name for column in model.columns]
    data.insert(0, info_header, "")
    return data


def _format_banner(worksheet, column_count: int) -> None:
    """Aplica o fundo azul ao banner e formata as confianças como porcentagens."""
    fill = PatternFill(fill_type="solid", fgColor=BANNER_COLOR)
    for excel_row in range(2, FIRST_DATA_EXCEL_ROW):
        for cell in worksheet[excel_row]:
            cell.fill = fill
    for excel_row, (_, field) in enumerate(MAPPING_METADATA, start=2):
        if field in CONFIDENCE_ROWS:
            for excel_column in range(2, column_count + 1):
                worksheet.cell(excel_row, excel_column).number_format = PERCENT_FORMAT


def _restore_time_cells(
    worksheet, client: SheetTable, selected: Sequence[ColumnRef | None]
) -> None:
    """Restaura o formato de hora do Excel para valores ``datetime.time`` gravados como texto pelo pandas."""
    for excel_column, column in enumerate(selected, start=2):
        if column is None:
            continue
        for excel_row, value in enumerate(client.series(column), start=FIRST_DATA_EXCEL_ROW):
            if isinstance(value, time):
                cell = worksheet.cell(excel_row, excel_column)
                cell.value = value
                cell.number_format = TIME_FORMAT


def _arrange_columns(worksheet, selected: Sequence[ColumnRef | None], column_count: int) -> None:
    """Ajusta as larguras, destaca as colunas entregues e oculta as vazias.

    Isso permite visualizar as colunas preenchidas lado a lado. As demais
    continuam no arquivo e podem ser reexibidas no Excel ou lidas pelo código.
    """
    worksheet.column_dimensions["A"].width = INFO_COLUMN_WIDTH
    for excel_column in range(2, min(column_count, LAST_FIXED_WIDTH_EXCEL_COLUMN) + 1):
        worksheet.column_dimensions[get_column_letter(excel_column)].width = MODEL_COLUMN_WIDTH

    delivered_fill = PatternFill(fill_type="solid", fgColor=DELIVERED_COLOR)
    for excel_column, column in enumerate(selected, start=2):
        letter = get_column_letter(excel_column)
        if column is None:
            worksheet.column_dimensions[letter].hidden = True
        else:
            worksheet.column_dimensions[letter].width = MODEL_COLUMN_WIDTH
            worksheet.cell(1, excel_column).fill = delivered_fill


def _coverage_row(
    model_column: ColumnRef, column: ColumnRef | None, row: dict
) -> list[object]:
    """Monta a linha de uma coluna do modelo na aba Cobertura.

    Usa a decisão após a revisão. Nas linhas entregues, registra o número
    da coluna de origem efetivamente copiada.
    """
    if column is not None:
        state = STATE_DELIVERED
        client_number, client_name = column.index + 1, column.display_name
    else:
        state = STATE_PENDING if effective_status(row) == "review" else STATE_ABSENT
        client_number = _blank_if_missing(row["client_column_number"])
        client_name = _blank_if_missing(row["client_column"])
    return [
        state,
        model_column.index + 1,
        model_column.name,
        get_column_letter(model_column.index + 2),
        client_number,
        _display_name(client_name),
        _blank_if_missing(row["final_confidence"]),
        _blank_if_missing(row["reason"]),
    ]


def _write_coverage_sheet(
    book,
    *,
    model: SheetTable,
    selected: Sequence[ColumnRef | None],
    mapping: pd.DataFrame,
) -> None:
    """Grava a situação de cada coluna na aba Cobertura, com as entregues primeiro.

    A aba diferencia o que foi entregue, o que aguarda revisão e o que não
    consta na exportação do cliente.
    """
    sheet = book.create_sheet(report_schema.sheet_name("coverage"))
    sheet.append([report_schema.header_label("coverage", field) for field in COVERAGE_FIELDS])
    rows = [
        _coverage_row(model_column, column, row)
        for model_column, column, row in zip(model.columns, selected, mapping.to_dict("records"))
    ]
    for values in sorted(rows, key=lambda values: COVERAGE_STATE_ORDER.index(values[0])):
        sheet.append(values)

    last_letter = get_column_letter(len(COVERAGE_FIELDS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{last_letter}{sheet.max_row}"
    for position, width in enumerate(COVERAGE_FIELDS.values(), start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width
    header_fill = PatternFill(fill_type="solid", fgColor=BANNER_COLOR)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
    confidence_column = list(COVERAGE_FIELDS).index("final_confidence") + 1
    for (cell,) in sheet.iter_rows(
        min_row=2, min_col=confidence_column, max_col=confidence_column
    ):
        cell.number_format = PERCENT_FORMAT


def _write_unused_sheet(book, mapping: pd.DataFrame, client: SheetTable) -> None:
    """Lista as origens sem aceite após a revisão, inclusive as colunas vazias.

    Uma correção libera a origem anterior. Propostas pendentes continuam na
    lista, mesmo quando --include-review as inclui na saída de conferência.
    """
    decisions = [resolve_mapping(row, client) for row in mapping.to_dict("records")]
    consumed = {d.column.index for d in decisions if d.delivered}
    pending = {d.column.index for d in decisions if d.status == "review" and d.column is not None}
    fields = ("client_column_number", "client_column", "disposition", "rows_with_value", "distinct_values")
    sheet = book.create_sheet(report_schema.sheet_name("unmatched"))
    sheet.append([report_schema.header_label("unmatched", field) for field in fields])
    for column in client.columns:
        if column.index in consumed:
            continue
        values = client.series(column).dropna()
        values = values[values.astype(str).str.strip() != ""]
        row = (column.index + 1, column.display_name,
               "proposed_not_accepted" if column.index in pending else "never_referenced",
               len(values), values.nunique())
        sheet.append([report_schema.external_value("unmatched", field, value)
                      for field, value in zip(fields, row, strict=True)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell, width in zip(sheet[1], (24, 48, 30, 24, 24), strict=True):
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width


def build_mapped_workbook(
    client_path: str | Path,
    mapping_path: str | Path,
    output_path: str | Path,
    model_path: str | Path = "data/model.xlsx",
    model_sheet: str | int = 0,
    client_sheet: str | int = 0,
    model_header_row: int = 0,
    client_header_row: int = 0,
    accepted_only: bool = True,
) -> None:
    """Aplica as escolhas e grava as abas Dados Mapeados e Cobertura.

    Por padrão, inclui os mapeamentos aceitos sem parecer contrário e os
    confirmados ou corrigidos pelo revisor. ``accepted_only=False`` também
    inclui propostas pendentes; nesse caso, a saída serve apenas para conferência.
    """
    report = _read_mapping_report(mapping_path)
    _validate_input_hashes(
        _run_metadata(mapping_path), client_path=client_path, model_path=model_path
    )
    model = SheetTable.from_excel(model_path, model_sheet, model_header_row)
    client = SheetTable.from_excel(client_path, client_sheet, client_header_row)
    mapping = _ordered_mapping(report, model)
    selected = _effective_columns(mapping, client, model, accepted_only)

    info_header = report_schema.header_label("mapped_data", "__mapping_info__")
    output_columns = [info_header, *[column.name for column in model.columns]]
    banner = pd.DataFrame(_banner_rows(mapping, selected), columns=output_columns)
    data = _client_data(client, model, selected, info_header)
    output_frame = pd.concat([banner, data], axis=0, ignore_index=True)

    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("O arquivo de saída deve terminar em .xlsx")
    mapped_sheet = report_schema.sheet_name("mapped_data")
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        output_frame.to_excel(writer, sheet_name=mapped_sheet, index=False)
        worksheet = writer.book[mapped_sheet]
        worksheet.freeze_panes = f"B{FIRST_DATA_EXCEL_ROW}"
        _format_banner(worksheet, len(output_columns))
        _restore_time_cells(worksheet, client, selected)
        _arrange_columns(worksheet, selected, len(output_columns))
        _write_coverage_sheet(writer.book, model=model, selected=selected, mapping=mapping)
        _write_unused_sheet(writer.book, mapping, client)

    print(f"Arquivo gravado: {output}")
    print(
        f"colunas_modelo={len(model.columns)} linhas_cliente={len(client.data)} "
        f"colunas_mapeadas={sum(column is not None for column in selected)}"
    )


def build_parser() -> argparse.ArgumentParser:
    """Configura os argumentos e a ajuda da CLI."""
    parser = PortugueseArgumentParser(
        description=(
            "Aplica um relatório de mapeamento e gera a planilha do cliente no layout do "
            "modelo."
        )
    )
    parser.add_argument("client", metavar="EXPORTACAO", help="exportação .xlsx do cliente")
    parser.add_argument(
        "mapping",
        metavar="RELATORIO",
        help="relatório de mapeamento gerado pelo invoice_mapper.py",
    )
    parser.add_argument(
        "--model",
        default="data/model.xlsx",
        metavar="ARQUIVO",
        help="planilha do modelo base (padrão: %(default)s)",
    )
    parser.add_argument(
        "--output", required=True, metavar="ARQUIVO", help="planilha .xlsx de saída"
    )
    parser.add_argument(
        "--model-sheet",
        default="0",
        metavar="ABA",
        help="índice (a partir de 0) ou nome da aba do modelo",
    )
    parser.add_argument(
        "--client-sheet",
        default="0",
        metavar="ABA",
        help="índice (a partir de 0) ou nome da aba do cliente",
    )
    parser.add_argument(
        "--model-header-row",
        type=int,
        default=1,
        metavar="LINHA",
        help="linha do cabeçalho do modelo (contada a partir de 1)",
    )
    parser.add_argument(
        "--client-header-row",
        type=int,
        default=1,
        metavar="LINHA",
        help="linha do cabeçalho do cliente (contada a partir de 1)",
    )
    parser.add_argument(
        "--accepted-only",
        action="store_true",
        help="obsoleta: mapeamentos aceitos e confirmados pelo revisor já são o padrão",
    )
    parser.add_argument(
        "--include-review",
        action="store_true",
        help=(
            "inclui propostas ainda não confirmadas pelo revisor; gera uma planilha "
            "para conferência, que não deve ser entregue ao cliente"
        ),
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    """Executa a CLI; em caso de erro de entrada ou de arquivo, exibe a mensagem e retorna 1."""
    args = build_parser().parse_args(argv)
    if args.accepted_only and args.include_review:
        raise SystemExit("--accepted-only e --include-review não podem ser usados juntos")
    if args.model_header_row < 1 or args.client_header_row < 1:
        raise SystemExit(
            "--model-header-row e --client-header-row devem ser positivos "
            "(a contagem começa em 1)"
        )
    try:
        build_mapped_workbook(
            client_path=args.client,
            mapping_path=args.mapping,
            output_path=args.output,
            model_path=args.model,
            model_sheet=_sheet_argument(args.model_sheet),
            client_sheet=_sheet_argument(args.client_sheet),
            model_header_row=args.model_header_row - 1,
            client_header_row=args.client_header_row - 1,
            accepted_only=not args.include_review,
        )
    except OSError as exc:
        raise SystemExit(_error_message(exc)) from exc
    except ValueError as exc:
        # As mensagens locais estão em português. Exceções do pandas mantêm
        # o texto original, útil para o diagnóstico técnico.
        raise SystemExit(str(exc)) from exc
    return 0


if __name__ == "__main__":
    sys.exit(main())
