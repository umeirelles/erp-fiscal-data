"""Avaliação com os gabaritos de-para (comando ``evaluate``).

Lê relatórios já gravados, sem executar o mapper nem chamar a API, e compara a
decisão efetiva de cada campo, após os pareceres, com o gabarito do cliente:
acertos, perdas, colunas erradas, fabricações e o que foi entregue.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
import enum
from functools import lru_cache
from pathlib import Path
import re

import openpyxl

from mapper import schema
from mapper.apply import normalize_review, resolve_mapping
from mapper.cli import PortugueseArgumentParser
from mapper.decide import TRANSFERRING_STEP_IDS
from mapper.spreadsheet import SheetTable, positive_column_number

# O recall usa como denominador as colunas anotadas no gabarito. Os resultados
# separam acertos, perdas, escolhas erradas e aceites de colunas vazias: uma taxa
# única de acerto esconderia essas diferenças.

#: Descrição exibida por ``run.py evaluate --help``.
DESCRIPTION = "Compara os relatórios arquivados com os gabaritos e apresenta as métricas de mapeamento."


# ---------------------------------------------------------------------------
# Letras de coluna do Excel e anotações dos gabaritos
# ---------------------------------------------------------------------------

_COLUNA = re.compile(r"coluna\s*([A-Za-z]+)", re.IGNORECASE)


def column_letter_to_index(letters: str) -> int:
    """Converte letras de coluna do Excel em um índice a partir de 1.

    >>> column_letter_to_index("A"), column_letter_to_index("Z")
    (1, 26)
    >>> column_letter_to_index("AA"), column_letter_to_index("FY")
    (27, 181)
    """
    text = letters.strip()
    if not text or not text.isalpha():
        raise ValueError(f"não é uma letra de coluna do Excel: {letters!r}")
    index = 0
    for char in text.upper():
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def index_to_column_letter(index: int) -> str:
    """Converte um índice a partir de 1 nas letras correspondentes da coluna do Excel."""
    if index < 1:
        raise ValueError(f"o índice da coluna deve ser >= 1; recebido {index}")
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def parse_coluna_annotation(cell: object) -> int | None:
    """Extrai a coluna indicada pela primeira anotação "Coluna <letras>" da célula.

    Retorna o índice a partir de 1, ou ``None`` se não houver anotação.
    Aceita espaços adicionais, como em "Coluna BV ", presente no gabarito do datasul.
    """
    if cell is None:
        return None
    match = _COLUNA.search(str(cell))
    if match is None:
        return None
    return column_letter_to_index(match.group(1))


# ---------------------------------------------------------------------------
# Resultado de cada campo
#
#
# As categorias separam uma coluna não encontrada de uma escolha errada ou
# de uma coluna vazia aceita como origem, para avaliar o impacto de cada erro.
# ---------------------------------------------------------------------------


class Outcome(enum.Enum):
    """Resultado de uma coluna do modelo.

    Os valores são identificadores internos e nunca vão para arquivo; a CLI
    imprime o rótulo em português de ``schema.ENUMS["outcome"]``.
    """

    #: O mapeador escolheu a coluna do cliente que o gabarito indica.
    CORRECT = "correct"
    #: O gabarito indica uma coluna e o mapeador não propôs nenhuma.
    MISSED = "missed"
    #: O gabarito indica uma coluna e o mapeador propôs outra.
    WRONG_COLUMN = "wrong_column"
    #: Mapeamento aceito cuja coluna de origem está vazia.
    FABRICATED = "fabricated"
    #: Campo fora do gabarito; não há referência para confirmar a escolha.
    UNDETERMINED = "undetermined"


def classify(
    *,
    chosen_client_column: int | None,
    expected_client_column: int | None,
    chosen_is_vacuous: bool,
    delivered: bool = True,
) -> Outcome:
    """Compara a escolha com o gabarito e verifica se houve aceitação de uma coluna vazia.

    ``expected_client_column=None`` indica que o campo não consta no gabarito.
    Antes dessa comparação, uma coluna vazia entregue é classificada como
    fabricação, mesmo que o gabarito aponte para ela. Propostas ainda em revisão
    não entram nessa categoria. Uma coluna preenchida com zeros contém dados
    e não é considerada vazia.
    """
    if delivered and chosen_client_column is not None and chosen_is_vacuous:
        return Outcome.FABRICATED
    if expected_client_column is None:
        # Fora do gabarito, não é possível confirmar nem a escolha nem a ausência.
        # Contabilizar esses casos como acertos aumentaria as métricas sem validação.
        return Outcome.UNDETERMINED
    if chosen_client_column is None:
        return Outcome.MISSED
    if chosen_client_column == expected_client_column:
        return Outcome.CORRECT
    return Outcome.WRONG_COLUMN


# ---------------------------------------------------------------------------
# Gabaritos de-para
#
#
# Só datasul e oracle têm gabarito de colunas: de_para_procosa.xlsx é uma
# especificação no nível de registro, e sap_terumo não tem de-para.
# ---------------------------------------------------------------------------

# Correções do gabarito verificadas nos dados do cliente. Sem elas, a avaliação
# classificaria essas escolhas corretas como erros. --strict-key-only usa o
# gabarito original para permitir a comparação. A CLI mostra apenas os números
# das colunas corrigidas; as justificativas ficam neste registro.
KEY_CORRECTIONS: dict[str, dict[int, tuple[int, str]]] = {
    "oracle": {
        213: (36, "o gabarito aponta c37 ALIQ COFINS; c36 ALIQ PIS é a alíquota do PIS"),
        220: (40, "o gabarito aponta c39 CST-PIS; o par está trocado no gabarito"),
        221: (39, "o gabarito aponta c40 CST-COFINS; o par está trocado no gabarito"),
    },
    "datasul": {
        184: (188, "o gabarito aponta c99 (bloco A), mas é incoerente consigo mesmo; "
                   "c188 fecha o trio fiscal com as colunas m52-m55 do próprio gabarito"),
    },
}


@dataclass(frozen=True)
class AnswerKey:
    """Reúne as correspondências anotadas pelo cliente e as correções verificadas do gabarito."""

    client: str
    #: número da coluna do modelo (a partir de 1) -> número da coluna do cliente (a partir de 1)
    present: dict[int, int]
    #: número da coluna do modelo -> por que a anotação do próprio cliente foi substituída
    corrections: dict[int, str] = field(default_factory=dict)

    @property
    def source_conflicts(self) -> dict[int, list[int]]:
        """Lista origens atribuídas a mais de um destino; exige revisão do gabarito."""
        targets = defaultdict(list)
        for target, source in self.present.items():
            targets[source].append(target)
        return {source: sorted(columns) for source, columns in sorted(targets.items())
                if len(columns) > 1}

    @property
    def denominator(self) -> int:
        """Retorna o número de colunas anotadas no gabarito, usado como denominador do recall."""
        return len(self.present)


def load_answer_key(
    path: str | Path,
    client: str,
    *,
    apply_corrections: bool = True,
) -> AnswerKey:
    """Lê a primeira linha do de-para e associa os números das colunas do modelo aos do cliente."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    finally:
        workbook.close()

    present: dict[int, int] = {}
    for model_column, cell in enumerate(header_row, start=1):
        client_column = parse_coluna_annotation(cell)
        if client_column is not None:
            present[model_column] = client_column

    corrections: dict[int, str] = {}
    if apply_corrections:
        for model_column, (correct_column, why) in KEY_CORRECTIONS.get(client, {}).items():
            if model_column in present:
                present[model_column] = correct_column
                corrections[model_column] = why

    return AnswerKey(client=client, present=present, corrections=corrections)


# ---------------------------------------------------------------------------
# Score de um relatório
#
#
# A leitura passa por ``schema``, que devolve os cabeçalhos com os nomes
# internos; situação e método voltam ao vocabulário interno (``accepted``,
# ``catalog``...) antes de qualquer comparação.
# ---------------------------------------------------------------------------


@dataclass
class Row:
    model_column: int
    model_name: str
    client_column: int | None
    #: Situação no vocabulário interno (``accepted``, ``review``, ``no_match``).
    status: str
    #: Método no vocabulário interno (``name``, ``data``, ``catalog``...).
    method: str
    final_confidence: float | None
    outcome: Outcome
    #: Relação aritmética nos dados do cliente que sustenta a escolha da coluna.
    arithmetic_proof: str | None = None

    #: Evidência que decidiu o aceite automático; ausente em relatórios antigos ou correções humanas.
    decided_by: str | None = None

    @property
    def objective_decision(self) -> bool | None:
        """Informa se o método vencedor usa dados; None significa método não registrado."""
        if self.decided_by in TRANSFERRING_STEP_IDS:
            return True
        if self.decided_by in {"lexical", "semantic"}:
            return False
        return None


@dataclass
class Score:
    client: str
    rows: list[Row] = field(default_factory=list)
    present: int = 0
    #: Mapeamentos aceitos para colunas vazias; fazem a avaliação retornar exit code 1.
    vacuous_targets: list[Row] = field(default_factory=list)
    #: Propostas para colunas vazias ainda não entregues; não alteram o exit code.
    vacuous_proposals: list[Row] = field(default_factory=list)

    def count(self, outcome: Outcome) -> int:
        return sum(1 for row in self.rows if row.outcome is outcome)

    @property
    def recall_on_present(self) -> float | None:
        """Calcula o recall usando apenas as colunas presentes no gabarito como denominador."""
        return self.count(Outcome.CORRECT) / self.present if self.present else None

    @property
    def accepted(self) -> list[Row]:
        return [r for r in self.rows if r.status == "accepted"]

    @property
    def delivered_correct(self) -> int:
        """Conta os acertos entregues após aplicar os pareceres do revisor."""
        return sum(
            1
            for r in self.rows
            if r.outcome is Outcome.CORRECT and r.status == "accepted"
        )

    @property
    def delivered_recall(self) -> float | None:
        """Calcula o recall dos mapeamentos corretos aceitos; não ultrapassa ``recall_on_present``."""
        return self.delivered_correct / self.present if self.present else None

    @property
    def withheld_correct(self) -> list[Row]:
        """Retorna as propostas corretas ainda não aceitas, por confiança decrescente.

        Essa ordenação ajuda o revisor a priorizar a análise.
        """
        held = [
            r
            for r in self.rows
            if r.outcome is Outcome.CORRECT and r.status != "accepted"
        ]
        return sorted(held, key=lambda r: -(r.final_confidence or 0.0))

    @property
    def accept_precision_lower(self) -> float | None:
        """Calcula o limite inferior de precisão considerando apenas os acertos confirmados.

        Mapeamentos fora do gabarito entram no denominador, mas não no numerador.
        Assim, uma nova escolha sem validação reduz o limite inferior."""
        accepts = self.accepted
        if not accepts:
            return None
        good = sum(1 for r in accepts if r.outcome is Outcome.CORRECT)
        return good / len(accepts)

    @property
    def undetermined_accepts(self) -> list[Row]:
        """Aceites em colunas do modelo fora do gabarito, que ele não consegue julgar.

        O exit code considera apenas fabricações. Uma escolha errada em coluna
        preenchida fora do gabarito pode passar pela avaliação com exit code 0;
        esta contagem permite acompanhar os casos sem validação.
        """
        return [r for r in self.accepted if r.outcome is Outcome.UNDETERMINED]

    @property
    def objective_decision_share(self) -> float | None:
        """Proporção de evidência objetiva entre aceites com método automático conhecido."""
        known = [r.objective_decision for r in self.accepted if r.objective_decision is not None]
        return sum(known) / len(known) if known else None

    @property
    def correct_pairs(self) -> set[tuple[int, int]]:
        """Os pares (coluna do modelo, coluna do cliente) corretos.

        Permite comparar os pares entre execuções. A mesma contagem pode esconder
        a perda de um acerto compensada por outro. O procedimento de
        docs/EVALUATION.md usa essa comparação.
        """
        return {
            (r.model_column, r.client_column)
            for r in self.rows
            if r.outcome is Outcome.CORRECT and r.client_column is not None
        }


@lru_cache(maxsize=None)
def _client_structure(client_path: str | Path) -> tuple[SheetTable, frozenset[int]]:
    """Lê os cabeçalhos e identifica as colunas vazias (a partir de 1) na primeira aba.

    Valores zero contam como dados. O resultado fica em cache porque a leitura
    completa da exportação concentra a maior parte do custo da avaliação.
    """
    workbook = openpyxl.load_workbook(client_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows)
        populated = [0] * len(header)
        for row in rows:
            for index, value in enumerate(row[: len(header)]):
                if value is not None and str(value).strip() != "":
                    populated[index] += 1
    finally:
        workbook.close()
    return SheetTable.from_headers(header), frozenset(
        i + 1 for i, count in enumerate(populated) if count == 0
    )


def read_mappings(report_path: str | Path) -> tuple[list[tuple], dict[str, int]]:
    """Lê a aba de mapeamentos de um relatório.

    Retorna as linhas sem conversão (a primeira contém o cabeçalho) e o índice {nome
    interno do campo: posição da coluna}. Os valores das células não são
    convertidos; quem compara situação ou método usa
    ``schema.internal_value``.

    Levanta ``ValueError`` se falta a aba ``Mapeamentos`` ou se o relatório
    está, inteiro ou em parte, no formato antigo em inglês; a mensagem diz
    como convertê-lo.
    """
    workbook = openpyxl.load_workbook(report_path, read_only=True, data_only=True)
    try:
        name = schema.find_sheet(list(workbook.sheetnames), "mappings")
        if name is None:
            raise ValueError(
                f"{report_path}: o relatório não tem a aba "
                f"{schema.sheet_name('mappings')!r}."
            )
        raw = list(workbook[name].iter_rows(values_only=True))
    finally:
        workbook.close()
    return raw, schema.internal_header_index(raw[0], "mappings")


def score_report(
    report_path: str | Path,
    client_path: str | Path,
    key: AnswerKey,
) -> Score:
    """Avalia um relatório já gravado, sem executar o mapeador."""
    client, vacuous = _client_structure(Path(client_path))
    raw, column = read_mappings(report_path)

    result = Score(client=key.client, present=key.denominator)
    numbers = [positive_column_number(cells[column["model_column_number"]],
                                      "numero_coluna_modelo") for cells in raw[1:]]
    if len(set(numbers)) != len(numbers):
        raise ValueError("O relatório tem numero_coluna_modelo repetido")
    if sorted(numbers) != list(range(1, len(numbers) + 1)):
        raise ValueError("O relatório deve conter cada numero_coluna_modelo de 1 a N exatamente uma vez")
    delivered_sources: set[int] = set()
    for position, cells in sorted(zip(numbers, raw[1:])):
        row = _score_row(position, cells, column, key, vacuous, client)
        if row.status == "accepted" and row.client_column is not None:
            if row.client_column in delivered_sources:
                raise ValueError(f"A coluna do cliente {row.client_column} foi confirmada para mais de um destino")
            delivered_sources.add(row.client_column)
        result.rows.append(row)
        if row.outcome is Outcome.FABRICATED:
            result.vacuous_targets.append(row)
        elif row.client_column in vacuous:
            result.vacuous_proposals.append(row)
    return result


def _score_row(
    position: int,
    cells: tuple,
    column: dict[str, int],
    key: AnswerKey,
    vacuous: frozenset[int],
    client: SheetTable,
) -> Row:
    """Lê a proposta de uma linha, converte seus valores internos e classifica o resultado."""
    record = {field: schema.internal_value("mappings", field, cells[index])
              for field, index in column.items()}
    record.setdefault("reviewer_status", "")
    record.setdefault("correct_client_column", "")
    record["reviewer_status"] = normalize_review(record)
    decision = resolve_mapping(record, client)
    chosen = decision.column.index + 1 if decision.column is not None else None
    status = decision.status
    method = str(
        schema.internal_value("mappings", "method", cells[column["method"]]) or ""
    )
    confidence = cells[column["final_confidence"]]
    # Relatórios antigos não têm a coluna prova_aritmetica.
    proof = cells[column["arithmetic_proof"]] if "arithmetic_proof" in column else None
    return Row(
        model_column=position,
        model_name=str(cells[column["model_column"]]),
        client_column=chosen,
        status=status,
        method=method,
        final_confidence=float(confidence) if isinstance(confidence, (int, float)) else None,
        arithmetic_proof=None if proof in (None, "") else str(proof),
        decided_by=(
            str(record.get("decided_by") or "")
            if record["reviewer_status"] != "incorrect" else None
        ),
        outcome=classify(
            chosen_client_column=chosen,
            expected_client_column=key.present.get(position),
            chosen_is_vacuous=chosen in vacuous,
            delivered=status == "accepted",
        ),
    )


def decision_census(score: Score) -> Counter[str]:
    """Conta o método vencedor; não atribui a uma correção humana o método da proposta antiga."""
    return Counter(r.decided_by if r.objective_decision is not None else "unrecorded"
                   for r in score.accepted)


# ---------------------------------------------------------------------------
# Comando evaluate (python run.py evaluate)
# ---------------------------------------------------------------------------

#: cliente -> (exportação do cliente, relatório de mapeamento, gabarito de-para ou None)
CLIENTS: dict[str, tuple[str, str, str | None]] = {
    "datasul": (
        "data/clients/datasul_maio.xlsx",
        "data/reports/mapping_datasul_full.xlsx",
        "data/keys/de_para_datasul.xlsx",
    ),
    "oracle": (
        "data/clients/oracle.xlsx",
        "data/reports/mapping_oracle_full.xlsx",
        "data/keys/de_para_oracle.xlsx",
    ),
    # de_para_procosa.xlsx é uma especificação no nível de registro, não um
    # gabarito de colunas.
    "sap_procosa": (
        "data/clients/sap_procosa.XLSX",
        "data/reports/mapping_sap_procosa_full.xlsx",
        None,
    ),
    "sap_terumo": (
        "data/clients/sap_terumo.XLSX",
        "data/reports/mapping_sap_terumo_full.xlsx",
        None,
    ),
}

#: Rótulo impresso de cada resultado. ``Outcome.value`` continua em inglês.
OUTCOME_LABELS: dict[Outcome, str] = {
    outcome: schema.ENUMS["outcome"][outcome.value] for outcome in Outcome
}

#: Início da linha de totais; a linha seguinte é recuada até depois dele.
_TOTAL_LABEL = "TOTAL COM GABARITO  "


def _pct(value: float | None) -> str:
    """Formata uma porcentagem com vírgula decimal e largura mínima de 6; usa ``n/d`` se ausente."""
    return "  n/d " if value is None else schema.format_percent_br(value).rjust(6)


def _decimal(value: float | None) -> str:
    """Formata um número com vírgula decimal, preservando os dígitos; usa ``n/d`` se ausente."""
    return "n/d" if value is None else schema.format_decimal_br(value)


def build_parser() -> PortugueseArgumentParser:
    parser = PortugueseArgumentParser(prog="run.py evaluate", description=DESCRIPTION)
    parser.add_argument("--root", default=".", metavar="PASTA",
                        help="raiz do repositório com as planilhas")
    parser.add_argument("--client", action="append", metavar="CLIENTE",
                        help="limita a avaliação a este cliente; repita a opção para mais de um")
    parser.add_argument("--strict-key-only", action="store_true",
                        help="não aplica as quatro correções verificadas do gabarito")
    parser.add_argument("--show-losses", action="store_true",
                        help="lista cada acerto retido e cada linha perdida ou com coluna errada")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    root = Path(args.root)

    found = delivered = present = fabrications = 0
    for name in args.client or list(CLIENTS):
        # Um nome fora de CLIENTS ainda termina em KeyError; veja docs/KNOWN_ISSUES.md.
        client_file, report_file, key_file = CLIENTS[name]
        if key_file is None:
            fabrications += _report_unkeyed_client(name, root / client_file, root / report_file)
            continue

        key = load_answer_key(root / key_file, name,
                              apply_corrections=not args.strict_key_only)
        score = score_report(root / report_file, root / client_file, key)
        _report_keyed_client(name, score, key, show_losses=args.show_losses)
        found += score.count(Outcome.CORRECT)
        delivered += score.delivered_correct
        present += score.present
        fabrications += score.count(Outcome.FABRICATED)

    if present:
        _report_totals(found, delivered, present, fabrications)
    # Somente fabricações alteram o exit code. Escolhas fora das 56 colunas
    # cobertas pelo gabarito podem aumentar "sem_veredito" sem causar falha.
    # Essa contagem também precisa ser considerada ao avaliar uma mudança.
    return 1 if fabrications else 0


def _report_unkeyed_client(name: str, client_path: Path, report_path: Path) -> int:
    """Exibe os totais de um cliente sem gabarito e retorna a contagem de aceites em colunas vazias.

    Se o relatório não existir, retorna 0 sem exibir aviso.
    """
    if not report_path.exists():
        return 0
    print(f"\n{name}  SEM GABARITO — não é possível calcular precisão ou recall.")
    score = score_report(report_path, client_path, AnswerKey(name, {}))
    pending = sum(row.status == "review" for row in score.rows)
    print(f"  aceitos {len(score.accepted):3}   pendências de revisão {pending:3}"
          f"   aceites em coluna vazia {len(score.vacuous_targets)}")
    return len(score.vacuous_targets)


def _report_keyed_client(name: str, score: Score, key: AnswerKey, *, show_losses: bool) -> None:
    fabricated = score.count(Outcome.FABRICATED)
    # Os métodos saem como estão gravados no relatório em português.
    census = {
        schema.external_value("mappings", "decided_by", method): count
        for method, count in decision_census(score).items()
    }

    print(f"\n{name}  presentes={score.present}  (colunas anotadas no gabarito)")
    print(f"  CORRETO          {score.count(Outcome.CORRECT):3}      recall sobre presentes "
          f"{_pct(score.recall_on_present)}")
    print(f"  COLUNA_ERRADA    {score.count(Outcome.WRONG_COLUMN):3}")
    print(f"  PERDIDO          {score.count(Outcome.MISSED):3}")
    print(f"  FABRICADO        {fabricated:3}"
          + ("   <-- BLOQUEANTE" if fabricated else ""))
    print(f"  aceitos          {len(score.accepted):3}      precisão (limite inferior) "
          f"{_pct(score.accept_precision_lower)}")
    print(f"  ENTREGUES        {score.delivered_correct:3}      recall entregue "
          f"{_pct(score.delivered_recall)}   <-- o que o cliente recebe")
    if score.withheld_correct:
        print(f"  retidos          {len(score.withheld_correct):3}      coluna certa "
              "encontrada, pendente de revisão")
    print(f"  sem_veredito     {len(score.undetermined_accepts):3}      aceitos que o "
          "gabarito não consegue verificar")
    print(f"  aceites decididos por evidência objetiva {_pct(score.objective_decision_share)}"
          "   (entre os aceites com método automático registrado)")
    print(f"  aceitos por método vencedor {census}")
    for source, targets in key.source_conflicts.items():
        labels = ", ".join(f"m{target}" for target in targets)
        print(f"  REVISAR GABARITO: c{source} está indicada para {labels}; "
              "a aplicação exige uma origem distinta por destino.")
    if key.corrections:
        print(f"  correções do gabarito aplicadas: {sorted(key.corrections)}")
    for row in score.vacuous_targets:
        print(f"    FABRICAÇÃO m{row.model_column} {row.model_name!r} <- c{row.client_column}"
              f" (0 linhas preenchidas, final={_decimal(row.final_confidence)})")
    if show_losses:
        _report_losses(score, key)


def _report_losses(score: Score, key: AnswerKey) -> None:
    """Lista primeiro os acertos pendentes de revisão, depois as perdas e escolhas de coluna erradas."""
    for row in score.withheld_correct:
        status = schema.external_value("mappings", "status", row.status)
        print(f"    {'retido':14} m{row.model_column:<4} "
              f"{row.model_name[:30]:30} <- c{row.client_column} "
              f"final={_decimal(row.final_confidence)} ({status})")
    for row in score.rows:
        if row.outcome in (Outcome.MISSED, Outcome.WRONG_COLUMN):
            want = key.present.get(row.model_column)
            chosen = "nenhuma" if row.client_column is None else row.client_column
            print(f"    {OUTCOME_LABELS[row.outcome]:14} m{row.model_column:<4} "
                  f"{row.model_name[:30]:30} escolhida={chosen} esperada=c{want}")


def _report_totals(found: int, delivered: int, present: int, fabrications: int) -> None:
    print(f"\n{_TOTAL_LABEL}encontrados {found}/{present} "
          f"({schema.format_percent_br(found / present)})   "
          f"ENTREGUES {delivered}/{present} "
          f"({schema.format_percent_br(delivered / present)})   "
          f"fabricações {fabrications}")
    if found != delivered:
        # A dica também aparece quando --show-losses já foi passado.
        print(f"{' ' * len(_TOTAL_LABEL)}{found - delivered} mapeamentos "
              "corretos encontrados, mas não entregues - rode com --show-losses para "
              "listá-los")
